import csv
import errno
import io
import json
import os
import queue
import re
import shutil
import tempfile
import threading
from collections import Counter
from collections.abc import AsyncIterator
from contextlib import asynccontextmanager
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any
from urllib.parse import quote
from uuid import uuid4

from fastapi import BackgroundTasks, Body, Depends, FastAPI, File, Form, HTTPException, Query, Request, Response, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook
from PIL import Image
from pydantic import ValidationError
from sqlalchemy import func, or_
from sqlalchemy.orm import Session
from starlette.datastructures import FormData
from starlette.exceptions import HTTPException as StarletteHTTPException

from app.audit import log_audit_event
from app.config import PROJECT_ROOT, get_settings, parse_cors_allowed_origins, resolved_cors_allow_origin_regex
from app.database import SessionLocal, get_db, init_db
from app.document_processor import (
    DocumentProcessingError,
    is_supported_image,
    rasterize_document,
    rasterize_image_page,
    read_image_size,
    save_upload_file,
)
from app.document_modules import (
    classification_result_to_dict,
    required_field_result_to_dict,
    run_classification_batch,
    run_classification_job,
    run_required_field_check_batch,
    run_required_field_check_job,
)
from app.domain.module_job import ModuleJobLifecycle
from app.domain.workflow_run import WorkflowRunLifecycle
from app.extraction import result_to_dict, run_batch_jobs, run_extraction_job
from app.models import (
    AuditEvent,
    Batch,
    BatchItem,
    ClassificationBatch,
    ClassificationBatchItem,
    DocumentConversionJob,
    ClassificationJob,
    ClassificationResult,
    Document,
    DocumentClassifier,
    DocumentLibraryFolder,
    DocumentPage,
    ExportJob,
    ExportPreset,
    ExtractionJob,
    ExtractionResult,
    RawExtraction,
    RequiredFieldCheckBatch,
    RequiredFieldCheckBatchItem,
    RequiredFieldCheckJob,
    RequiredFieldCheckResult,
    RequiredFieldChecklist,
    Schema,
    WorkflowDefinition,
    WorkflowRun,
    WorkflowRunItem,
)
from app.raw_extractor import RawExtractionError, RawExtractionOptions, create_raw_outputs, save_raw_upload, validate_raw_upload
from app.repositories.sqlalchemy import SqlAlchemyWorkflowRunRepository
from app.routers.system import router as system_router
from app.services.workflow_runs import WorkflowRunApplicationService
from app.schemas import (
    ArchiveSearchResult,
    BatchInitRequest,
    AuditEventRead,
    BatchItemRead,
    BatchRead,
    ClassificationBatchInitRequest,
    ClassificationBatchItemRead,
    ClassificationBatchRead,
    ClassificationBatchFromDocumentsRequest,
    ClassificationJobCreate,
    ClassificationJobRead,
    ClassificationResultPatch,
    ClassificationResultRead,
    DocumentPageRead,
    DocumentBatchUploadRead,
    DocumentLibraryActionRead,
    DocumentLibraryClipboardRequest,
    DocumentLibraryCreateFolderRequest,
    DocumentRead,
    DocumentSelectionRequest,
    DocumentTreeRead,
    DocumentTreeFolderRead,
    DocumentClassifierCreate,
    DocumentClassifierRead,
    DocumentClassifierUpdate,
    DraftExtractionJobCreate,
    ExportJobCreate,
    ExportJobRead,
    ExportPresetCreate,
    ExportPresetRead,
    ExportPresetUpdate,
    ExtractionJobCreate,
    ExtractionJobRead,
    ExtractionResultPatch,
    RawExtractionRead,
    RawExtractionFromDocumentRequest,
    BatchFromDocumentsRequest,
    RequiredFieldCheckBatchInitRequest,
    RequiredFieldCheckBatchItemRead,
    RequiredFieldCheckBatchRead,
    RequiredFieldCheckBatchFromDocumentsRequest,
    RequiredFieldCheckJobCreate,
    RequiredFieldCheckJobRead,
    RequiredFieldCheckResultPatch,
    RequiredFieldCheckResultRead,
    RequiredFieldChecklistCreate,
    RequiredFieldChecklistRecommendationRead,
    RequiredFieldChecklistRecommendationRequest,
    RequiredFieldChecklistRead,
    RequiredFieldChecklistUpdate,
    SchemaCreate,
    SchemaDescriptionRecommendationRead,
    SchemaDescriptionRecommendationRequest,
    SchemaRecommendationRead,
    SchemaRecommendationRequest,
    SchemaRead,
    SchemaUpdate,
    WorkflowAiDraftRead,
    WorkflowRunEnqueueRequest,
    WorkflowDefinitionCreate,
    WorkflowDefinitionRead,
    WorkflowDefinitionUpdate,
    WorkflowRunInitRequest,
    WorkflowRunRead,
    WorkflowRunFromDocumentsRequest,
    WorkflowRunRestartRequest,
)
from app.vlm import (
    recommend_required_field_checklist_with_vlm,
    recommend_schema_description_with_vlm,
    recommend_schema_with_vlm,
    vlm_error_detail,
)
from app.workflows import (
    WorkflowDefinitionError,
    run_workflow_run,
    validate_workflow_definition,
    workflow_definition_to_read,
    workflow_run_export_csv,
    workflow_run_export_payload,
    workflow_run_to_read,
    workflow_runs_to_read,
)
from app.storage import delete_local_tree, delete_storage_ref, is_s3_ref, materialize_storage_ref, persist_artifact, scratch_dir_for_ref
from app.workspace import (
    current_workspace_id as _current_workspace_id,
    ensure_workspace_scope as _ensure_workspace_scope,
    require_workspace_admin_mode as _require_workspace_admin_mode,
    scope_query as _scope_query,
)


@asynccontextmanager
async def lifespan(_: FastAPI) -> AsyncIterator[None]:
    init_db()
    stop_conversion = _start_document_conversion_workers()
    stop_export = _start_export_job_worker()
    stop_cleanup = _start_retention_cleanup_worker()
    yield
    if stop_conversion:
        stop_conversion.set()
    if stop_export:
        stop_export.set()
    if stop_cleanup:
        stop_cleanup.set()


app = FastAPI(title="Document Automation Workspace API", version="0.1.0", lifespan=lifespan)

WORKFLOW_RUN_TERMINAL_STATUSES = {"completed", "completed_with_errors", "needs_review", "failed", "canceled"}
WORKFLOW_AI_DRAFT_MAX_IMAGES = 10
WORKFLOW_AI_DRAFT_IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg"}
WORKFLOW_ENQUEUE_BLOCKED_STATUSES = {"waiting", "failed", "canceled"}
DOCUMENT_READY_STATUSES = {"ready"}
DOCUMENT_CONVERTING_STATUSES = {"queued", "preprocessing"}
DOCUMENT_CONVERSION_TERMINAL_STATUSES = {"completed", "failed", "canceled"}
EXPORT_JOB_OWNER_TYPES = {"workflow_run", "batch", "classification_batch", "required_field_check_batch"}
EXPORT_JOB_WORKER_INTERVAL_SECONDS = 1.0
XLSX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_document_conversion_queue: queue.Queue[str] = queue.Queue()
_document_conversion_enqueued: set[str] = set()
_document_conversion_queue_lock = threading.Lock()


@dataclass(frozen=True)
class ExportArtifact:
    content: bytes
    filename: str
    content_type: str


@dataclass
class LocalUpload:
    filename: str
    content_type: str
    file: Any

settings = get_settings()
app.add_middleware(
    CORSMiddleware,
    allow_origins=parse_cors_allowed_origins(settings.cors_allowed_origins),
    allow_origin_regex=resolved_cors_allow_origin_regex(settings.cors_allow_origin_regex),
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)
app.include_router(system_router)


@app.middleware("http")
async def security_headers_middleware(request: Request, call_next):
    settings = get_settings()
    response = await call_next(request)
    if settings.security_headers_enabled:
        _apply_security_headers(response)
    return response


@app.post("/api/templates/bank-documents-poc/seed")
def seed_bank_documents_poc(request: Request, db: Session = Depends(get_db)) -> dict[str, Any]:
    workspace_id = _current_workspace_id(request)
    template = _load_bank_poc_template()
    created: dict[str, bool] = {}

    schema, created["schema"] = _seed_bank_poc_schema(db, template["schema"], workspace_id=workspace_id)
    classifier, created["classifier"] = _seed_bank_poc_classifier(db, template["classifier"], workspace_id=workspace_id)
    checklist, created["checklist"] = _seed_bank_poc_checklist(db, template["required_checklist"], workspace_id=workspace_id)
    workflow, created["workflow"] = _seed_bank_poc_workflow(db, template["workflow"], schema, classifier, checklist, workspace_id=workspace_id)
    documents, document_created = _seed_bank_poc_sample_documents(db, template.get("sample_documents", []), workspace_id=workspace_id)
    created["sample_documents"] = any(document_created.values())

    log_audit_event(
        db,
        entity_type="template",
        entity_id="bank_documents_poc",
        action="seeded",
        message="Seeded bank documents PoC template",
        metadata=created,
    )
    db.commit()
    db.refresh(schema)
    db.refresh(classifier)
    db.refresh(checklist)
    db.refresh(workflow)
    for document in documents:
        db.refresh(document)
    return {
        "template_key": template["name"],
        "created": created,
        "schema": _schema_read(schema),
        "classifier": _classifier_read(classifier),
        "checklist": _checklist_read(checklist),
        "workflow": WorkflowDefinitionRead(**workflow_definition_to_read(workflow, db)),
        "sample_document": _document_read(documents[0]) if documents else None,
        "sample_documents": [_bank_poc_sample_document_read(document) for document in documents],
    }

@app.post("/api/raw-extractions", response_model=RawExtractionRead)
def upload_raw_extraction(
    request: Request,
    file: UploadFile = File(...),
    include_images: bool = Form(default=True),
    include_formulas: bool = Form(default=False),
    db: Session = Depends(get_db),
) -> RawExtractionRead:
    workspace_id = _current_workspace_id(request)
    try:
        source_format = validate_raw_upload(file.filename or "")[1:]
    except RawExtractionError as exc:
        raise HTTPException(status_code=exc.status_code, detail=str(exc)) from exc

    raw = RawExtraction(
        workspace_id=workspace_id,
        filename=file.filename or "uploaded_document",
        source_format=source_format,
        size_bytes=0,
        status="processing",
    )
    db.add(raw)
    db.flush()

    try:
        filename, source_format, original_path, size_bytes = save_raw_upload(file, raw.id)
        raw.filename = filename
        raw.source_format = source_format
        raw.storage_path = str(original_path)
        raw.size_bytes = size_bytes
        pdf_path, html_path, warnings = create_raw_outputs(
            original_path,
            source_format,
            RawExtractionOptions(include_images=include_images, include_formulas=include_formulas),
        )
        if get_settings().storage_backend.strip().lower() == "s3":
            raw.storage_path = persist_artifact(original_path, f"raw/{raw.id}/original.{source_format}")
            raw.pdf_path = persist_artifact(pdf_path, f"raw/{raw.id}/preview.pdf", "application/pdf")
            raw.html_path = persist_artifact(html_path, f"raw/{raw.id}/content.html", "text/html; charset=utf-8")
        else:
            raw.pdf_path = str(pdf_path)
            raw.html_path = str(html_path)
        raw.warnings = json.dumps(warnings, ensure_ascii=False)
        raw.status = "completed"
        raw.error_message = None
    except RawExtractionError as exc:
        db.rollback()
        raise HTTPException(status_code=exc.status_code, detail=str(exc)) from exc
    except Exception as exc:
        raw.status = "failed"
        raw.error_message = str(exc)
    db.commit()
    db.refresh(raw)
    return _raw_extraction_read(raw)


@app.post("/api/raw-extractions/from-document", response_model=RawExtractionRead)
def create_raw_extraction_from_document(
    payload: RawExtractionFromDocumentRequest,
    request: Request,
    db: Session = Depends(get_db),
) -> RawExtractionRead:
    workspace_id = _current_workspace_id(request)
    document = db.get(Document, payload.document_id)
    _ensure_workspace_scope(document, workspace_id, "Document not found")
    if document.status == "deleted":
        raise HTTPException(status_code=410, detail="Original document was deleted")
    if document.status != "ready":
        raise HTTPException(status_code=409, detail={"message": "Document is not ready yet", "status": document.status})
    try:
        source_format = validate_raw_upload(document.filename)[1:]
    except RawExtractionError as exc:
        raise HTTPException(status_code=exc.status_code, detail=str(exc)) from exc

    raw = RawExtraction(
        workspace_id=workspace_id,
        filename=document.filename,
        source_format=source_format,
        size_bytes=document.size_bytes,
        storage_path=document.storage_path,
        status="processing",
    )
    db.add(raw)
    db.flush()

    try:
        source_path = materialize_storage_ref(document.storage_path)
        if not source_path.exists():
            raise RawExtractionError("Original document file is missing", status_code=404)
        work_dir = get_settings().resolved_raw_storage_dir / raw.id
        work_dir.mkdir(parents=True, exist_ok=True)
        suffix = source_path.suffix or f".{source_format}"
        working_original = work_dir / f"original{suffix}"
        if source_path != working_original:
            shutil.copy2(source_path, working_original)
        pdf_path, html_path, warnings = create_raw_outputs(
            working_original,
            source_format,
            RawExtractionOptions(include_images=payload.include_images, include_formulas=payload.include_formulas),
        )
        if get_settings().storage_backend.strip().lower() == "s3":
            raw.storage_path = persist_artifact(working_original, f"raw/{raw.id}/original.{source_format}")
            raw.pdf_path = persist_artifact(pdf_path, f"raw/{raw.id}/preview.pdf", "application/pdf")
            raw.html_path = persist_artifact(html_path, f"raw/{raw.id}/content.html", "text/html; charset=utf-8")
        else:
            raw.storage_path = str(working_original)
            raw.pdf_path = str(pdf_path)
            raw.html_path = str(html_path)
        raw.warnings = json.dumps(warnings, ensure_ascii=False)
        raw.status = "completed"
        raw.error_message = None
    except RawExtractionError as exc:
        db.rollback()
        raise HTTPException(status_code=exc.status_code, detail=str(exc)) from exc
    except Exception as exc:
        raw.status = "failed"
        raw.error_message = str(exc)
    db.commit()
    db.refresh(raw)
    return _raw_extraction_read(raw)


@app.get("/api/raw-extractions", response_model=list[RawExtractionRead])
def list_raw_extractions(
    request: Request,
    limit: int = Query(default=20, ge=1, le=100),
    db: Session = Depends(get_db),
) -> list[RawExtractionRead]:
    workspace_id = _current_workspace_id(request)
    rows = _scope_query(db.query(RawExtraction), RawExtraction, workspace_id).order_by(RawExtraction.created_at.desc()).limit(limit).all()
    return [_raw_extraction_read(row) for row in rows]


@app.get("/api/raw-extractions/{raw_id}/pdf")
def get_raw_extraction_pdf(raw_id: str, request: Request, db: Session = Depends(get_db)) -> FileResponse:
    workspace_id = _current_workspace_id(request)
    raw = db.get(RawExtraction, raw_id)
    _ensure_workspace_scope(raw, workspace_id, "Raw extraction not found")
    if not raw.pdf_path:
        raise HTTPException(status_code=404, detail="Raw extraction PDF preview is not available")
    path = materialize_storage_ref(raw.pdf_path)
    if not path.exists():
        raise HTTPException(status_code=404, detail="Raw extraction PDF preview is missing")
    return FileResponse(path, media_type="application/pdf")


@app.get("/api/raw-extractions/{raw_id}/html")
def get_raw_extraction_html(raw_id: str, request: Request, db: Session = Depends(get_db)) -> FileResponse:
    workspace_id = _current_workspace_id(request)
    raw = db.get(RawExtraction, raw_id)
    _ensure_workspace_scope(raw, workspace_id, "Raw extraction not found")
    if not raw.html_path:
        raise HTTPException(status_code=404, detail="Raw extraction HTML is not available")
    path = materialize_storage_ref(raw.html_path)
    if not path.exists():
        raise HTTPException(status_code=404, detail="Raw extraction HTML is missing")
    return FileResponse(path, media_type="text/html")


@app.get("/api/raw-extractions/{raw_id}", response_model=RawExtractionRead)
def get_raw_extraction(raw_id: str, request: Request, db: Session = Depends(get_db)) -> RawExtractionRead:
    workspace_id = _current_workspace_id(request)
    raw = db.get(RawExtraction, raw_id)
    _ensure_workspace_scope(raw, workspace_id, "Raw extraction not found")
    return _raw_extraction_read(raw)


@app.post("/api/documents", response_model=DocumentRead)
def upload_document(request: Request, file: UploadFile = File(...), db: Session = Depends(get_db)) -> DocumentRead:
    workspace_id = _current_workspace_id(request)
    document = _create_document_from_upload(file, db, workspace_id=workspace_id)
    if not document.library_path:
        document.library_path = document.filename
    log_audit_event(
        db,
        entity_type="document",
        entity_id=document.id,
        action="upload_processed",
        message=f"Processed {document.filename} in the document library",
        metadata={"filename": document.filename, "library_path": document.library_path},
    )
    db.commit()
    db.refresh(document)
    return _document_read(document)


@app.post("/api/library/uploads", response_model=DocumentBatchUploadRead)
async def upload_library_documents(request: Request, db: Session = Depends(get_db)) -> DocumentBatchUploadRead:
    workspace_id = _current_workspace_id(request)
    form, files = await _read_batch_upload_form(request)
    if not files:
        raise HTTPException(status_code=400, detail="At least one file is required")
    _validate_upload_file_count(files)
    library_paths = _upload_library_paths(form, files)
    documents: list[Document] = []
    conversion_job_ids: list[str] = []
    try:
        for index, file in enumerate(files):
            document, conversion_job_id = _create_queued_library_document(file, db, library_path=library_paths[index], workspace_id=workspace_id)
            documents.append(document)
            conversion_job_ids.append(conversion_job_id)
            log_audit_event(
                db,
                entity_type="document",
                entity_id=document.id,
                action="upload_queued",
                message=f"Queued {document.filename} for document conversion",
                metadata={"filename": document.filename, "library_path": document.library_path, "conversion_job_id": conversion_job_id},
            )
        db.commit()
    except DocumentProcessingError as exc:
        db.rollback()
        raise HTTPException(status_code=exc.status_code, detail=str(exc)) from exc
    except Exception:
        db.rollback()
        raise
    finally:
        for file in files:
            await file.close()
    for conversion_job_id in conversion_job_ids:
        _enqueue_document_conversion_job(conversion_job_id)
    for document in documents:
        db.refresh(document)
    return DocumentBatchUploadRead(documents=[_document_read(document) for document in documents])


@app.get("/api/documents", response_model=list[DocumentRead])
def list_documents(
    request: Request,
    limit: int = Query(default=50, ge=1, le=500),
    offset: int = Query(default=0, ge=0),
    q: str | None = Query(default=None),
    library_path: str | None = Query(default=None),
    status: str | None = Query(default=None),
    include_deleted: bool = Query(default=False),
    db: Session = Depends(get_db),
) -> list[DocumentRead]:
    workspace_id = _current_workspace_id(request)
    query = _document_library_query(db, include_deleted=include_deleted, status=status, q=q, library_path=library_path, workspace_id=workspace_id)
    documents = query.order_by(Document.created_at.desc(), Document.id.desc()).offset(offset).limit(limit).all()
    return [_document_read(document) for document in documents]


@app.get("/api/documents/ids", response_model=list[str])
def list_document_ids(
    request: Request,
    limit: int = Query(default=10000, ge=1, le=20000),
    q: str | None = Query(default=None),
    library_path: str | None = Query(default=None),
    status: str | None = Query(default=None),
    include_deleted: bool = Query(default=False),
    db: Session = Depends(get_db),
) -> list[str]:
    workspace_id = _current_workspace_id(request)
    query = _document_library_query(db, include_deleted=include_deleted, status=status, q=q, library_path=library_path, workspace_id=workspace_id)
    rows = query.order_by(Document.created_at.desc(), Document.id.desc()).with_entities(Document.id).limit(limit).all()
    return [row[0] for row in rows]


@app.post("/api/documents/selection", response_model=DocumentBatchUploadRead)
def get_selected_documents(payload: DocumentSelectionRequest, request: Request, db: Session = Depends(get_db)) -> DocumentBatchUploadRead:
    workspace_id = _current_workspace_id(request)
    documents = _selected_library_documents(db, list(dict.fromkeys(payload.document_ids)), workspace_id=workspace_id)
    return DocumentBatchUploadRead(documents=[_document_read(document) for document in documents])


@app.post("/api/documents/delete", response_model=DocumentBatchUploadRead)
def delete_selected_documents_from_library(payload: DocumentSelectionRequest, request: Request, db: Session = Depends(get_db)) -> DocumentBatchUploadRead:
    workspace_id = _current_workspace_id(request)
    document_ids = list(dict.fromkeys(payload.document_ids))
    documents = _scope_query(db.query(Document), Document, workspace_id).filter(Document.id.in_(document_ids)).all()
    by_id = {document.id: document for document in documents}
    missing = [document_id for document_id in document_ids if document_id not in by_id]
    if missing:
        raise HTTPException(status_code=404, detail=f"Document not found: {missing[0]}")
    deleted_documents = [_delete_library_document_payload(by_id[document_id], db) for document_id in document_ids]
    log_audit_event(
        db,
        entity_type="document",
        entity_id="bulk",
        action="bulk_deleted",
        message=f"Deleted original payloads for {len(deleted_documents)} document(s)",
        metadata={"document_ids": document_ids},
    )
    db.commit()
    for document in deleted_documents:
        db.refresh(document)
    return DocumentBatchUploadRead(documents=[_document_read(document) for document in deleted_documents])


@app.get("/api/library/tree", response_model=DocumentTreeRead)
def get_document_library_tree(request: Request, include_deleted: bool = Query(default=False), db: Session = Depends(get_db)) -> DocumentTreeRead:
    workspace_id = _current_workspace_id(request)
    return _document_tree_read(db, include_deleted=include_deleted, workspace_id=workspace_id)


@app.post("/api/library/folders", response_model=DocumentTreeRead)
def create_library_folder(payload: DocumentLibraryCreateFolderRequest, request: Request, db: Session = Depends(get_db)) -> DocumentTreeRead:
    workspace_id = _current_workspace_id(request)
    folder_path = _normalize_library_path(payload.folder_path)
    if not folder_path:
        raise HTTPException(status_code=422, detail="Folder path is required")
    _ensure_library_folder_records(db, folder_path, workspace_id=workspace_id)
    log_audit_event(
        db,
        entity_type="document_library_folder",
        entity_id=folder_path,
        action="created",
        message=f"Created document library folder {folder_path}",
        metadata={"path": folder_path},
    )
    db.commit()
    return _document_tree_read(db, workspace_id=workspace_id)


@app.post("/api/library/move", response_model=DocumentLibraryActionRead)
def move_library_entries(payload: DocumentLibraryClipboardRequest, request: Request, db: Session = Depends(get_db)) -> DocumentLibraryActionRead:
    workspace_id = _current_workspace_id(request)
    result = _move_library_entries(payload, db, workspace_id=workspace_id)
    db.commit()
    return result


@app.post("/api/library/copy", response_model=DocumentLibraryActionRead)
def copy_library_entries(payload: DocumentLibraryClipboardRequest, request: Request, db: Session = Depends(get_db)) -> DocumentLibraryActionRead:
    workspace_id = _current_workspace_id(request)
    result, conversion_job_ids = _copy_library_entries(payload, db, workspace_id=workspace_id)
    db.commit()
    for conversion_job_id in conversion_job_ids:
        _enqueue_document_conversion_job(conversion_job_id)
    return result


def _document_tree_read(db: Session, *, include_deleted: bool = False, workspace_id: str | None = None) -> DocumentTreeRead:
    query = _scope_query(db.query(Document), Document, workspace_id)
    if not include_deleted:
        query = query.filter(Document.status != "deleted")
    documents = query.all()
    folders: dict[str, dict[str, Any]] = {}

    def ensure_folder(path: str) -> dict[str, Any]:
        normalized = _normalize_library_path(path)
        if normalized not in folders:
            parent = normalized.rsplit("/", 1)[0] if "/" in normalized else None
            folders[normalized] = {
                "path": normalized,
                "name": normalized.rsplit("/", 1)[-1] if normalized else "문서 보관함",
                "parent": parent,
                "total_count": 0,
                "ready_count": 0,
                "converting_count": 0,
                "failed_count": 0,
                "deleted_count": 0,
            }
            if parent is not None:
                ensure_folder(parent)
        return folders[normalized]

    ensure_folder("")
    for explicit_folder in _scope_query(db.query(DocumentLibraryFolder), DocumentLibraryFolder, workspace_id).order_by(DocumentLibraryFolder.path.asc()).all():
        ensure_folder(explicit_folder.path)

    for document in documents:
        for folder_path in _library_folder_ancestors(_library_folder_path(document.library_path)):
            folder = ensure_folder(folder_path)
            folder["total_count"] += 1
            if document.status == "ready":
                folder["ready_count"] += 1
            elif document.status in DOCUMENT_CONVERTING_STATUSES:
                folder["converting_count"] += 1
            elif document.status == "deleted":
                folder["deleted_count"] += 1
            elif document.status == "failed":
                folder["failed_count"] += 1
    return DocumentTreeRead(folders=[DocumentTreeFolderRead(**folders[path]) for path in sorted(folders)])


@app.get("/api/documents/{document_id}", response_model=DocumentRead)
def get_document(document_id: str, request: Request, db: Session = Depends(get_db)) -> DocumentRead:
    workspace_id = _current_workspace_id(request)
    document = db.get(Document, document_id)
    _ensure_workspace_scope(document, workspace_id, "Document not found")
    if document.status == "ready":
        _repair_image_document_if_needed(document, db)
    return _document_read(document)


@app.delete("/api/documents/{document_id}", response_model=DocumentRead)
def delete_document_from_library(document_id: str, request: Request, db: Session = Depends(get_db)) -> DocumentRead:
    workspace_id = _current_workspace_id(request)
    document = db.get(Document, document_id)
    _ensure_workspace_scope(document, workspace_id, "Document not found")
    _delete_library_document_payload(document, db)
    log_audit_event(
        db,
        entity_type="document",
        entity_id=document.id,
        action="deleted",
        message="Deleted original document payload; historical results were kept",
        metadata={"filename": document.filename, "library_path": document.library_path},
    )
    db.commit()
    db.refresh(document)
    return _document_read(document)


@app.get("/api/documents/{document_id}/pages/{page_number}/image")
def get_document_page_image(document_id: str, page_number: int, request: Request, db: Session = Depends(get_db)) -> FileResponse:
    workspace_id = _current_workspace_id(request)
    document = db.get(Document, document_id)
    _ensure_workspace_scope(document, workspace_id, "Document not found")
    if document and document.status == "deleted":
        raise HTTPException(status_code=410, detail="Original document was deleted")
    page = (
        db.query(DocumentPage)
        .filter(DocumentPage.document_id == document_id, DocumentPage.page_number == page_number)
        .one_or_none()
    )
    if not page:
        raise HTTPException(status_code=404, detail="Document page not found")
    if document and document.status == "ready":
        _repair_image_document_if_needed(document, db)
        db.refresh(page)
    path = materialize_storage_ref(page.image_path)
    if not path.exists():
        raise HTTPException(status_code=404, detail="Document page image missing")
    return FileResponse(path, media_type=_image_media_type(path))


@app.get("/api/documents/{document_id}/pages/{page_number}/thumbnail")
def get_document_page_thumbnail(
    document_id: str,
    page_number: int,
    request: Request,
    width: int = Query(default=96, ge=48, le=512),
    db: Session = Depends(get_db),
) -> FileResponse:
    workspace_id = _current_workspace_id(request)
    document = db.get(Document, document_id)
    _ensure_workspace_scope(document, workspace_id, "Document not found")
    if document and document.status == "deleted":
        raise HTTPException(status_code=410, detail="Original document was deleted")
    page = (
        db.query(DocumentPage)
        .filter(DocumentPage.document_id == document_id, DocumentPage.page_number == page_number)
        .one_or_none()
    )
    if not page:
        raise HTTPException(status_code=404, detail="Document page not found")
    if document and document.status == "ready":
        _repair_image_document_if_needed(document, db)
        db.refresh(page)
    source_path = materialize_storage_ref(page.image_path)
    if not source_path.exists():
        raise HTTPException(status_code=404, detail="Document page image missing")

    thumbnail_path = source_path.with_name(f"{source_path.stem}_thumb_{width}.jpg")
    if not thumbnail_path.exists() or thumbnail_path.stat().st_mtime < source_path.stat().st_mtime:
        with Image.open(source_path) as source:
            image = source.convert("RGB")
            ratio = width / max(1, image.width)
            target_size = (width, max(1, round(image.height * ratio)))
            image = image.resize(target_size, Image.Resampling.LANCZOS)
            image.save(thumbnail_path, format="JPEG", quality=82, optimize=True)

    return FileResponse(thumbnail_path, media_type="image/jpeg")


@app.post("/api/schemas", response_model=SchemaRead)
def create_schema(payload: SchemaCreate, request: Request, db: Session = Depends(get_db)) -> SchemaRead:
    workspace_id = _current_workspace_id(request)
    _raise_if_schema_name_conflicts(db, payload.name, workspace_id=workspace_id)
    schema = Schema(
        workspace_id=workspace_id,
        name=payload.name,
        display_name=payload.display_name,
        description=payload.description,
        current_version=1,
        schema_json=json.dumps(payload.model_dump(), ensure_ascii=False),
        is_template=payload.is_template,
        template_category=payload.template_category,
        pinned=payload.pinned,
        ephemeral=False,
        archived=False,
    )
    db.add(schema)
    db.flush()
    schema_json = payload.model_dump()
    _validate_schema_region_references(schema_json)
    log_audit_event(
        db,
        entity_type="schema",
        entity_id=schema.id,
        action="created",
        message=f"Created schema {schema.name}",
        metadata={"is_template": schema.is_template, "field_count": len(payload.fields)},
    )
    db.commit()
    db.refresh(schema)
    return _schema_read(schema)


@app.get("/api/schemas", response_model=list[SchemaRead])
def list_schemas(
    request: Request,
    templates: bool | None = None,
    include_ephemeral: bool = False,
    include_archived: bool = False,
    db: Session = Depends(get_db),
) -> list[SchemaRead]:
    workspace_id = _current_workspace_id(request)
    query = db.query(Schema)
    if workspace_id is not None:
        if templates is True:
            query = query.filter(or_(Schema.workspace_id == workspace_id, Schema.workspace_id.is_(None)))
        else:
            query = query.filter(Schema.workspace_id == workspace_id)
    if not include_ephemeral:
        query = query.filter(Schema.ephemeral == False)  # noqa: E712
    if not include_archived:
        query = query.filter(Schema.archived == False)  # noqa: E712
    if templates is not None:
        query = query.filter(Schema.is_template == templates)
    schemas = query.order_by(Schema.pinned.desc(), Schema.created_at.desc()).all()
    return [_schema_read(schema) for schema in schemas]


@app.post("/api/schemas/recommendations", response_model=SchemaRecommendationRead)
def recommend_schema(
    payload: SchemaRecommendationRequest,
    request: Request,
    db: Session = Depends(get_db),
) -> SchemaRecommendationRead:
    workspace_id = _current_workspace_id(request)
    document = db.get(Document, payload.document_id)
    _ensure_workspace_scope(document, workspace_id, "Document not found")
    try:
        recommendation = recommend_schema_with_vlm([page.image_path for page in document.pages])
        recommendation_read = _schema_recommendation_read(recommendation)
        document.document_type = recommendation_read.document_type
        document.language = recommendation_read.language
        document.ai_summary = recommendation_read.description
        document.recommendation_reasoning = recommendation_read.reasoning
        log_audit_event(
            db,
            entity_type="document",
            entity_id=document.id,
            action="schema_recommended",
            message="AI schema recommendation generated",
            metadata={
                "document_type": recommendation_read.document_type,
                "language": recommendation_read.language,
                "field_count": len(recommendation_read.fields),
            },
        )
        db.commit()
        return recommendation_read
    except ValidationError as exc:
        raise HTTPException(status_code=502, detail=f"VLM returned an invalid schema recommendation: {exc}") from exc
    except RuntimeError as exc:
        raise HTTPException(status_code=400, detail=vlm_error_detail(exc)) from exc


@app.post("/api/schemas/description-recommendations", response_model=SchemaDescriptionRecommendationRead)
def recommend_schema_description(
    payload: SchemaDescriptionRecommendationRequest,
    request: Request,
    db: Session = Depends(get_db),
) -> SchemaDescriptionRecommendationRead:
    workspace_id = _current_workspace_id(request)
    image_paths: list[str] = []
    if payload.document_id:
        document = db.get(Document, payload.document_id)
        _ensure_workspace_scope(document, workspace_id, "Document not found")
        image_paths = [page.image_path for page in document.pages]
    try:
        recommendation = recommend_schema_description_with_vlm(
            image_paths,
            schema_name=payload.name,
            current_description=payload.current_description,
            fields=payload.fields,
        )
        return SchemaDescriptionRecommendationRead(**recommendation)
    except ValidationError as exc:
        raise HTTPException(status_code=502, detail=f"VLM returned an invalid schema description recommendation: {exc}") from exc
    except RuntimeError as exc:
        raise HTTPException(status_code=400, detail=vlm_error_detail(exc)) from exc


@app.get("/api/schemas/{schema_id}", response_model=SchemaRead)
def get_schema(schema_id: str, request: Request, db: Session = Depends(get_db)) -> SchemaRead:
    workspace_id = _current_workspace_id(request)
    schema = db.get(Schema, schema_id)
    _ensure_workspace_scope(schema, workspace_id, "Schema not found")
    return _schema_read(schema)


@app.post("/api/schemas/{schema_id}/duplicate", response_model=SchemaRead)
def duplicate_schema(schema_id: str, request: Request, db: Session = Depends(get_db)) -> SchemaRead:
    workspace_id = _current_workspace_id(request)
    schema = db.get(Schema, schema_id)
    if not schema:
        raise HTTPException(status_code=404, detail="Schema not found")
    if workspace_id is not None and schema.workspace_id not in {workspace_id, None}:
        raise HTTPException(status_code=404, detail="Schema not found")
    if schema.ephemeral:
        raise HTTPException(status_code=400, detail="Draft schemas cannot be duplicated from the library")
    schema_data = _schema_data(schema)
    existing_names = {
        row[0]
        for row in _scope_query(
            db.query(Schema.name).filter(Schema.ephemeral == False, Schema.archived == False),  # noqa: E712
            Schema,
            workspace_id,
        ).all()
    }
    duplicated_name = _duplicate_name(schema.name, existing_names)
    duplicated_data = {
        **schema_data,
        "name": duplicated_name,
        "display_name": duplicated_name,
    }
    duplicated = Schema(
        workspace_id=workspace_id,
        name=duplicated_name,
        display_name=duplicated_name,
        description=schema.description,
        current_version=1,
        schema_json=json.dumps(duplicated_data, ensure_ascii=False),
        is_template=schema.is_template,
        template_category=schema.template_category,
        pinned=schema.pinned,
        ephemeral=False,
        archived=False,
    )
    db.add(duplicated)
    db.flush()
    log_audit_event(
        db,
        entity_type="schema",
        entity_id=duplicated.id,
        action="duplicated",
        message=f"Duplicated schema {schema.name} to {duplicated.name}",
        metadata={"source_schema_id": schema.id, "field_count": len(duplicated_data["fields"])},
    )
    db.commit()
    db.refresh(duplicated)
    return _schema_read(duplicated)


@app.patch("/api/schemas/{schema_id}", response_model=SchemaRead)
def update_schema(schema_id: str, payload: SchemaUpdate, request: Request, db: Session = Depends(get_db)) -> SchemaRead:
    workspace_id = _current_workspace_id(request)
    schema = db.get(Schema, schema_id)
    _ensure_workspace_scope(schema, workspace_id, "Schema not found")

    current = _schema_data(schema)
    next_schema_data = {
        "name": payload.name if payload.name is not None else schema.name,
        "display_name": (
            payload.display_name if "display_name" in payload.model_fields_set else schema.display_name
        ),
        "description": payload.description if "description" in payload.model_fields_set else schema.description,
        "is_template": payload.is_template if payload.is_template is not None else schema.is_template,
        "template_category": (
            payload.template_category if "template_category" in payload.model_fields_set else schema.template_category
        ),
        "pinned": payload.pinned if payload.pinned is not None else schema.pinned,
        "regions": [region.model_dump() for region in payload.regions] if payload.regions is not None else current.get("regions", []),
        "fields": [field.model_dump() for field in payload.fields] if payload.fields is not None else current["fields"],
    }
    _validate_schema_region_references(next_schema_data)
    if next_schema_data["name"].strip() == schema.name.strip():
        _merge_duplicate_schema_names_into(db, schema, next_schema_data["name"], workspace_id=workspace_id)
    else:
        _raise_if_schema_name_conflicts(db, next_schema_data["name"], schema_id=schema.id, workspace_id=workspace_id)

    schema.name = next_schema_data["name"]
    schema.display_name = next_schema_data["display_name"]
    schema.description = next_schema_data["description"]
    schema.is_template = next_schema_data["is_template"]
    schema.template_category = next_schema_data["template_category"]
    schema.pinned = next_schema_data["pinned"]
    schema.schema_json = json.dumps(next_schema_data, ensure_ascii=False)
    log_audit_event(
        db,
        entity_type="schema",
        entity_id=schema.id,
        action="updated",
        message=f"Updated schema {schema.name}",
        metadata={
            "is_template": schema.is_template,
            "field_count": len(next_schema_data["fields"]),
        },
    )
    db.commit()
    db.refresh(schema)
    return _schema_read(schema)


@app.delete("/api/schemas/{schema_id}", response_model=SchemaRead)
def delete_schema(schema_id: str, request: Request, db: Session = Depends(get_db)) -> SchemaRead:
    workspace_id = _current_workspace_id(request)
    schema = db.get(Schema, schema_id)
    _ensure_workspace_scope(schema, workspace_id, "Schema not found")
    if schema.ephemeral:
        raise HTTPException(status_code=400, detail="Draft schemas cannot be archived from the library")

    schema.archived = True
    schema.pinned = False
    schema.is_template = False
    log_audit_event(
        db,
        entity_type="schema",
        entity_id=schema.id,
        action="archived",
        message=f"Archived schema {schema.name}",
        metadata={"name": schema.name},
    )
    db.commit()
    db.refresh(schema)
    return _schema_read(schema)


@app.post("/api/document-classifiers", response_model=DocumentClassifierRead)
def create_document_classifier(payload: DocumentClassifierCreate, request: Request, db: Session = Depends(get_db)) -> DocumentClassifierRead:
    workspace_id = _current_workspace_id(request)
    classifier = DocumentClassifier(
        workspace_id=workspace_id,
        name=payload.name,
        description=payload.description,
        allow_unknown=payload.allow_unknown,
        config_json=json.dumps(payload.model_dump(), ensure_ascii=False),
        archived=False,
    )
    db.add(classifier)
    db.flush()
    log_audit_event(
        db,
        entity_type="document_classifier",
        entity_id=classifier.id,
        action="created",
        message=f"Created document classifier {classifier.name}",
        metadata={"class_count": len(payload.classes)},
    )
    db.commit()
    db.refresh(classifier)
    return _classifier_read(classifier)


@app.get("/api/document-classifiers", response_model=list[DocumentClassifierRead])
def list_document_classifiers(
    request: Request,
    include_archived: bool = False,
    db: Session = Depends(get_db),
) -> list[DocumentClassifierRead]:
    workspace_id = _current_workspace_id(request)
    query = db.query(DocumentClassifier)
    query = _scope_query(query, DocumentClassifier, workspace_id)
    if not include_archived:
        query = query.filter(DocumentClassifier.archived == False)  # noqa: E712
    rows = query.order_by(DocumentClassifier.created_at.desc()).all()
    return [_classifier_read(row) for row in rows]


@app.get("/api/document-classifiers/{classifier_id}", response_model=DocumentClassifierRead)
def get_document_classifier(classifier_id: str, request: Request, db: Session = Depends(get_db)) -> DocumentClassifierRead:
    workspace_id = _current_workspace_id(request)
    classifier = db.get(DocumentClassifier, classifier_id)
    _ensure_workspace_scope(classifier, workspace_id, "Document classifier not found")
    return _classifier_read(classifier)


@app.post("/api/document-classifiers/{classifier_id}/duplicate", response_model=DocumentClassifierRead)
def duplicate_document_classifier(classifier_id: str, request: Request, db: Session = Depends(get_db)) -> DocumentClassifierRead:
    workspace_id = _current_workspace_id(request)
    classifier = db.get(DocumentClassifier, classifier_id)
    _ensure_workspace_scope(classifier, workspace_id, "Document classifier not found")
    config = _classifier_data(classifier)
    existing_names = {
        row[0]
        for row in _scope_query(
            db.query(DocumentClassifier.name).filter(DocumentClassifier.archived == False),  # noqa: E712
            DocumentClassifier,
            workspace_id,
        ).all()
    }
    duplicated_name = _duplicate_name(classifier.name, existing_names)
    duplicated_config = {
        **config,
        "name": duplicated_name,
    }
    duplicated = DocumentClassifier(
        workspace_id=workspace_id,
        name=duplicated_name,
        description=classifier.description,
        allow_unknown=classifier.allow_unknown,
        config_json=json.dumps(duplicated_config, ensure_ascii=False),
        archived=False,
    )
    db.add(duplicated)
    db.flush()
    log_audit_event(
        db,
        entity_type="document_classifier",
        entity_id=duplicated.id,
        action="duplicated",
        message=f"Duplicated document classifier {classifier.name} to {duplicated.name}",
        metadata={"source_classifier_id": classifier.id, "class_count": len(duplicated_config["classes"])},
    )
    db.commit()
    db.refresh(duplicated)
    return _classifier_read(duplicated)


@app.patch("/api/document-classifiers/{classifier_id}", response_model=DocumentClassifierRead)
def update_document_classifier(
    classifier_id: str,
    payload: DocumentClassifierUpdate,
    request: Request,
    db: Session = Depends(get_db),
) -> DocumentClassifierRead:
    workspace_id = _current_workspace_id(request)
    classifier = db.get(DocumentClassifier, classifier_id)
    _ensure_workspace_scope(classifier, workspace_id, "Document classifier not found")
    current = _classifier_data(classifier)
    next_config = {
        "name": payload.name if payload.name is not None else classifier.name,
        "description": payload.description if "description" in payload.model_fields_set else classifier.description,
        "allow_unknown": payload.allow_unknown if payload.allow_unknown is not None else classifier.allow_unknown,
        "classes": [item.model_dump() for item in payload.classes] if payload.classes is not None else current["classes"],
    }
    classifier.name = next_config["name"]
    classifier.description = next_config["description"]
    classifier.allow_unknown = bool(next_config["allow_unknown"])
    classifier.config_json = json.dumps(next_config, ensure_ascii=False)
    log_audit_event(
        db,
        entity_type="document_classifier",
        entity_id=classifier.id,
        action="updated",
        message=f"Updated document classifier {classifier.name}",
        metadata={"class_count": len(next_config["classes"])},
    )
    db.commit()
    db.refresh(classifier)
    return _classifier_read(classifier)


@app.delete("/api/document-classifiers/{classifier_id}", response_model=DocumentClassifierRead)
def delete_document_classifier(classifier_id: str, request: Request, db: Session = Depends(get_db)) -> DocumentClassifierRead:
    workspace_id = _current_workspace_id(request)
    classifier = db.get(DocumentClassifier, classifier_id)
    _ensure_workspace_scope(classifier, workspace_id, "Document classifier not found")
    classifier.archived = True
    log_audit_event(
        db,
        entity_type="document_classifier",
        entity_id=classifier.id,
        action="archived",
        message=f"Archived document classifier {classifier.name}",
        metadata={"name": classifier.name},
    )
    db.commit()
    db.refresh(classifier)
    return _classifier_read(classifier)


@app.post("/api/classification-jobs", response_model=ClassificationJobRead)
def create_classification_job(
    payload: ClassificationJobCreate,
    background_tasks: BackgroundTasks,
    request: Request,
    db: Session = Depends(get_db),
) -> ClassificationJobRead:
    workspace_id = _current_workspace_id(request)
    document = db.get(Document, payload.document_id)
    classifier = db.get(DocumentClassifier, payload.classifier_id)
    _ensure_workspace_scope(document, workspace_id, "Document not found")
    if not classifier or classifier.archived or (workspace_id is not None and classifier.workspace_id != workspace_id):
        raise HTTPException(status_code=404, detail="Document classifier not found")
    status = _execution_job_status_for_document(document)
    if status == "blocked":
        _raise_document_not_executable(document)
    if document.status == "ready":
        _repair_image_document_if_needed(document, db)
    job = ClassificationJob(workspace_id=workspace_id, document_id=document.id, classifier_id=classifier.id, status=status)
    db.add(job)
    db.flush()
    log_audit_event(
        db,
        entity_type="classification_job",
        entity_id=job.id,
        action="created",
        message="Classification job created",
        metadata={"document_id": document.id, "classifier_id": classifier.id},
    )
    db.commit()
    db.refresh(job)
    response = _classification_job_read(job)
    if status == "queued":
        db.close()
        background_tasks.add_task(run_classification_job, job.id)
    return response


@app.get("/api/classification-jobs/{job_id}", response_model=ClassificationJobRead)
def get_classification_job(job_id: str, request: Request, db: Session = Depends(get_db)) -> ClassificationJobRead:
    workspace_id = _current_workspace_id(request)
    job = db.get(ClassificationJob, job_id)
    _ensure_workspace_scope(job, workspace_id, "Classification job not found")
    return _classification_job_read(job)


@app.patch("/api/classification-results/{result_id}", response_model=ClassificationResultRead)
def patch_classification_result(
    result_id: str,
    payload: ClassificationResultPatch,
    request: Request,
    db: Session = Depends(get_db),
) -> ClassificationResultRead:
    workspace_id = _current_workspace_id(request)
    result = db.get(ClassificationResult, result_id)
    if not result:
        raise HTTPException(status_code=404, detail="Classification result not found")
    _ensure_workspace_scope(result.job, workspace_id, "Classification result not found")
    if payload.corrected_output is not None:
        result.corrected_output = json.dumps(payload.corrected_output, ensure_ascii=False)
    if payload.reviewed is not None:
        result.reviewed = payload.reviewed
    db.commit()
    db.refresh(result)
    return ClassificationResultRead(**classification_result_to_dict(result))


@app.post("/api/required-field-checklists", response_model=RequiredFieldChecklistRead)
def create_required_field_checklist(payload: RequiredFieldChecklistCreate, request: Request, db: Session = Depends(get_db)) -> RequiredFieldChecklistRead:
    workspace_id = _current_workspace_id(request)
    checklist = RequiredFieldChecklist(
        workspace_id=workspace_id,
        name=payload.name,
        description=payload.description,
        config_json=json.dumps(payload.model_dump(), ensure_ascii=False),
        archived=False,
    )
    db.add(checklist)
    db.flush()
    log_audit_event(
        db,
        entity_type="required_field_checklist",
        entity_id=checklist.id,
        action="created",
        message=f"Created required field checklist {checklist.name}",
        metadata={"item_count": len(payload.items)},
    )
    db.commit()
    db.refresh(checklist)
    return _checklist_read(checklist)


@app.get("/api/required-field-checklists", response_model=list[RequiredFieldChecklistRead])
def list_required_field_checklists(
    request: Request,
    include_archived: bool = False,
    db: Session = Depends(get_db),
) -> list[RequiredFieldChecklistRead]:
    workspace_id = _current_workspace_id(request)
    query = db.query(RequiredFieldChecklist)
    query = _scope_query(query, RequiredFieldChecklist, workspace_id)
    if not include_archived:
        query = query.filter(RequiredFieldChecklist.archived == False)  # noqa: E712
    rows = query.order_by(RequiredFieldChecklist.created_at.desc()).all()
    return [_checklist_read(row) for row in rows]


@app.post("/api/required-field-checklists/recommendations", response_model=RequiredFieldChecklistRecommendationRead)
def recommend_required_field_checklist(
    payload: RequiredFieldChecklistRecommendationRequest,
    request: Request,
    db: Session = Depends(get_db),
) -> RequiredFieldChecklistRecommendationRead:
    workspace_id = _current_workspace_id(request)
    document = db.get(Document, payload.document_id)
    _ensure_workspace_scope(document, workspace_id, "Document not found")
    try:
        recommendation = recommend_required_field_checklist_with_vlm([page.image_path for page in document.pages])
        recommendation_read = _required_field_checklist_recommendation_read(recommendation)
        log_audit_event(
            db,
            entity_type="document",
            entity_id=document.id,
            action="required_field_checklist_recommended",
            message="AI required field checklist recommendation generated",
            metadata={
                "item_count": len(recommendation_read.items),
                "region_count": len(recommendation_read.regions),
            },
        )
        db.commit()
        return recommendation_read
    except (ValidationError, ValueError) as exc:
        raise HTTPException(status_code=502, detail=f"VLM returned an invalid checklist recommendation: {exc}") from exc
    except RuntimeError as exc:
        raise HTTPException(status_code=400, detail=vlm_error_detail(exc)) from exc


@app.get("/api/required-field-checklists/{checklist_id}", response_model=RequiredFieldChecklistRead)
def get_required_field_checklist(checklist_id: str, request: Request, db: Session = Depends(get_db)) -> RequiredFieldChecklistRead:
    workspace_id = _current_workspace_id(request)
    checklist = db.get(RequiredFieldChecklist, checklist_id)
    _ensure_workspace_scope(checklist, workspace_id, "Required field checklist not found")
    return _checklist_read(checklist)


@app.post("/api/required-field-checklists/{checklist_id}/duplicate", response_model=RequiredFieldChecklistRead)
def duplicate_required_field_checklist(checklist_id: str, request: Request, db: Session = Depends(get_db)) -> RequiredFieldChecklistRead:
    workspace_id = _current_workspace_id(request)
    checklist = db.get(RequiredFieldChecklist, checklist_id)
    _ensure_workspace_scope(checklist, workspace_id, "Required field checklist not found")
    config = _checklist_data(checklist)
    existing_names = {
        row[0]
        for row in _scope_query(
            db.query(RequiredFieldChecklist.name).filter(RequiredFieldChecklist.archived == False),  # noqa: E712
            RequiredFieldChecklist,
            workspace_id,
        ).all()
    }
    duplicated_name = _duplicate_name(checklist.name, existing_names)
    duplicated_config = {
        **config,
        "name": duplicated_name,
    }
    duplicated = RequiredFieldChecklist(
        workspace_id=workspace_id,
        name=duplicated_name,
        description=checklist.description,
        config_json=json.dumps(duplicated_config, ensure_ascii=False),
        archived=False,
    )
    db.add(duplicated)
    db.flush()
    log_audit_event(
        db,
        entity_type="required_field_checklist",
        entity_id=duplicated.id,
        action="duplicated",
        message=f"Duplicated required field checklist {checklist.name} to {duplicated.name}",
        metadata={"source_checklist_id": checklist.id, "item_count": len(duplicated_config["items"])},
    )
    db.commit()
    db.refresh(duplicated)
    return _checklist_read(duplicated)


@app.patch("/api/required-field-checklists/{checklist_id}", response_model=RequiredFieldChecklistRead)
def update_required_field_checklist(
    checklist_id: str,
    payload: RequiredFieldChecklistUpdate,
    request: Request,
    db: Session = Depends(get_db),
) -> RequiredFieldChecklistRead:
    workspace_id = _current_workspace_id(request)
    checklist = db.get(RequiredFieldChecklist, checklist_id)
    _ensure_workspace_scope(checklist, workspace_id, "Required field checklist not found")
    current = _checklist_data(checklist)
    next_config = {
        "name": payload.name if payload.name is not None else checklist.name,
        "description": payload.description if "description" in payload.model_fields_set else checklist.description,
        "regions": [region.model_dump() for region in payload.regions] if payload.regions is not None else current.get("regions", []),
        "items": [item.model_dump() for item in payload.items] if payload.items is not None else current["items"],
    }
    _validate_checklist_region_references(next_config)
    checklist.name = next_config["name"]
    checklist.description = next_config["description"]
    checklist.config_json = json.dumps(next_config, ensure_ascii=False)
    log_audit_event(
        db,
        entity_type="required_field_checklist",
        entity_id=checklist.id,
        action="updated",
        message=f"Updated required field checklist {checklist.name}",
        metadata={"item_count": len(next_config["items"])},
    )
    db.commit()
    db.refresh(checklist)
    return _checklist_read(checklist)


@app.delete("/api/required-field-checklists/{checklist_id}", response_model=RequiredFieldChecklistRead)
def delete_required_field_checklist(checklist_id: str, request: Request, db: Session = Depends(get_db)) -> RequiredFieldChecklistRead:
    workspace_id = _current_workspace_id(request)
    checklist = db.get(RequiredFieldChecklist, checklist_id)
    _ensure_workspace_scope(checklist, workspace_id, "Required field checklist not found")
    checklist.archived = True
    log_audit_event(
        db,
        entity_type="required_field_checklist",
        entity_id=checklist.id,
        action="archived",
        message=f"Archived required field checklist {checklist.name}",
        metadata={"name": checklist.name},
    )
    db.commit()
    db.refresh(checklist)
    return _checklist_read(checklist)


@app.post("/api/required-field-check-jobs", response_model=RequiredFieldCheckJobRead)
def create_required_field_check_job(
    payload: RequiredFieldCheckJobCreate,
    background_tasks: BackgroundTasks,
    request: Request,
    db: Session = Depends(get_db),
) -> RequiredFieldCheckJobRead:
    workspace_id = _current_workspace_id(request)
    document = db.get(Document, payload.document_id)
    checklist = db.get(RequiredFieldChecklist, payload.checklist_id)
    _ensure_workspace_scope(document, workspace_id, "Document not found")
    if not checklist or checklist.archived or (workspace_id is not None and checklist.workspace_id != workspace_id):
        raise HTTPException(status_code=404, detail="Required field checklist not found")
    status = _execution_job_status_for_document(document)
    if status == "blocked":
        _raise_document_not_executable(document)
    if document.status == "ready":
        _repair_image_document_if_needed(document, db)
    job = RequiredFieldCheckJob(workspace_id=workspace_id, document_id=document.id, checklist_id=checklist.id, status=status)
    db.add(job)
    db.flush()
    log_audit_event(
        db,
        entity_type="required_field_check_job",
        entity_id=job.id,
        action="created",
        message="Required field check job created",
        metadata={"document_id": document.id, "checklist_id": checklist.id},
    )
    db.commit()
    db.refresh(job)
    response = _required_field_job_read(job)
    if status == "queued":
        db.close()
        background_tasks.add_task(run_required_field_check_job, job.id)
    return response


@app.get("/api/required-field-check-jobs/{job_id}", response_model=RequiredFieldCheckJobRead)
def get_required_field_check_job(job_id: str, request: Request, db: Session = Depends(get_db)) -> RequiredFieldCheckJobRead:
    workspace_id = _current_workspace_id(request)
    job = db.get(RequiredFieldCheckJob, job_id)
    _ensure_workspace_scope(job, workspace_id, "Required field check job not found")
    return _required_field_job_read(job)


@app.patch("/api/required-field-check-results/{result_id}", response_model=RequiredFieldCheckResultRead)
def patch_required_field_check_result(
    result_id: str,
    payload: RequiredFieldCheckResultPatch,
    request: Request,
    db: Session = Depends(get_db),
) -> RequiredFieldCheckResultRead:
    workspace_id = _current_workspace_id(request)
    result = db.get(RequiredFieldCheckResult, result_id)
    if not result:
        raise HTTPException(status_code=404, detail="Required field check result not found")
    _ensure_workspace_scope(result.job, workspace_id, "Required field check result not found")
    if payload.corrected_output is not None:
        result.corrected_output = json.dumps(payload.corrected_output, ensure_ascii=False)
    if payload.reviewed is not None:
        result.reviewed = payload.reviewed
    db.commit()
    db.refresh(result)
    return RequiredFieldCheckResultRead(**required_field_result_to_dict(result))


@app.post("/api/extraction-jobs", response_model=ExtractionJobRead)
def create_extraction_job(
    payload: ExtractionJobCreate,
    background_tasks: BackgroundTasks,
    request: Request,
    db: Session = Depends(get_db),
) -> ExtractionJobRead:
    workspace_id = _current_workspace_id(request)
    document = db.get(Document, payload.document_id)
    schema = db.get(Schema, payload.schema_id)
    _ensure_workspace_scope(document, workspace_id, "Document not found")
    if not schema or (workspace_id is not None and schema.workspace_id != workspace_id):
        raise HTTPException(status_code=404, detail="Schema not found")
    status = _execution_job_status_for_document(document)
    if status == "blocked":
        _raise_document_not_executable(document)
    if document.status == "ready":
        _repair_image_document_if_needed(document, db)

    job = ExtractionJob(
        workspace_id=workspace_id,
        document_id=document.id,
        schema_id=schema.id,
        schema_version=1,
        status=status,
    )
    db.add(job)
    db.flush()
    log_audit_event(
        db,
        entity_type="extraction_job",
        entity_id=job.id,
        action="created",
        message="Extraction job created",
        metadata={"document_id": document.id, "schema_id": schema.id},
    )
    db.commit()
    db.refresh(job)
    response = _job_read(job)
    if status == "queued":
        db.close()
        background_tasks.add_task(run_extraction_job, job.id)
    return response


@app.post("/api/extraction-jobs/draft", response_model=ExtractionJobRead)
def create_draft_extraction_job(
    payload: DraftExtractionJobCreate,
    background_tasks: BackgroundTasks,
    request: Request,
    db: Session = Depends(get_db),
) -> ExtractionJobRead:
    workspace_id = _current_workspace_id(request)
    document = db.get(Document, payload.document_id)
    _ensure_workspace_scope(document, workspace_id, "Document not found")
    status = _execution_job_status_for_document(document)
    if status == "blocked":
        _raise_document_not_executable(document)
    if document.status == "ready":
        _repair_image_document_if_needed(document, db)

    draft_schema = payload.schema_definition
    schema_data = draft_schema.model_dump()
    schema_data["is_template"] = False
    schema_data["template_category"] = None
    schema_data["pinned"] = False
    _validate_schema_region_references(schema_data)

    schema = Schema(
        workspace_id=workspace_id,
        name=draft_schema.name,
        display_name=draft_schema.display_name or draft_schema.name,
        description=draft_schema.description,
        current_version=1,
        schema_json=json.dumps(schema_data, ensure_ascii=False),
        is_template=False,
        template_category=None,
        pinned=False,
        ephemeral=True,
        archived=False,
    )
    db.add(schema)
    db.flush()
    job = ExtractionJob(
        workspace_id=workspace_id,
        document_id=document.id,
        schema_id=schema.id,
        schema_version=1,
        status=status,
    )
    db.add(job)
    db.flush()
    log_audit_event(
        db,
        entity_type="extraction_job",
        entity_id=job.id,
        action="created",
        message="Draft extraction job created",
        metadata={"document_id": document.id, "schema_id": schema.id, "schema_mode": "draft"},
    )
    db.commit()
    db.refresh(job)
    response = _job_read(job)
    if status == "queued":
        db.close()
        background_tasks.add_task(run_extraction_job, job.id)
    return response


@app.get("/api/extraction-jobs", response_model=list[ExtractionJobRead])
def list_extraction_jobs(
    request: Request,
    document_id: str | None = None,
    limit: int = Query(default=20, ge=1, le=100),
    db: Session = Depends(get_db),
) -> list[ExtractionJobRead]:
    workspace_id = _current_workspace_id(request)
    query = _scope_query(db.query(ExtractionJob), ExtractionJob, workspace_id)
    if document_id:
        query = query.filter(ExtractionJob.document_id == document_id)
    jobs = query.order_by(ExtractionJob.created_at.desc()).limit(limit).all()
    return [_job_read(job) for job in jobs]


@app.get("/api/extraction-jobs/{job_id}", response_model=ExtractionJobRead)
def get_extraction_job(job_id: str, request: Request, db: Session = Depends(get_db)) -> ExtractionJobRead:
    workspace_id = _current_workspace_id(request)
    job = db.get(ExtractionJob, job_id)
    _ensure_workspace_scope(job, workspace_id, "Extraction job not found")
    return _job_read(job)


@app.patch("/api/extraction-results/{result_id}")
def patch_extraction_result(
    result_id: str,
    payload: ExtractionResultPatch,
    request: Request,
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    workspace_id = _current_workspace_id(request)
    result = db.get(ExtractionResult, result_id)
    if not result:
        raise HTTPException(status_code=404, detail="Extraction result not found")
    _ensure_workspace_scope(result.job, workspace_id, "Extraction result not found")
    if payload.corrected_output is not None:
        result.corrected_output = json.dumps(payload.corrected_output, ensure_ascii=False)
    if payload.reviewed_fields is not None:
        result.reviewed_fields = json.dumps(payload.reviewed_fields, ensure_ascii=False)
    log_audit_event(
        db,
        entity_type="extraction_result",
        entity_id=result.id,
        action="review_saved",
        message="Review changes saved",
        metadata={
            "has_corrections": payload.corrected_output is not None,
            "reviewed_count": len(payload.reviewed_fields or []),
        },
    )
    db.commit()
    db.refresh(result)
    return result_to_dict(result)


@app.get("/api/extraction-results/{result_id}/export")
def export_extraction_result(
    result_id: str,
    request: Request,
    format: str = Query(default="json", pattern="^(json|csv|xlsx)$"),
    preset_id: str | None = None,
    db: Session = Depends(get_db),
) -> Response:
    workspace_id = _current_workspace_id(request)
    result = db.get(ExtractionResult, result_id)
    if not result:
        raise HTTPException(status_code=404, detail="Extraction result not found")

    job = result.job or db.get(ExtractionJob, result.job_id)
    _ensure_workspace_scope(job, workspace_id, "Extraction result not found")
    schema = db.get(Schema, job.schema_id) if job else None
    payload = json.loads(result.corrected_output) if result.corrected_output else json.loads(result.validated_output)
    original_payload = json.loads(result.validated_output)
    reviewed_fields = set(json.loads(result.reviewed_fields or "[]"))
    preset = db.get(ExportPreset, preset_id) if preset_id else None
    if preset_id:
        _ensure_workspace_scope(preset, workspace_id, "Export preset not found")
    export_payload = _apply_export_preset(payload, preset) if preset else payload
    original_export_payload = _apply_export_preset(original_payload, preset) if preset else original_payload
    log_audit_event(
        db,
        entity_type="extraction_result",
        entity_id=result.id,
        action="exported",
        message=f"Exported {format.upper()}",
        metadata={"format": format, "preset_id": preset_id},
    )
    db.commit()
    filename = _export_filename("KIE", schema.name if schema else "schema", job.id if job else result_id, format)
    if format == "json":
        return JSONResponse(export_payload, headers=_download_headers(filename))

    fieldnames = [
        "key_name",
        "value",
        "normalized_value",
        "page",
        "confidence",
        "evidence",
        "warnings",
        "original_value",
        "changed",
        "reviewed",
        "ai_review_enabled",
        "ai_review_status",
        "ai_corrected",
        "ai_review_reason",
        "ai_review_confidence",
        "ai_initial_value",
        "ai_initial_evidence",
        "ai_correction_reason",
    ]
    rows: list[dict[str, Any]] = []
    original_values = original_export_payload.get("values", {}) if isinstance(original_export_payload.get("values"), dict) else {}
    for key, value in export_payload.get("values", {}).items():
        value_dict = value if isinstance(value, dict) else {}
        original_value = original_values.get(key)
        ai_review = value_dict.get("ai_review") if isinstance(value_dict.get("ai_review"), dict) else {}
        current_cell = _extract_kie_cell_value(value)
        original_cell = _extract_kie_cell_value(original_value) if original_value is not None else current_cell
        rows.append(
            {
                "key_name": key,
                "value": current_cell,
                "normalized_value": value_dict.get("normalized_value"),
                "page": value_dict.get("page"),
                "confidence": value_dict.get("confidence"),
                "evidence": value_dict.get("evidence"),
                "warnings": ";".join(value_dict.get("warnings", [])),
                "original_value": original_cell,
                "changed": current_cell != original_cell,
                "reviewed": key in reviewed_fields,
                "ai_review_enabled": bool(ai_review.get("enabled")),
                "ai_review_status": ai_review.get("judgement_status"),
                "ai_corrected": bool(ai_review.get("corrected")),
                "ai_review_reason": ai_review.get("judgement_reason"),
                "ai_review_confidence": ai_review.get("judgement_confidence"),
                "ai_initial_value": ai_review.get("initial_value"),
                "ai_initial_evidence": ai_review.get("initial_evidence"),
                "ai_correction_reason": ai_review.get("correction_reason"),
            }
        )
    if format == "xlsx":
        return _xlsx_download_response(rows, fieldnames, filename)

    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=fieldnames)
    writer.writeheader()
    for row in rows:
        writer.writerow({key: _csv_cell(row.get(key)) for key in fieldnames})
    return _csv_download_response(output.getvalue(), filename)


@app.post("/api/export-presets", response_model=ExportPresetRead)
def create_export_preset(payload: ExportPresetCreate, request: Request, db: Session = Depends(get_db)) -> ExportPresetRead:
    workspace_id = _current_workspace_id(request)
    if payload.schema_id:
        schema = db.get(Schema, payload.schema_id)
        _ensure_workspace_scope(schema, workspace_id, "Schema not found")
    preset = ExportPreset(
        workspace_id=workspace_id,
        schema_id=payload.schema_id,
        name=payload.name.strip(),
        fields_json=json.dumps([field.model_dump() for field in payload.fields], ensure_ascii=False),
    )
    db.add(preset)
    db.flush()
    log_audit_event(
        db,
        entity_type="export_preset",
        entity_id=preset.id,
        action="created",
        message=f"Created export preset {preset.name}",
        metadata={"schema_id": preset.schema_id},
    )
    db.commit()
    db.refresh(preset)
    return _export_preset_read(preset)


@app.get("/api/export-presets", response_model=list[ExportPresetRead])
def list_export_presets(request: Request, schema_id: str | None = None, db: Session = Depends(get_db)) -> list[ExportPresetRead]:
    workspace_id = _current_workspace_id(request)
    query = _scope_query(db.query(ExportPreset), ExportPreset, workspace_id)
    if schema_id:
        schema = db.get(Schema, schema_id)
        _ensure_workspace_scope(schema, workspace_id, "Schema not found")
        query = query.filter((ExportPreset.schema_id == schema_id) | (ExportPreset.schema_id.is_(None)))
    presets = query.order_by(ExportPreset.created_at.desc()).all()
    return [_export_preset_read(preset) for preset in presets]


@app.patch("/api/export-presets/{preset_id}", response_model=ExportPresetRead)
def update_export_preset(
    preset_id: str,
    payload: ExportPresetUpdate,
    request: Request,
    db: Session = Depends(get_db),
) -> ExportPresetRead:
    workspace_id = _current_workspace_id(request)
    preset = db.get(ExportPreset, preset_id)
    _ensure_workspace_scope(preset, workspace_id, "Export preset not found")
    if payload.schema_id:
        schema = db.get(Schema, payload.schema_id)
        _ensure_workspace_scope(schema, workspace_id, "Schema not found")
    if payload.name is not None:
        preset.name = payload.name.strip()
    if "schema_id" in payload.model_fields_set:
        preset.schema_id = payload.schema_id
    if payload.fields is not None:
        preset.fields_json = json.dumps([field.model_dump() for field in payload.fields], ensure_ascii=False)
    log_audit_event(
        db,
        entity_type="export_preset",
        entity_id=preset.id,
        action="updated",
        message=f"Updated export preset {preset.name}",
        metadata={"schema_id": preset.schema_id},
    )
    db.commit()
    db.refresh(preset)
    return _export_preset_read(preset)


@app.delete("/api/export-presets/{preset_id}")
def delete_export_preset(preset_id: str, request: Request, db: Session = Depends(get_db)) -> dict[str, str]:
    workspace_id = _current_workspace_id(request)
    preset = db.get(ExportPreset, preset_id)
    _ensure_workspace_scope(preset, workspace_id, "Export preset not found")
    log_audit_event(
        db,
        entity_type="export_preset",
        entity_id=preset.id,
        action="deleted",
        message=f"Deleted export preset {preset.name}",
        metadata={"schema_id": preset.schema_id},
    )
    db.delete(preset)
    db.commit()
    return {"status": "deleted"}


@app.post("/api/workflows", response_model=WorkflowDefinitionRead)
def create_workflow(payload: WorkflowDefinitionCreate, request: Request, db: Session = Depends(get_db)) -> WorkflowDefinitionRead:
    workspace_id = _current_workspace_id(request)
    try:
        validate_workflow_definition(payload.definition, db, workspace_id=workspace_id)
    except WorkflowDefinitionError as exc:
        raise HTTPException(status_code=422, detail={"errors": exc.errors}) from exc
    workflow = WorkflowDefinition(
        workspace_id=workspace_id,
        name=payload.name.strip(),
        description=payload.description,
        definition_json=json.dumps(payload.definition, ensure_ascii=False),
        archived=False,
    )
    db.add(workflow)
    db.flush()
    log_audit_event(
        db,
        entity_type="workflow_definition",
        entity_id=workflow.id,
        action="created",
        message=f"Created workflow {workflow.name}",
        metadata={"node_count": len(payload.definition.get("nodes", [])), "edge_count": len(payload.definition.get("edges", []))},
    )
    db.commit()
    db.refresh(workflow)
    return WorkflowDefinitionRead(**workflow_definition_to_read(workflow, db))


@app.get("/api/workflows", response_model=list[WorkflowDefinitionRead])
def list_workflows(request: Request, include_archived: bool = False, db: Session = Depends(get_db)) -> list[WorkflowDefinitionRead]:
    workspace_id = _current_workspace_id(request)
    query = db.query(WorkflowDefinition)
    query = _scope_query(query, WorkflowDefinition, workspace_id)
    if not include_archived:
        query = query.filter(WorkflowDefinition.archived == False)  # noqa: E712
    workflows = query.order_by(WorkflowDefinition.created_at.desc()).all()
    return [WorkflowDefinitionRead(**workflow_definition_to_read(workflow, db)) for workflow in workflows]


@app.post("/api/workflows/ai-draft", response_model=WorkflowAiDraftRead)
def create_workflow_ai_draft(
    request: Request,
    files: list[UploadFile] = File(...),
    include_checklist: bool = Form(True),
    db: Session = Depends(get_db),
) -> WorkflowAiDraftRead:
    workspace_id = _current_workspace_id(request)
    image_paths = _save_workflow_ai_draft_images(files)
    try:
        schema_recommendation = _schema_recommendation_read(recommend_schema_with_vlm([str(path) for path in image_paths]))
        schema_name = _available_scoped_name(
            db,
            Schema,
            schema_recommendation.name or "AI 추천 schema",
            workspace_id=workspace_id,
        )
        schema_draft = SchemaCreate(
            name=schema_name,
            display_name=schema_recommendation.display_name or schema_name,
            description=schema_recommendation.description,
            fields=schema_recommendation.fields,
            regions=[],
        )

        checklist_draft = None
        if include_checklist:
            checklist_recommendation = _required_field_checklist_recommendation_read(
                recommend_required_field_checklist_with_vlm([str(path) for path in image_paths])
            )
            checklist_draft = RequiredFieldChecklistCreate(
                name=_available_scoped_name(
                    db,
                    RequiredFieldChecklist,
                    checklist_recommendation.name or "AI 추천 checklist",
                    workspace_id=workspace_id,
                ),
                description=checklist_recommendation.description,
                regions=checklist_recommendation.regions,
                items=checklist_recommendation.items,
            )

        definition = _workflow_ai_draft_definition(include_checklist=checklist_draft is not None)
        workflow_name = _workflow_ai_draft_name(schema_draft, sample_count=len(image_paths))
        draft_event_id = f"workflow_ai_draft_{uuid4().hex}"
        log_audit_event(
            db,
            entity_type="workflow_ai_draft",
            entity_id=draft_event_id,
            action="created",
            message="AI workflow draft generated without persisting sample images",
            metadata={
                "sample_count": len(image_paths),
                "schema_field_count": len(schema_draft.fields),
                "checklist_item_count": len(checklist_draft.items) if checklist_draft else 0,
                "images_persisted": False,
            },
        )
        db.commit()
        return WorkflowAiDraftRead(
            workflow_name=workflow_name,
            schema_draft=schema_draft,
            checklist_draft=checklist_draft,
            definition=definition,
            sample_count=len(image_paths),
            images_persisted=False,
            reasoning=schema_recommendation.reasoning,
        )
    except (ValidationError, ValueError) as exc:
        raise HTTPException(status_code=502, detail=f"VLM returned an invalid workflow draft recommendation: {exc}") from exc
    except RuntimeError as exc:
        raise HTTPException(status_code=400, detail=vlm_error_detail(exc)) from exc
    finally:
        for path in image_paths:
            shutil.rmtree(path.parent, ignore_errors=True)


@app.get("/api/workflows/{workflow_id}", response_model=WorkflowDefinitionRead)
def get_workflow(workflow_id: str, request: Request, db: Session = Depends(get_db)) -> WorkflowDefinitionRead:
    workspace_id = _current_workspace_id(request)
    workflow = db.get(WorkflowDefinition, workflow_id)
    if not workflow or workflow.archived or (workspace_id is not None and workflow.workspace_id != workspace_id):
        raise HTTPException(status_code=404, detail="Workflow not found")
    return WorkflowDefinitionRead(**workflow_definition_to_read(workflow, db))


@app.patch("/api/workflows/{workflow_id}", response_model=WorkflowDefinitionRead)
def update_workflow(
    workflow_id: str,
    payload: WorkflowDefinitionUpdate,
    request: Request,
    db: Session = Depends(get_db),
) -> WorkflowDefinitionRead:
    workspace_id = _current_workspace_id(request)
    workflow = db.get(WorkflowDefinition, workflow_id)
    if not workflow or workflow.archived or (workspace_id is not None and workflow.workspace_id != workspace_id):
        raise HTTPException(status_code=404, detail="Workflow not found")
    if payload.definition is not None:
        try:
            validate_workflow_definition(payload.definition, db, workspace_id=workspace_id)
        except WorkflowDefinitionError as exc:
            raise HTTPException(status_code=422, detail={"errors": exc.errors}) from exc
        workflow.definition_json = json.dumps(payload.definition, ensure_ascii=False)
    if payload.name is not None:
        workflow.name = payload.name.strip()
    if "description" in payload.model_fields_set:
        workflow.description = payload.description
    log_audit_event(
        db,
        entity_type="workflow_definition",
        entity_id=workflow.id,
        action="updated",
        message=f"Updated workflow {workflow.name}",
        metadata={},
    )
    db.commit()
    db.refresh(workflow)
    return WorkflowDefinitionRead(**workflow_definition_to_read(workflow, db))


@app.delete("/api/workflows/{workflow_id}", response_model=WorkflowDefinitionRead)
def delete_workflow(workflow_id: str, request: Request, db: Session = Depends(get_db)) -> WorkflowDefinitionRead:
    workspace_id = _current_workspace_id(request)
    workflow = db.get(WorkflowDefinition, workflow_id)
    if not workflow or (workspace_id is not None and workflow.workspace_id != workspace_id):
        raise HTTPException(status_code=404, detail="Workflow not found")
    workflow.archived = True
    log_audit_event(
        db,
        entity_type="workflow_definition",
        entity_id=workflow.id,
        action="archived",
        message=f"Archived workflow {workflow.name}",
        metadata={},
    )
    db.commit()
    db.refresh(workflow)
    return WorkflowDefinitionRead(**workflow_definition_to_read(workflow, db))


@app.post("/api/workflows/{workflow_id}/runs", response_model=WorkflowRunRead)
async def create_workflow_run(
    workflow_id: str,
    request: Request,
    background_tasks: BackgroundTasks,
    db: Session = Depends(get_db),
) -> WorkflowRunRead:
    workspace_id = _current_workspace_id(request)
    form, files = await _read_batch_upload_form(request)
    workflow = db.get(WorkflowDefinition, workflow_id)
    if not workflow or workflow.archived or (workspace_id is not None and workflow.workspace_id != workspace_id):
        raise HTTPException(status_code=404, detail="Workflow not found")
    if not files:
        raise HTTPException(status_code=400, detail="At least one file is required")
    _validate_upload_file_count(files)
    try:
        validate_workflow_definition(json.loads(workflow.definition_json), db, workspace_id=workspace_id)
    except WorkflowDefinitionError as exc:
        raise HTTPException(status_code=422, detail={"errors": exc.errors}) from exc

    run = WorkflowRun(
        workspace_id=workspace_id,
        workflow_id=workflow.id,
        workflow_name=workflow.name,
        workflow_definition_json=workflow.definition_json,
        status="uploading",
        total_count=len(files),
        created_at=datetime.utcnow(),
    )
    db.add(run)
    db.flush()
    log_audit_event(
        db,
        entity_type="workflow_run",
        entity_id=run.id,
        action="created",
        message=f"Created workflow run with {len(files)} file(s)",
        metadata={"workflow_id": workflow.id, "file_count": len(files)},
    )
    db.commit()
    await _append_workflow_upload_items(run, form, files, db)
    db.refresh(run)
    _validate_owner_can_start(run, run.items)
    now = datetime.utcnow()
    run.execution_generation = (run.execution_generation or 0) + 1
    run.status = "running"
    run.upload_duration_ms = _workflow_upload_duration_ms(run)
    run.started_at = run.started_at or now
    run.inference_started_at = now
    for item in run.items:
        if item.status == "queued":
            item.execution_generation = run.execution_generation
    log_audit_event(
        db,
        entity_type="workflow_run",
        entity_id=run.id,
        action="started",
        message=f"Started workflow run with {run.total_count} file(s)",
        metadata={"workflow_id": run.workflow_id, "file_count": run.total_count},
    )
    db.commit()
    db.refresh(run)
    response = WorkflowRunRead(**workflow_run_to_read(run))
    execution_generation = run.execution_generation
    db.close()
    background_tasks.add_task(run_workflow_run, run.id, execution_generation)
    return response


@app.post("/api/workflows/{workflow_id}/runs/init", response_model=WorkflowRunRead)
def init_workflow_run(
    workflow_id: str,
    payload: WorkflowRunInitRequest,
    request: Request,
    db: Session = Depends(get_db),
) -> WorkflowRunRead:
    workspace_id = _current_workspace_id(request)
    workflow = db.get(WorkflowDefinition, workflow_id)
    if not workflow or workflow.archived or (workspace_id is not None and workflow.workspace_id != workspace_id):
        raise HTTPException(status_code=404, detail="Workflow not found")
    _validate_declared_batch_file_count(payload.total_count)
    try:
        validate_workflow_definition(json.loads(workflow.definition_json), db, workspace_id=workspace_id)
    except WorkflowDefinitionError as exc:
        raise HTTPException(status_code=422, detail={"errors": exc.errors}) from exc

    run = WorkflowRun(
        workspace_id=workspace_id,
        workflow_id=workflow.id,
        workflow_name=workflow.name,
        workflow_definition_json=workflow.definition_json,
        status="uploading",
        total_count=payload.total_count,
        created_at=datetime.utcnow(),
    )
    db.add(run)
    db.flush()
    log_audit_event(
        db,
        entity_type="workflow_run",
        entity_id=run.id,
        action="upload_initialized",
        message=f"Initialized workflow run upload with {payload.total_count} file(s)",
        metadata={"workflow_id": workflow.id, "file_count": payload.total_count},
    )
    db.commit()
    db.refresh(run)
    return WorkflowRunRead(**workflow_run_to_read(run))


@app.post("/api/workflows/{workflow_id}/runs/from-documents", response_model=WorkflowRunRead)
def create_workflow_run_from_documents(
    workflow_id: str,
    payload: WorkflowRunFromDocumentsRequest,
    background_tasks: BackgroundTasks,
    request: Request,
    db: Session = Depends(get_db),
) -> WorkflowRunRead:
    workspace_id = _current_workspace_id(request)
    workflow = db.get(WorkflowDefinition, workflow_id)
    if not workflow or workflow.archived or (workspace_id is not None and workflow.workspace_id != workspace_id):
        raise HTTPException(status_code=404, detail="Workflow not found")
    try:
        validate_workflow_definition(json.loads(workflow.definition_json), db, workspace_id=workspace_id)
    except WorkflowDefinitionError as exc:
        raise HTTPException(status_code=422, detail={"errors": exc.errors}) from exc
    documents = _documents_for_selection(payload.document_ids, db, workspace_id=workspace_id)
    run, has_ready_items = _create_workflow_run_from_documents(workflow, documents, db)
    response = WorkflowRunRead(**workflow_run_to_read(run))
    if has_ready_items:
        execution_generation = run.execution_generation
        run_id = run.id
        db.close()
        background_tasks.add_task(run_workflow_run, run_id, execution_generation)
    return response


@app.post("/api/workflow-runs/{run_id}/items", response_model=WorkflowRunRead)
async def append_workflow_run_items(
    run_id: str,
    request: Request,
    db: Session = Depends(get_db),
) -> WorkflowRunRead:
    workspace_id = _current_workspace_id(request)
    form, files = await _read_batch_upload_form(request)
    run_repo = SqlAlchemyWorkflowRunRepository(db)
    run = run_repo.get(run_id)
    _ensure_workspace_scope(run, workspace_id, "Workflow run not found")
    if run.status == "paused" and len(run.items) < run.total_count:
        run.status = "uploading"
        run.error_message = None
        db.flush()
    if run.status not in {"uploading", "queued"}:
        raise HTTPException(status_code=409, detail="Workflow run already started")
    await _append_workflow_upload_items(run, form, files, db)
    run_repo.refresh(run)
    return WorkflowRunRead(**workflow_run_to_read(run))


@app.post("/api/workflow-runs/{run_id}/start", response_model=WorkflowRunRead)
def start_workflow_run(
    run_id: str,
    background_tasks: BackgroundTasks,
    request: Request,
    db: Session = Depends(get_db),
) -> WorkflowRunRead:
    workspace_id = _current_workspace_id(request)
    run_repo = SqlAlchemyWorkflowRunRepository(db)
    run = run_repo.get(run_id)
    _ensure_workspace_scope(run, workspace_id, "Workflow run not found")
    if run.status not in {"uploading", "queued", "waiting"}:
        raise HTTPException(status_code=409, detail=f"Workflow run cannot be started from status {run.status}")
    if run.status == "waiting":
        _validate_waiting_workflow_run_can_start(run, db)
    _validate_owner_can_start(run, run.items)
    now = datetime.utcnow()
    start_result = WorkflowRunApplicationService().start(run, now, upload_duration_ms=_workflow_upload_duration_ms(run))
    log_audit_event(
        db,
        entity_type="workflow_run",
        entity_id=run.id,
        action="started",
        message=f"Started workflow run with {run.total_count} file(s)",
        metadata={"workflow_id": run.workflow_id, "file_count": run.total_count},
    )
    db.commit()
    run_repo.refresh(run)
    response = WorkflowRunRead(**workflow_run_to_read(run))
    background_tasks.add_task(run_workflow_run, run.id, start_result.execution_generation)
    return response


@app.post("/api/workflow-runs/{run_id}/enqueue", response_model=WorkflowRunRead)
def enqueue_workflow_run(
    run_id: str,
    request: Request,
    payload: WorkflowRunEnqueueRequest | None = Body(default=None),
    db: Session = Depends(get_db),
) -> WorkflowRunRead:
    workspace_id = _current_workspace_id(request)
    run_repo = SqlAlchemyWorkflowRunRepository(db)
    source_run = run_repo.get(run_id)
    _ensure_workspace_scope(source_run, workspace_id, "Workflow run not found")
    if not source_run.items:
        raise HTTPException(status_code=422, detail="No uploaded workflow items are available to enqueue")
    _validate_workflow_enqueue_source(source_run)
    workflow_id = payload.workflow_id if payload and payload.workflow_id else source_run.workflow_id
    workflow = db.get(WorkflowDefinition, workflow_id)
    if not workflow or workflow.archived or (workspace_id is not None and workflow.workspace_id != workspace_id):
        raise HTTPException(status_code=404, detail="Workflow not found")
    try:
        validate_workflow_definition(json.loads(workflow.definition_json), db, workspace_id=workspace_id)
    except WorkflowDefinitionError as exc:
        raise HTTPException(status_code=422, detail={"errors": exc.errors}) from exc

    _validate_owner_can_start(source_run, source_run.items)
    now = datetime.utcnow()
    new_run, queued_count = _create_waiting_workflow_run(source_run, workflow, now, db)
    if not queued_count:
        new_run.status = "completed_with_errors"
        new_run.completed_at = now
    log_audit_event(
        db,
        entity_type="workflow_run",
        entity_id=new_run.id,
        action="queued",
        message=f"Queued workflow run after {source_run.id}",
        metadata={
            "workflow_id": new_run.workflow_id,
            "source_run_id": source_run.id,
            "queue_group_id": new_run.workflow_run_group_id,
            "queue_order": new_run.queue_order,
            "queued_count": queued_count,
        },
    )
    db.commit()
    db.refresh(new_run)
    return WorkflowRunRead(**workflow_run_to_read(new_run))


@app.post("/api/workflow-runs/{run_id}/cancel-waiting", response_model=WorkflowRunRead)
def cancel_waiting_workflow_run(run_id: str, request: Request, db: Session = Depends(get_db)) -> WorkflowRunRead:
    workspace_id = _current_workspace_id(request)
    run_repo = SqlAlchemyWorkflowRunRepository(db)
    run = run_repo.get(run_id)
    _ensure_workspace_scope(run, workspace_id, "Workflow run not found")
    if run.status != "waiting":
        raise HTTPException(status_code=409, detail="Workflow run is not waiting in the queue")
    _cancel_waiting_workflow_run(run, db)
    db.commit()
    run_repo.refresh(run)
    return WorkflowRunRead(**workflow_run_to_read(run))


@app.delete("/api/workflow-runs/{run_id}/queue-entry")
def delete_workflow_queue_entry(run_id: str, request: Request, db: Session = Depends(get_db)) -> dict[str, str]:
    workspace_id = _current_workspace_id(request)
    run = SqlAlchemyWorkflowRunRepository(db).get(run_id)
    _ensure_workspace_scope(run, workspace_id, "Workflow run not found")
    if not run.queued_from_run_id:
        raise HTTPException(status_code=409, detail="Only queued workflow run entries can be removed")
    if run.status in {"completed", "completed_with_errors", "needs_review", "failed"}:
        raise HTTPException(status_code=409, detail="Finished queued workflow runs cannot be removed")
    if run.status == "waiting":
        _cancel_waiting_workflow_run(run, db)
        db.flush()
    elif run.status != "canceled":
        _stop_workflow_run_without_deleting_documents(run, "Stopped and removed from workflow run queue", db)
        db.flush()
    item_count = len(run.items)
    queue_group_id = run.workflow_run_group_id
    queue_order = run.queue_order
    source_run_id = run.queued_from_run_id
    db.delete(run)
    log_audit_event(
        db,
        entity_type="workflow_run",
        entity_id=run_id,
        action="queue_entry_deleted",
        message=f"Removed queued workflow run entry with {item_count} item(s); shared documents were kept",
        metadata={
            "queue_group_id": queue_group_id,
            "queue_order": queue_order,
            "source_run_id": source_run_id,
            "removed_item_count": item_count,
        },
    )
    db.commit()
    return {"status": "deleted", "id": run_id}


@app.post("/api/workflow-runs/{run_id}/discard", response_model=WorkflowRunRead)
def discard_workflow_run(run_id: str, request: Request, db: Session = Depends(get_db)) -> WorkflowRunRead:
    workspace_id = _current_workspace_id(request)
    run_repo = SqlAlchemyWorkflowRunRepository(db)
    run = run_repo.get(run_id)
    _ensure_workspace_scope(run, workspace_id, "Workflow run not found")
    if run.status in WORKFLOW_RUN_TERMINAL_STATUSES:
        raise HTTPException(status_code=409, detail="Workflow run is already terminal")
    if run.status == "waiting":
        _cancel_waiting_workflow_run(run, db)
        db.commit()
        run_repo.refresh(run)
        return WorkflowRunRead(**workflow_run_to_read(run))
    _stop_workflow_run_without_deleting_documents(run, "Stopped by user; library documents were kept", db)
    log_audit_event(
        db,
        entity_type="workflow_run",
        entity_id=run.id,
        action="discarded",
        message="Stopped workflow run without deleting document library payloads",
        metadata={"workflow_id": run.workflow_id, "kept_document_count": len(run.items)},
    )
    db.commit()
    run_repo.refresh(run)
    return WorkflowRunRead(**workflow_run_to_read(run))


@app.post("/api/workflow-runs/{run_id}/resume", response_model=WorkflowRunRead)
def resume_workflow_run(
    run_id: str,
    background_tasks: BackgroundTasks,
    request: Request,
    db: Session = Depends(get_db),
) -> WorkflowRunRead:
    workspace_id = _current_workspace_id(request)
    run_repo = SqlAlchemyWorkflowRunRepository(db)
    run = run_repo.get(run_id)
    _ensure_workspace_scope(run, workspace_id, "Workflow run not found")
    if run.status in {"completed", "completed_with_errors", "needs_review", "failed", "canceled"}:
        raise HTTPException(status_code=409, detail="Workflow run is already terminal")
    if run.status == "waiting":
        raise HTTPException(status_code=409, detail="Waiting workflow run must be started or canceled from the queue")
    if not run.items:
        raise HTTPException(status_code=422, detail="No uploaded workflow items are available to continue")
    _validate_owner_upload_complete(run, run.items)
    now = datetime.utcnow()
    resume_result = WorkflowRunApplicationService().resume(run, now, fallback_status=lambda: workflow_run_to_read(run)["status"])
    log_audit_event(
        db,
        entity_type="workflow_run",
        entity_id=run.id,
        action="resumed",
        message=f"Continued workflow run with {resume_result.queued_count} queued item(s)",
        metadata={"workflow_id": run.workflow_id, "queued_count": resume_result.queued_count},
    )
    db.commit()
    run_repo.refresh(run)
    response = WorkflowRunRead(**workflow_run_to_read(run))
    if resume_result.queued_count:
        background_tasks.add_task(run_workflow_run, run.id, resume_result.execution_generation)
    return response


@app.post("/api/workflow-runs/{run_id}/pause", response_model=WorkflowRunRead)
def pause_workflow_run(run_id: str, request: Request, db: Session = Depends(get_db)) -> WorkflowRunRead:
    workspace_id = _current_workspace_id(request)
    run_repo = SqlAlchemyWorkflowRunRepository(db)
    run = run_repo.get(run_id)
    _ensure_workspace_scope(run, workspace_id, "Workflow run not found")
    if run.status in {"completed", "completed_with_errors", "needs_review", "failed", "canceled"}:
        raise HTTPException(status_code=409, detail="Workflow run is already terminal")
    if run.status == "waiting":
        raise HTTPException(status_code=409, detail="Waiting workflow run must be started or canceled from the queue")
    if run.status == "paused":
        return WorkflowRunRead(**workflow_run_to_read(run))

    now = datetime.utcnow()
    pause_result = WorkflowRunApplicationService().pause(
        run,
        now,
        cancel_active_job=lambda item: _cancel_workflow_item_active_jobs(item, db, now, "Paused by user"),
    )
    _accumulate_workflow_run_inference_duration(run, now)
    log_audit_event(
        db,
        entity_type="workflow_run",
        entity_id=run.id,
        action="paused",
        message=f"Paused workflow run; {pause_result.paused_count} active item(s) held",
        metadata={"workflow_id": run.workflow_id, "paused_count": pause_result.paused_count},
    )
    db.commit()
    run_repo.refresh(run)
    return WorkflowRunRead(**workflow_run_to_read(run))


@app.post("/api/workflow-runs/{run_id}/restart", response_model=WorkflowRunRead)
def restart_workflow_run(
    run_id: str,
    background_tasks: BackgroundTasks,
    request: Request,
    payload: WorkflowRunRestartRequest | None = Body(default=None),
    db: Session = Depends(get_db),
) -> WorkflowRunRead:
    workspace_id = _current_workspace_id(request)
    run_repo = SqlAlchemyWorkflowRunRepository(db)
    run = run_repo.get(run_id)
    _ensure_workspace_scope(run, workspace_id, "Workflow run not found")
    if not run.items:
        raise HTTPException(status_code=422, detail="No uploaded workflow items are available to restart")
    workflow_id = payload.workflow_id if payload and payload.workflow_id else run.workflow_id
    workflow = db.get(WorkflowDefinition, workflow_id)
    if not workflow or workflow.archived or (workspace_id is not None and workflow.workspace_id != workspace_id):
        raise HTTPException(status_code=404, detail="Workflow not found")
    try:
        validate_workflow_definition(json.loads(workflow.definition_json), db, workspace_id=workspace_id)
    except WorkflowDefinitionError as exc:
        raise HTTPException(status_code=422, detail={"errors": exc.errors}) from exc

    sealed_missing_count = _seal_missing_workflow_upload_items(run, db)
    if sealed_missing_count:
        db.flush()
        db.expire(run, ["items"])
    _validate_owner_upload_complete(run, run.items)
    now = datetime.utcnow()
    if run.status not in {"completed", "completed_with_errors", "needs_review", "failed", "canceled"}:
        _accumulate_workflow_run_inference_duration(run, now)
        run.execution_generation = (run.execution_generation or 0) + 1
        run.status = "canceled"
        run.error_message = "Replaced by restarted workflow run"
        run.completed_at = now
    new_run, queued_count = _create_restarted_workflow_run(run, workflow, now, db)
    if not queued_count:
        new_run.status = "completed_with_errors"
        new_run.completed_at = now
        new_run.inference_started_at = None
    log_audit_event(
        db,
        entity_type="workflow_run",
        entity_id=new_run.id,
        action="restarted",
        message=f"Created restarted workflow run with {queued_count} queued item(s)",
        metadata={
            "workflow_id": new_run.workflow_id,
            "source_run_id": run.id,
            "queued_count": queued_count,
            "sealed_missing_count": sealed_missing_count,
        },
    )
    db.commit()
    db.refresh(new_run)
    response = WorkflowRunRead(**workflow_run_to_read(new_run))
    if queued_count:
        background_tasks.add_task(run_workflow_run, new_run.id, new_run.execution_generation)
    return response


@app.post("/api/workflow-runs/{run_id}/retry-failed", response_model=WorkflowRunRead)
def retry_failed_workflow_run(
    run_id: str,
    background_tasks: BackgroundTasks,
    request: Request,
    db: Session = Depends(get_db),
) -> WorkflowRunRead:
    workspace_id = _current_workspace_id(request)
    run_repo = SqlAlchemyWorkflowRunRepository(db)
    run = run_repo.get(run_id)
    _ensure_workspace_scope(run, workspace_id, "Workflow run not found")
    if run.status == "canceled":
        raise HTTPException(status_code=409, detail="Canceled workflow run cannot retry failed items")
    blocking_statuses = {"uploading", "preprocessing", "queued", "running", "paused"}
    blocking_count = sum(1 for item in run.items if item.status in blocking_statuses)
    if blocking_count:
        raise HTTPException(
            status_code=409,
            detail={
                "message": "Failed items can be retried after active or paused items are finished.",
                "blocking_count": blocking_count,
            },
        )
    failed_count = sum(1 for item in run.items if item.status == "failed")
    if not failed_count:
        raise HTTPException(status_code=422, detail="No failed workflow items are available to retry")

    now = datetime.utcnow()
    _accumulate_workflow_run_inference_duration(run, now)
    run.execution_generation = (run.execution_generation or 0) + 1
    retry_result = WorkflowRunApplicationService().retry_failed(run, now, initial_payload=_initial_workflow_item_payload)
    if not retry_result.queued_count:
        raise HTTPException(
            status_code=422,
            detail={
                "message": "No failed workflow items can be retried because their documents are not ready.",
                "failed_count": failed_count,
            },
        )
    log_audit_event(
        db,
        entity_type="workflow_run",
        entity_id=run.id,
        action="retry_failed",
        message=f"Retried {retry_result.queued_count} failed workflow item(s)",
        metadata={"workflow_id": run.workflow_id, "queued_count": retry_result.queued_count, "failed_count": failed_count},
    )
    db.commit()
    run_repo.refresh(run)
    response = WorkflowRunRead(**workflow_run_to_read(run))
    background_tasks.add_task(run_workflow_run, run.id, run.execution_generation)
    return response


@app.get("/api/workflow-runs", response_model=list[WorkflowRunRead])
def list_workflow_runs(request: Request, limit: int = Query(default=20, ge=1, le=100), db: Session = Depends(get_db)) -> list[WorkflowRunRead]:
    workspace_id = _current_workspace_id(request)
    runs = SqlAlchemyWorkflowRunRepository(db).list_recent(limit=limit, workspace_id=workspace_id)
    return [WorkflowRunRead(**payload) for payload in workflow_runs_to_read(runs, include_items=False, db=db)]


@app.get("/api/workflow-runs/{run_id}", response_model=WorkflowRunRead)
def get_workflow_run(run_id: str, request: Request, db: Session = Depends(get_db)) -> WorkflowRunRead:
    workspace_id = _current_workspace_id(request)
    run = SqlAlchemyWorkflowRunRepository(db).get(run_id)
    _ensure_workspace_scope(run, workspace_id, "Workflow run not found")
    return WorkflowRunRead(**workflow_run_to_read(run))


@app.get("/api/workflow-runs/{run_id}/summary", response_model=WorkflowRunRead)
def get_workflow_run_summary(run_id: str, request: Request, db: Session = Depends(get_db)) -> WorkflowRunRead:
    workspace_id = _current_workspace_id(request)
    run = SqlAlchemyWorkflowRunRepository(db).get(run_id)
    _ensure_workspace_scope(run, workspace_id, "Workflow run not found")
    return WorkflowRunRead(**workflow_run_to_read(run, include_items=False, db=db))


@app.get("/api/workflow-runs/{run_id}/export")
def export_workflow_run(
    run_id: str,
    request: Request,
    format: str = Query(default="csv", pattern="^(json|csv|xlsx)$"),
    db: Session = Depends(get_db),
) -> Response:
    workspace_id = _current_workspace_id(request)
    run = db.get(WorkflowRun, run_id)
    _ensure_workspace_scope(run, workspace_id, "Workflow run not found")
    artifact = _build_export_artifact(db, "workflow_run", run.id, format)
    log_audit_event(
        db,
        entity_type="workflow_run",
        entity_id=run.id,
        action="exported",
        message=f"Exported workflow run {format.upper()}",
        metadata={"format": format},
    )
    db.commit()
    return _export_artifact_response(artifact)


@app.post("/api/export-jobs", response_model=ExportJobRead, status_code=202)
def create_export_job(
    payload: ExportJobCreate,
    background_tasks: BackgroundTasks,
    request: Request,
    db: Session = Depends(get_db),
) -> ExportJobRead:
    workspace_id = _current_workspace_id(request)
    job = _queue_export_job(
        db,
        background_tasks,
        owner_type=payload.owner_type,
        owner_id=payload.owner_id,
        format=payload.format,
        workspace_id=workspace_id,
    )
    return _export_job_read(job)


@app.get("/api/export-jobs", response_model=list[ExportJobRead])
def list_export_jobs(
    request: Request,
    owner_type: str | None = Query(default=None, pattern="^(workflow_run|batch|classification_batch|required_field_check_batch)$"),
    owner_id: str | None = Query(default=None),
    status: str | None = Query(default=None, pattern="^(queued|running|completed|failed)$"),
    limit: int = Query(default=20, ge=1, le=100),
    db: Session = Depends(get_db),
) -> list[ExportJobRead]:
    workspace_id = _current_workspace_id(request)
    query = _scope_query(db.query(ExportJob), ExportJob, workspace_id)
    if owner_type:
        query = query.filter(ExportJob.owner_type == owner_type)
    if owner_id:
        query = query.filter(ExportJob.owner_id == owner_id)
    if status:
        query = query.filter(ExportJob.status == status)
    jobs = query.order_by(ExportJob.created_at.desc(), ExportJob.id.desc()).limit(limit).all()
    return [_export_job_read(job) for job in jobs]


@app.post("/api/export-jobs/{job_id}/retry", response_model=ExportJobRead, status_code=202)
def retry_export_job(
    job_id: str,
    background_tasks: BackgroundTasks,
    request: Request,
    db: Session = Depends(get_db),
) -> ExportJobRead:
    workspace_id = _current_workspace_id(request)
    source = db.get(ExportJob, job_id)
    _ensure_workspace_scope(source, workspace_id, "Export job not found")
    if source.status != "failed":
        raise HTTPException(status_code=409, detail="Only failed export jobs can be retried")
    job = _queue_export_job(
        db,
        background_tasks,
        owner_type=source.owner_type,
        owner_id=source.owner_id,
        format=source.format,
        retry_source_job_id=source.id,
        workspace_id=workspace_id,
    )
    return _export_job_read(job)


@app.get("/api/export-jobs/{job_id}", response_model=ExportJobRead)
def get_export_job(job_id: str, request: Request, db: Session = Depends(get_db)) -> ExportJobRead:
    workspace_id = _current_workspace_id(request)
    job = db.get(ExportJob, job_id)
    _ensure_workspace_scope(job, workspace_id, "Export job not found")
    return _export_job_read(job)


@app.get("/api/export-jobs/{job_id}/download")
def download_export_job(job_id: str, request: Request, db: Session = Depends(get_db)) -> FileResponse:
    workspace_id = _current_workspace_id(request)
    job = db.get(ExportJob, job_id)
    _ensure_workspace_scope(job, workspace_id, "Export job not found")
    if job.status != "completed" or not job.storage_path or not job.filename:
        raise HTTPException(status_code=409, detail="Export job is not ready for download")
    path = materialize_storage_ref(job.storage_path, suffix=Path(job.filename).suffix)
    if not path.exists():
        raise HTTPException(status_code=404, detail="Export artifact not found")
    return FileResponse(
        path,
        media_type=job.content_type or "application/octet-stream",
        headers=_download_headers(job.filename),
    )


@app.post("/api/batches", response_model=BatchRead)
async def create_batch(
    request: Request,
    background_tasks: BackgroundTasks,
    db: Session = Depends(get_db),
) -> BatchRead:
    workspace_id = _current_workspace_id(request)
    form, files = await _read_batch_upload_form(request)
    schema_id = _required_form_value(form, "schema_id")
    schema = db.get(Schema, schema_id)
    if not schema or (workspace_id is not None and schema.workspace_id != workspace_id):
        raise HTTPException(status_code=404, detail="Schema not found")
    if not files:
        raise HTTPException(status_code=400, detail="At least one file is required")
    _validate_upload_file_count(files)

    batch = Batch(workspace_id=workspace_id, schema_id=schema.id, schema_version=1, status="uploading", total_count=len(files))
    db.add(batch)
    db.flush()
    log_audit_event(
        db,
        entity_type="batch",
        entity_id=batch.id,
        action="created",
        message=f"Created batch with {len(files)} file(s)",
        metadata={"schema_id": schema.id, "file_count": len(files)},
    )
    db.commit()
    await _append_extraction_batch_items(batch, form, files, db)
    db.refresh(batch)
    _validate_owner_can_start(batch, batch.items)
    batch.status = "running"
    job_ids = _queued_extraction_job_ids(batch)
    log_audit_event(
        db,
        entity_type="batch",
        entity_id=batch.id,
        action="started",
        message=f"Started batch with {len(job_ids)} queued job(s)",
        metadata={"schema_id": schema.id, "queued_count": len(job_ids)},
    )
    db.commit()
    db.refresh(batch)
    response = _batch_read(batch)
    db.close()
    background_tasks.add_task(run_batch_jobs, batch.id, job_ids)
    return response


@app.post("/api/batches/init", response_model=BatchRead)
def init_batch(payload: BatchInitRequest, request: Request, db: Session = Depends(get_db)) -> BatchRead:
    workspace_id = _current_workspace_id(request)
    schema = db.get(Schema, payload.schema_id)
    if not schema or (workspace_id is not None and schema.workspace_id != workspace_id):
        raise HTTPException(status_code=404, detail="Schema not found")
    _validate_declared_batch_file_count(payload.total_count)
    batch = Batch(workspace_id=workspace_id, schema_id=schema.id, schema_version=1, status="uploading", total_count=payload.total_count)
    db.add(batch)
    db.flush()
    log_audit_event(
        db,
        entity_type="batch",
        entity_id=batch.id,
        action="upload_initialized",
        message=f"Initialized batch upload with {payload.total_count} file(s)",
        metadata={"schema_id": schema.id, "file_count": payload.total_count},
    )
    db.commit()
    db.refresh(batch)
    return _batch_read(batch)


@app.post("/api/batches/from-documents", response_model=BatchRead)
def create_batch_from_documents(
    payload: BatchFromDocumentsRequest,
    background_tasks: BackgroundTasks,
    request: Request,
    db: Session = Depends(get_db),
) -> BatchRead:
    workspace_id = _current_workspace_id(request)
    schema = db.get(Schema, payload.schema_id)
    if not schema or (workspace_id is not None and schema.workspace_id != workspace_id):
        raise HTTPException(status_code=404, detail="Schema not found")
    documents = _documents_for_selection(payload.document_ids, db, workspace_id=workspace_id)
    batch, ready_job_ids = _create_extraction_batch_from_documents(schema, documents, db)
    response = _batch_read(batch)
    if ready_job_ids:
        db.close()
        background_tasks.add_task(run_batch_jobs, response.id, ready_job_ids)
    return response


@app.post("/api/batches/{batch_id}/items", response_model=BatchRead)
async def append_batch_items(batch_id: str, request: Request, db: Session = Depends(get_db)) -> BatchRead:
    workspace_id = _current_workspace_id(request)
    form, files = await _read_batch_upload_form(request)
    batch = db.get(Batch, batch_id)
    _ensure_workspace_scope(batch, workspace_id, "Batch not found")
    if batch.status not in {"uploading", "queued"}:
        raise HTTPException(status_code=409, detail="Batch already started")
    await _append_extraction_batch_items(batch, form, files, db)
    db.refresh(batch)
    return _batch_read(batch)


@app.post("/api/batches/{batch_id}/start", response_model=BatchRead)
def start_batch(batch_id: str, background_tasks: BackgroundTasks, request: Request, db: Session = Depends(get_db)) -> BatchRead:
    workspace_id = _current_workspace_id(request)
    batch = db.get(Batch, batch_id)
    _ensure_workspace_scope(batch, workspace_id, "Batch not found")
    if batch.status not in {"uploading", "queued"}:
        return _batch_read(batch)
    _validate_owner_can_start(batch, batch.items)
    batch.status = "running"
    job_ids = _queued_extraction_job_ids(batch)
    log_audit_event(
        db,
        entity_type="batch",
        entity_id=batch.id,
        action="started",
        message=f"Started batch with {len(job_ids)} queued job(s)",
        metadata={"schema_id": batch.schema_id, "queued_count": len(job_ids)},
    )
    db.commit()
    db.refresh(batch)
    response = _batch_read(batch)
    db.close()
    background_tasks.add_task(run_batch_jobs, batch.id, job_ids)
    return response


@app.post("/api/batches/{batch_id}/discard", response_model=BatchRead)
def discard_batch(batch_id: str, request: Request, db: Session = Depends(get_db)) -> BatchRead:
    workspace_id = _current_workspace_id(request)
    batch = db.get(Batch, batch_id)
    _ensure_workspace_scope(batch, workspace_id, "Batch not found")
    discarded_count = _discard_batch_items(batch, db)
    batch.status = "canceled"
    batch.completed_at = datetime.utcnow()
    log_audit_event(
        db,
        entity_type="batch",
        entity_id=batch.id,
        action="discarded",
        message=f"Discarded batch payloads for {discarded_count} item(s)",
        metadata={"schema_id": batch.schema_id, "discarded_count": discarded_count},
    )
    db.commit()
    db.refresh(batch)
    return _batch_read(batch)


@app.post("/api/batches/{batch_id}/resume", response_model=BatchRead)
def resume_batch(batch_id: str, background_tasks: BackgroundTasks, request: Request, db: Session = Depends(get_db)) -> BatchRead:
    workspace_id = _current_workspace_id(request)
    batch = db.get(Batch, batch_id)
    _ensure_workspace_scope(batch, workspace_id, "Batch not found")
    if batch.status in {"completed", "completed_with_errors", "failed", "canceled"}:
        raise HTTPException(status_code=409, detail="Batch is already terminal")
    if not batch.items:
        raise HTTPException(status_code=422, detail="No uploaded batch items are available to continue")
    _validate_owner_upload_complete(batch, batch.items)
    _prepare_job_batch_resume(batch.items)
    job_ids = _queued_extraction_job_ids(batch)
    batch.status = "running" if job_ids else _batch_read(batch).status
    batch.completed_at = None if job_ids else batch.completed_at
    log_audit_event(
        db,
        entity_type="batch",
        entity_id=batch.id,
        action="resumed",
        message=f"Continued batch with {len(job_ids)} queued job(s)",
        metadata={"schema_id": batch.schema_id, "queued_count": len(job_ids)},
    )
    db.commit()
    db.refresh(batch)
    response = _batch_read(batch)
    if job_ids:
        db.close()
        background_tasks.add_task(run_batch_jobs, batch.id, job_ids)
    return response


@app.get("/api/batches", response_model=list[BatchRead])
def list_batches(
    request: Request,
    limit: int = Query(default=20, ge=1, le=100),
    include_items: bool = Query(default=True),
    db: Session = Depends(get_db),
) -> list[BatchRead]:
    workspace_id = _current_workspace_id(request)
    batches = _scope_query(db.query(Batch), Batch, workspace_id).order_by(Batch.created_at.desc()).limit(limit).all()
    return [_batch_read(batch, include_items=include_items, db=db) for batch in batches]


@app.get("/api/batches/{batch_id}", response_model=BatchRead)
def get_batch(batch_id: str, request: Request, db: Session = Depends(get_db)) -> BatchRead:
    workspace_id = _current_workspace_id(request)
    batch = db.get(Batch, batch_id)
    _ensure_workspace_scope(batch, workspace_id, "Batch not found")
    return _batch_read(batch)


@app.get("/api/batches/{batch_id}/summary", response_model=BatchRead)
def get_batch_summary(batch_id: str, request: Request, db: Session = Depends(get_db)) -> BatchRead:
    workspace_id = _current_workspace_id(request)
    batch = db.get(Batch, batch_id)
    _ensure_workspace_scope(batch, workspace_id, "Batch not found")
    return _batch_read(batch, include_items=False, db=db)


@app.post("/api/batches/{batch_id}/cancel", response_model=BatchRead)
def cancel_batch(batch_id: str, request: Request, db: Session = Depends(get_db)) -> BatchRead:
    workspace_id = _current_workspace_id(request)
    batch = db.get(Batch, batch_id)
    _ensure_workspace_scope(batch, workspace_id, "Batch not found")

    canceled_count = 0
    now = datetime.utcnow()
    for item in batch.items:
        if item.job and item.job.status in {"queued", "running"}:
            item.job.status = "canceled"
            item.job.error_message = "Canceled by user"
            item.job.completed_at = now
            canceled_count += 1

    if canceled_count:
        _close_batch_if_all_jobs_terminal(batch, now)
        if batch.status not in {"canceled", "completed", "completed_with_errors"}:
            batch.status = "cancel_requested"
        log_audit_event(
            db,
            entity_type="batch",
            entity_id=batch.id,
            action="cancel_requested",
            message=f"Cancel requested for {canceled_count} running or queued job(s)",
            metadata={"canceled_count": canceled_count},
        )
    else:
        log_audit_event(
            db,
            entity_type="batch",
            entity_id=batch.id,
            action="cancel_skipped",
            message="No running or queued batch jobs to cancel",
            metadata={},
        )

    db.commit()
    db.refresh(batch)
    return _batch_read(batch)


@app.get("/api/batches/{batch_id}/export")
def export_batch(
    batch_id: str,
    request: Request,
    format: str = Query(default="csv", pattern="^(json|csv|xlsx)$"),
    db: Session = Depends(get_db),
) -> Response:
    workspace_id = _current_workspace_id(request)
    batch = db.get(Batch, batch_id)
    _ensure_workspace_scope(batch, workspace_id, "Batch not found")
    artifact = _build_export_artifact(db, "batch", batch.id, format)
    log_audit_event(
        db,
        entity_type="batch",
        entity_id=batch.id,
        action="exported",
        message=f"Exported batch {format.upper()}",
        metadata={"format": format},
    )
    db.commit()
    return _export_artifact_response(artifact)


@app.post("/api/classification-batches", response_model=ClassificationBatchRead)
async def create_classification_batch(
    request: Request,
    background_tasks: BackgroundTasks,
    db: Session = Depends(get_db),
) -> ClassificationBatchRead:
    workspace_id = _current_workspace_id(request)
    form, files = await _read_batch_upload_form(request)
    classifier_id = _required_form_value(form, "classifier_id")
    classifier = db.get(DocumentClassifier, classifier_id)
    if not classifier or classifier.archived or (workspace_id is not None and classifier.workspace_id != workspace_id):
        raise HTTPException(status_code=404, detail="Document classifier not found")
    if not files:
        raise HTTPException(status_code=400, detail="At least one file is required")
    _validate_upload_file_count(files)

    batch = ClassificationBatch(workspace_id=workspace_id, classifier_id=classifier.id, status="uploading", total_count=len(files))
    db.add(batch)
    db.flush()
    log_audit_event(
        db,
        entity_type="classification_batch",
        entity_id=batch.id,
        action="created",
        message=f"Created classification batch with {len(files)} file(s)",
        metadata={"classifier_id": classifier.id, "file_count": len(files)},
    )
    db.commit()
    await _append_classification_batch_items(batch, form, files, db)
    db.refresh(batch)
    _validate_owner_can_start(batch, batch.items)
    batch.status = "running"
    job_ids = _queued_classification_job_ids(batch)
    log_audit_event(
        db,
        entity_type="classification_batch",
        entity_id=batch.id,
        action="started",
        message=f"Started classification batch with {len(job_ids)} queued job(s)",
        metadata={"classifier_id": classifier.id, "queued_count": len(job_ids)},
    )
    db.commit()
    db.refresh(batch)
    response = _classification_batch_read(batch)
    db.close()
    background_tasks.add_task(run_classification_batch, batch.id, job_ids)
    return response


@app.post("/api/classification-batches/init", response_model=ClassificationBatchRead)
def init_classification_batch(
    payload: ClassificationBatchInitRequest,
    request: Request,
    db: Session = Depends(get_db),
) -> ClassificationBatchRead:
    workspace_id = _current_workspace_id(request)
    classifier = db.get(DocumentClassifier, payload.classifier_id)
    if not classifier or classifier.archived or (workspace_id is not None and classifier.workspace_id != workspace_id):
        raise HTTPException(status_code=404, detail="Document classifier not found")
    _validate_declared_batch_file_count(payload.total_count)
    batch = ClassificationBatch(workspace_id=workspace_id, classifier_id=classifier.id, status="uploading", total_count=payload.total_count)
    db.add(batch)
    db.flush()
    log_audit_event(
        db,
        entity_type="classification_batch",
        entity_id=batch.id,
        action="upload_initialized",
        message=f"Initialized classification batch upload with {payload.total_count} file(s)",
        metadata={"classifier_id": classifier.id, "file_count": payload.total_count},
    )
    db.commit()
    db.refresh(batch)
    return _classification_batch_read(batch)


@app.post("/api/classification-batches/from-documents", response_model=ClassificationBatchRead)
def create_classification_batch_from_documents(
    payload: ClassificationBatchFromDocumentsRequest,
    background_tasks: BackgroundTasks,
    request: Request,
    db: Session = Depends(get_db),
) -> ClassificationBatchRead:
    workspace_id = _current_workspace_id(request)
    classifier = db.get(DocumentClassifier, payload.classifier_id)
    if not classifier or classifier.archived or (workspace_id is not None and classifier.workspace_id != workspace_id):
        raise HTTPException(status_code=404, detail="Document classifier not found")
    documents = _documents_for_selection(payload.document_ids, db, workspace_id=workspace_id)
    batch, ready_job_ids = _create_classification_batch_from_documents(classifier, documents, db)
    response = _classification_batch_read(batch)
    if ready_job_ids:
        db.close()
        background_tasks.add_task(run_classification_batch, response.id, ready_job_ids)
    return response


@app.post("/api/classification-batches/{batch_id}/items", response_model=ClassificationBatchRead)
async def append_classification_batch_items(batch_id: str, request: Request, db: Session = Depends(get_db)) -> ClassificationBatchRead:
    workspace_id = _current_workspace_id(request)
    form, files = await _read_batch_upload_form(request)
    batch = db.get(ClassificationBatch, batch_id)
    _ensure_workspace_scope(batch, workspace_id, "Classification batch not found")
    if batch.status not in {"uploading", "queued"}:
        raise HTTPException(status_code=409, detail="Classification batch already started")
    await _append_classification_batch_items(batch, form, files, db)
    db.refresh(batch)
    return _classification_batch_read(batch)


@app.post("/api/classification-batches/{batch_id}/start", response_model=ClassificationBatchRead)
def start_classification_batch(
    batch_id: str,
    background_tasks: BackgroundTasks,
    request: Request,
    db: Session = Depends(get_db),
) -> ClassificationBatchRead:
    workspace_id = _current_workspace_id(request)
    batch = db.get(ClassificationBatch, batch_id)
    _ensure_workspace_scope(batch, workspace_id, "Classification batch not found")
    if batch.status not in {"uploading", "queued"}:
        return _classification_batch_read(batch)
    _validate_owner_can_start(batch, batch.items)
    batch.status = "running"
    job_ids = _queued_classification_job_ids(batch)
    log_audit_event(
        db,
        entity_type="classification_batch",
        entity_id=batch.id,
        action="started",
        message=f"Started classification batch with {len(job_ids)} queued job(s)",
        metadata={"classifier_id": batch.classifier_id, "queued_count": len(job_ids)},
    )
    db.commit()
    db.refresh(batch)
    response = _classification_batch_read(batch)
    db.close()
    background_tasks.add_task(run_classification_batch, batch.id, job_ids)
    return response


@app.post("/api/classification-batches/{batch_id}/discard", response_model=ClassificationBatchRead)
def discard_classification_batch(batch_id: str, request: Request, db: Session = Depends(get_db)) -> ClassificationBatchRead:
    workspace_id = _current_workspace_id(request)
    batch = db.get(ClassificationBatch, batch_id)
    _ensure_workspace_scope(batch, workspace_id, "Classification batch not found")
    discarded_count = _discard_batch_items(batch, db)
    batch.status = "canceled"
    batch.completed_at = datetime.utcnow()
    log_audit_event(
        db,
        entity_type="classification_batch",
        entity_id=batch.id,
        action="discarded",
        message=f"Discarded classification batch payloads for {discarded_count} item(s)",
        metadata={"classifier_id": batch.classifier_id, "discarded_count": discarded_count},
    )
    db.commit()
    db.refresh(batch)
    return _classification_batch_read(batch)


@app.post("/api/classification-batches/{batch_id}/resume", response_model=ClassificationBatchRead)
def resume_classification_batch(
    batch_id: str,
    background_tasks: BackgroundTasks,
    request: Request,
    db: Session = Depends(get_db),
) -> ClassificationBatchRead:
    workspace_id = _current_workspace_id(request)
    batch = db.get(ClassificationBatch, batch_id)
    _ensure_workspace_scope(batch, workspace_id, "Classification batch not found")
    if batch.status in {"completed", "completed_with_errors", "failed", "canceled"}:
        raise HTTPException(status_code=409, detail="Classification batch is already terminal")
    if not batch.items:
        raise HTTPException(status_code=422, detail="No uploaded classification items are available to continue")
    _validate_owner_upload_complete(batch, batch.items)
    _prepare_job_batch_resume(batch.items)
    job_ids = _queued_classification_job_ids(batch)
    batch.status = "running" if job_ids else _classification_batch_read(batch).status
    batch.completed_at = None if job_ids else batch.completed_at
    log_audit_event(
        db,
        entity_type="classification_batch",
        entity_id=batch.id,
        action="resumed",
        message=f"Continued classification batch with {len(job_ids)} queued job(s)",
        metadata={"classifier_id": batch.classifier_id, "queued_count": len(job_ids)},
    )
    db.commit()
    db.refresh(batch)
    response = _classification_batch_read(batch)
    if job_ids:
        db.close()
        background_tasks.add_task(run_classification_batch, batch.id, job_ids)
    return response


@app.get("/api/classification-batches", response_model=list[ClassificationBatchRead])
def list_classification_batches(
    request: Request,
    limit: int = Query(default=20, ge=1, le=100),
    include_items: bool = Query(default=True),
    db: Session = Depends(get_db),
) -> list[ClassificationBatchRead]:
    workspace_id = _current_workspace_id(request)
    batches = _scope_query(db.query(ClassificationBatch), ClassificationBatch, workspace_id).order_by(ClassificationBatch.created_at.desc()).limit(limit).all()
    return [_classification_batch_read(batch, include_items=include_items, db=db) for batch in batches]


@app.get("/api/classification-batches/{batch_id}", response_model=ClassificationBatchRead)
def get_classification_batch(batch_id: str, request: Request, db: Session = Depends(get_db)) -> ClassificationBatchRead:
    workspace_id = _current_workspace_id(request)
    batch = db.get(ClassificationBatch, batch_id)
    _ensure_workspace_scope(batch, workspace_id, "Classification batch not found")
    return _classification_batch_read(batch)


@app.get("/api/classification-batches/{batch_id}/summary", response_model=ClassificationBatchRead)
def get_classification_batch_summary(batch_id: str, request: Request, db: Session = Depends(get_db)) -> ClassificationBatchRead:
    workspace_id = _current_workspace_id(request)
    batch = db.get(ClassificationBatch, batch_id)
    _ensure_workspace_scope(batch, workspace_id, "Classification batch not found")
    return _classification_batch_read(batch, include_items=False, db=db)


@app.post("/api/classification-batches/{batch_id}/cancel", response_model=ClassificationBatchRead)
def cancel_classification_batch(batch_id: str, request: Request, db: Session = Depends(get_db)) -> ClassificationBatchRead:
    workspace_id = _current_workspace_id(request)
    batch = db.get(ClassificationBatch, batch_id)
    _ensure_workspace_scope(batch, workspace_id, "Classification batch not found")
    _cancel_module_batch(batch, "classification_batch", db)
    db.commit()
    db.refresh(batch)
    return _classification_batch_read(batch)


@app.get("/api/classification-batches/{batch_id}/export")
def export_classification_batch(
    batch_id: str,
    request: Request,
    format: str = Query(default="csv", pattern="^(json|csv|xlsx)$"),
    db: Session = Depends(get_db),
) -> Response:
    workspace_id = _current_workspace_id(request)
    batch = db.get(ClassificationBatch, batch_id)
    _ensure_workspace_scope(batch, workspace_id, "Classification batch not found")
    artifact = _build_export_artifact(db, "classification_batch", batch.id, format)
    log_audit_event(
        db,
        entity_type="classification_batch",
        entity_id=batch.id,
        action="exported",
        message=f"Exported classification batch {format.upper()}",
        metadata={"format": format},
    )
    db.commit()
    return _export_artifact_response(artifact)


@app.post("/api/required-field-check-batches", response_model=RequiredFieldCheckBatchRead)
async def create_required_field_check_batch(
    request: Request,
    background_tasks: BackgroundTasks,
    db: Session = Depends(get_db),
) -> RequiredFieldCheckBatchRead:
    workspace_id = _current_workspace_id(request)
    form, files = await _read_batch_upload_form(request)
    checklist_id = _required_form_value(form, "checklist_id")
    checklist = db.get(RequiredFieldChecklist, checklist_id)
    if not checklist or checklist.archived or (workspace_id is not None and checklist.workspace_id != workspace_id):
        raise HTTPException(status_code=404, detail="Required field checklist not found")
    if not files:
        raise HTTPException(status_code=400, detail="At least one file is required")
    _validate_upload_file_count(files)

    batch = RequiredFieldCheckBatch(workspace_id=workspace_id, checklist_id=checklist.id, status="uploading", total_count=len(files))
    db.add(batch)
    db.flush()
    log_audit_event(
        db,
        entity_type="required_field_check_batch",
        entity_id=batch.id,
        action="created",
        message=f"Created required field check batch with {len(files)} file(s)",
        metadata={"checklist_id": checklist.id, "file_count": len(files)},
    )
    db.commit()
    await _append_required_field_batch_items(batch, form, files, db)
    db.refresh(batch)
    _validate_owner_can_start(batch, batch.items)
    batch.status = "running"
    job_ids = _queued_required_field_job_ids(batch)
    log_audit_event(
        db,
        entity_type="required_field_check_batch",
        entity_id=batch.id,
        action="started",
        message=f"Started required field check batch with {len(job_ids)} queued job(s)",
        metadata={"checklist_id": checklist.id, "queued_count": len(job_ids)},
    )
    db.commit()
    db.refresh(batch)
    response = _required_field_batch_read(batch)
    db.close()
    background_tasks.add_task(run_required_field_check_batch, batch.id, job_ids)
    return response


@app.post("/api/required-field-check-batches/init", response_model=RequiredFieldCheckBatchRead)
def init_required_field_check_batch(
    payload: RequiredFieldCheckBatchInitRequest,
    request: Request,
    db: Session = Depends(get_db),
) -> RequiredFieldCheckBatchRead:
    workspace_id = _current_workspace_id(request)
    checklist = db.get(RequiredFieldChecklist, payload.checklist_id)
    if not checklist or checklist.archived or (workspace_id is not None and checklist.workspace_id != workspace_id):
        raise HTTPException(status_code=404, detail="Required field checklist not found")
    _validate_declared_batch_file_count(payload.total_count)
    batch = RequiredFieldCheckBatch(workspace_id=workspace_id, checklist_id=checklist.id, status="uploading", total_count=payload.total_count)
    db.add(batch)
    db.flush()
    log_audit_event(
        db,
        entity_type="required_field_check_batch",
        entity_id=batch.id,
        action="upload_initialized",
        message=f"Initialized required field check batch upload with {payload.total_count} file(s)",
        metadata={"checklist_id": checklist.id, "file_count": payload.total_count},
    )
    db.commit()
    db.refresh(batch)
    return _required_field_batch_read(batch)


@app.post("/api/required-field-check-batches/from-documents", response_model=RequiredFieldCheckBatchRead)
def create_required_field_check_batch_from_documents(
    payload: RequiredFieldCheckBatchFromDocumentsRequest,
    background_tasks: BackgroundTasks,
    request: Request,
    db: Session = Depends(get_db),
) -> RequiredFieldCheckBatchRead:
    workspace_id = _current_workspace_id(request)
    checklist = db.get(RequiredFieldChecklist, payload.checklist_id)
    if not checklist or checklist.archived or (workspace_id is not None and checklist.workspace_id != workspace_id):
        raise HTTPException(status_code=404, detail="Required field checklist not found")
    documents = _documents_for_selection(payload.document_ids, db, workspace_id=workspace_id)
    batch, ready_job_ids = _create_required_field_batch_from_documents(checklist, documents, db)
    response = _required_field_batch_read(batch)
    if ready_job_ids:
        db.close()
        background_tasks.add_task(run_required_field_check_batch, response.id, ready_job_ids)
    return response


@app.post("/api/required-field-check-batches/{batch_id}/items", response_model=RequiredFieldCheckBatchRead)
async def append_required_field_check_batch_items(
    batch_id: str,
    request: Request,
    db: Session = Depends(get_db),
) -> RequiredFieldCheckBatchRead:
    workspace_id = _current_workspace_id(request)
    form, files = await _read_batch_upload_form(request)
    batch = db.get(RequiredFieldCheckBatch, batch_id)
    _ensure_workspace_scope(batch, workspace_id, "Required field check batch not found")
    if batch.status not in {"uploading", "queued"}:
        raise HTTPException(status_code=409, detail="Required field check batch already started")
    await _append_required_field_batch_items(batch, form, files, db)
    db.refresh(batch)
    return _required_field_batch_read(batch)


@app.post("/api/required-field-check-batches/{batch_id}/start", response_model=RequiredFieldCheckBatchRead)
def start_required_field_check_batch(
    batch_id: str,
    background_tasks: BackgroundTasks,
    request: Request,
    db: Session = Depends(get_db),
) -> RequiredFieldCheckBatchRead:
    workspace_id = _current_workspace_id(request)
    batch = db.get(RequiredFieldCheckBatch, batch_id)
    _ensure_workspace_scope(batch, workspace_id, "Required field check batch not found")
    if batch.status not in {"uploading", "queued"}:
        return _required_field_batch_read(batch)
    _validate_owner_can_start(batch, batch.items)
    batch.status = "running"
    job_ids = _queued_required_field_job_ids(batch)
    log_audit_event(
        db,
        entity_type="required_field_check_batch",
        entity_id=batch.id,
        action="started",
        message=f"Started required field check batch with {len(job_ids)} queued job(s)",
        metadata={"checklist_id": batch.checklist_id, "queued_count": len(job_ids)},
    )
    db.commit()
    db.refresh(batch)
    response = _required_field_batch_read(batch)
    db.close()
    background_tasks.add_task(run_required_field_check_batch, batch.id, job_ids)
    return response


@app.post("/api/required-field-check-batches/{batch_id}/discard", response_model=RequiredFieldCheckBatchRead)
def discard_required_field_check_batch(batch_id: str, request: Request, db: Session = Depends(get_db)) -> RequiredFieldCheckBatchRead:
    workspace_id = _current_workspace_id(request)
    batch = db.get(RequiredFieldCheckBatch, batch_id)
    _ensure_workspace_scope(batch, workspace_id, "Required field check batch not found")
    discarded_count = _discard_batch_items(batch, db)
    batch.status = "canceled"
    batch.completed_at = datetime.utcnow()
    log_audit_event(
        db,
        entity_type="required_field_check_batch",
        entity_id=batch.id,
        action="discarded",
        message=f"Discarded required field check batch payloads for {discarded_count} item(s)",
        metadata={"checklist_id": batch.checklist_id, "discarded_count": discarded_count},
    )
    db.commit()
    db.refresh(batch)
    return _required_field_batch_read(batch)


@app.post("/api/required-field-check-batches/{batch_id}/resume", response_model=RequiredFieldCheckBatchRead)
def resume_required_field_check_batch(
    batch_id: str,
    background_tasks: BackgroundTasks,
    request: Request,
    db: Session = Depends(get_db),
) -> RequiredFieldCheckBatchRead:
    workspace_id = _current_workspace_id(request)
    batch = db.get(RequiredFieldCheckBatch, batch_id)
    _ensure_workspace_scope(batch, workspace_id, "Required field check batch not found")
    if batch.status in {"completed", "completed_with_errors", "failed", "canceled"}:
        raise HTTPException(status_code=409, detail="Required field check batch is already terminal")
    if not batch.items:
        raise HTTPException(status_code=422, detail="No uploaded required field items are available to continue")
    _validate_owner_upload_complete(batch, batch.items)
    _prepare_job_batch_resume(batch.items)
    job_ids = _queued_required_field_job_ids(batch)
    batch.status = "running" if job_ids else _required_field_batch_read(batch).status
    batch.completed_at = None if job_ids else batch.completed_at
    log_audit_event(
        db,
        entity_type="required_field_check_batch",
        entity_id=batch.id,
        action="resumed",
        message=f"Continued required field check batch with {len(job_ids)} queued job(s)",
        metadata={"checklist_id": batch.checklist_id, "queued_count": len(job_ids)},
    )
    db.commit()
    db.refresh(batch)
    response = _required_field_batch_read(batch)
    if job_ids:
        db.close()
        background_tasks.add_task(run_required_field_check_batch, batch.id, job_ids)
    return response


@app.get("/api/required-field-check-batches", response_model=list[RequiredFieldCheckBatchRead])
def list_required_field_check_batches(
    request: Request,
    limit: int = Query(default=20, ge=1, le=100),
    include_items: bool = Query(default=True),
    db: Session = Depends(get_db),
) -> list[RequiredFieldCheckBatchRead]:
    workspace_id = _current_workspace_id(request)
    batches = _scope_query(db.query(RequiredFieldCheckBatch), RequiredFieldCheckBatch, workspace_id).order_by(RequiredFieldCheckBatch.created_at.desc()).limit(limit).all()
    return [_required_field_batch_read(batch, include_items=include_items, db=db) for batch in batches]


@app.get("/api/required-field-check-batches/{batch_id}", response_model=RequiredFieldCheckBatchRead)
def get_required_field_check_batch(batch_id: str, request: Request, db: Session = Depends(get_db)) -> RequiredFieldCheckBatchRead:
    workspace_id = _current_workspace_id(request)
    batch = db.get(RequiredFieldCheckBatch, batch_id)
    _ensure_workspace_scope(batch, workspace_id, "Required field check batch not found")
    return _required_field_batch_read(batch)


@app.get("/api/required-field-check-batches/{batch_id}/summary", response_model=RequiredFieldCheckBatchRead)
def get_required_field_check_batch_summary(batch_id: str, request: Request, db: Session = Depends(get_db)) -> RequiredFieldCheckBatchRead:
    workspace_id = _current_workspace_id(request)
    batch = db.get(RequiredFieldCheckBatch, batch_id)
    _ensure_workspace_scope(batch, workspace_id, "Required field check batch not found")
    return _required_field_batch_read(batch, include_items=False, db=db)


@app.post("/api/required-field-check-batches/{batch_id}/cancel", response_model=RequiredFieldCheckBatchRead)
def cancel_required_field_check_batch(batch_id: str, request: Request, db: Session = Depends(get_db)) -> RequiredFieldCheckBatchRead:
    workspace_id = _current_workspace_id(request)
    batch = db.get(RequiredFieldCheckBatch, batch_id)
    _ensure_workspace_scope(batch, workspace_id, "Required field check batch not found")
    _cancel_module_batch(batch, "required_field_check_batch", db)
    db.commit()
    db.refresh(batch)
    return _required_field_batch_read(batch)


@app.get("/api/required-field-check-batches/{batch_id}/export")
def export_required_field_check_batch(
    batch_id: str,
    request: Request,
    format: str = Query(default="csv", pattern="^(json|csv|xlsx)$"),
    db: Session = Depends(get_db),
) -> Response:
    workspace_id = _current_workspace_id(request)
    batch = db.get(RequiredFieldCheckBatch, batch_id)
    _ensure_workspace_scope(batch, workspace_id, "Required field check batch not found")
    artifact = _build_export_artifact(db, "required_field_check_batch", batch.id, format)
    log_audit_event(
        db,
        entity_type="required_field_check_batch",
        entity_id=batch.id,
        action="exported",
        message=f"Exported required field check batch {format.upper()}",
        metadata={"format": format},
    )
    db.commit()
    return _export_artifact_response(artifact)


@app.get("/api/archive/search", response_model=list[ArchiveSearchResult])
def archive_search(
    request: Request,
    q: str | None = None,
    status: str | None = None,
    schema_id: str | None = None,
    document_type: str | None = None,
    limit: int = Query(default=20, ge=1, le=100),
    db: Session = Depends(get_db),
) -> list[ArchiveSearchResult]:
    workspace_id = _current_workspace_id(request)
    return _archive_search(db, workspace_id=workspace_id, q=q, status=status, schema_id=schema_id, document_type=document_type, limit=limit)


@app.delete("/api/maintenance/parsing-history")
def clear_parsing_history(db: Session = Depends(get_db)) -> dict[str, Any]:
    _require_workspace_admin_mode()
    settings = get_settings()
    counts = {
        "batches": db.query(Batch).count(),
        "batch_items": db.query(BatchItem).count(),
        "classification_batches": db.query(ClassificationBatch).count(),
        "classification_batch_items": db.query(ClassificationBatchItem).count(),
        "classification_jobs": db.query(ClassificationJob).count(),
        "classification_results": db.query(ClassificationResult).count(),
        "required_field_check_batches": db.query(RequiredFieldCheckBatch).count(),
        "required_field_check_batch_items": db.query(RequiredFieldCheckBatchItem).count(),
        "required_field_check_jobs": db.query(RequiredFieldCheckJob).count(),
        "required_field_check_results": db.query(RequiredFieldCheckResult).count(),
        "workflow_runs": db.query(WorkflowRun).count(),
        "workflow_run_items": db.query(WorkflowRunItem).count(),
        "export_jobs": db.query(ExportJob).count(),
        "documents": db.query(Document).count(),
        "document_pages": db.query(DocumentPage).count(),
        "extraction_jobs": db.query(ExtractionJob).count(),
        "extraction_results": db.query(ExtractionResult).count(),
        "raw_extractions": db.query(RawExtraction).count(),
        "audit_events": db.query(AuditEvent).count(),
        "draft_schemas": db.query(Schema).filter(Schema.ephemeral == True).count(),  # noqa: E712
    }

    ephemeral_schema_ids = [row[0] for row in db.query(Schema.id).filter(Schema.ephemeral == True).all()]  # noqa: E712
    db.query(ClassificationBatchItem).delete(synchronize_session=False)
    db.query(ClassificationBatch).delete(synchronize_session=False)
    db.query(ClassificationResult).delete(synchronize_session=False)
    db.query(ClassificationJob).delete(synchronize_session=False)
    db.query(RequiredFieldCheckBatchItem).delete(synchronize_session=False)
    db.query(RequiredFieldCheckBatch).delete(synchronize_session=False)
    db.query(RequiredFieldCheckResult).delete(synchronize_session=False)
    db.query(RequiredFieldCheckJob).delete(synchronize_session=False)
    db.query(WorkflowRunItem).delete(synchronize_session=False)
    db.query(WorkflowRun).delete(synchronize_session=False)
    db.query(ExportJob).delete(synchronize_session=False)
    db.query(BatchItem).delete(synchronize_session=False)
    db.query(Batch).delete(synchronize_session=False)
    db.query(ExtractionResult).delete(synchronize_session=False)
    db.query(ExtractionJob).delete(synchronize_session=False)
    db.query(DocumentPage).delete(synchronize_session=False)
    db.query(Document).delete(synchronize_session=False)
    db.query(RawExtraction).delete(synchronize_session=False)
    db.query(AuditEvent).delete(synchronize_session=False)
    if ephemeral_schema_ids:
        db.query(Schema).filter(Schema.id.in_(ephemeral_schema_ids)).delete(synchronize_session=False)
    db.commit()

    removed_paths: list[str] = []
    for path in [settings.resolved_storage_dir, settings.resolved_raw_storage_dir]:
        if path.exists():
            shutil.rmtree(path)
            removed_paths.append(str(path))
        path.mkdir(parents=True, exist_ok=True)

    return {"status": "cleared", "counts": counts, "removed_paths": removed_paths}


@app.post("/api/maintenance/retention-cleanup")
def run_retention_cleanup() -> dict[str, Any]:
    _require_workspace_admin_mode()
    return _cleanup_expired_upload_data()


@app.get("/api/audit-events", response_model=list[AuditEventRead])
def list_audit_events(
    request: Request,
    entity_type: str | None = None,
    entity_id: str | None = None,
    limit: int = Query(default=50, ge=1, le=200),
    db: Session = Depends(get_db),
) -> list[AuditEventRead]:
    workspace_id = _current_workspace_id(request)
    query = _scope_query(db.query(AuditEvent), AuditEvent, workspace_id)
    if entity_type:
        query = query.filter(AuditEvent.entity_type == entity_type)
    if entity_id:
        query = query.filter(AuditEvent.entity_id == entity_id)
    events = query.order_by(AuditEvent.created_at.desc()).limit(limit).all()
    return [_audit_event_read(event) for event in events]


def _document_read(document: Document) -> DocumentRead:
    return DocumentRead(
        document_id=document.id,
        filename=document.filename,
        library_path=document.library_path,
        mime_type=document.mime_type,
        size_bytes=document.size_bytes,
        page_count=document.page_count,
        status=document.status,
        error_message=document.error_message,
        document_type=document.document_type,
        language=document.language,
        ai_summary=document.ai_summary,
        recommendation_reasoning=document.recommendation_reasoning,
        created_at=document.created_at,
        deleted_at=document.deleted_at,
        pages=[
            DocumentPageRead(
                id=page.id,
                page=page.page_number,
                image_url=f"/api/documents/{document.id}/pages/{page.page_number}/image",
                width=page.width,
                height=page.height,
            )
            for page in document.pages
        ],
    )


def _load_bank_poc_template() -> dict[str, Any]:
    template_path = PROJECT_ROOT / "docs" / "templates" / "bank-documents-poc.json"
    try:
        return json.loads(template_path.read_text(encoding="utf-8"))
    except OSError as exc:
        raise HTTPException(status_code=500, detail="Bank PoC template file is missing") from exc
    except json.JSONDecodeError as exc:
        raise HTTPException(status_code=500, detail="Bank PoC template file is invalid") from exc


def _seed_bank_poc_schema(db: Session, payload: dict[str, Any], *, workspace_id: str | None) -> tuple[Schema, bool]:
    schema_json = {
        "name": payload["name"],
        "display_name": payload.get("display_name"),
        "description": payload.get("description"),
        "regions": payload.get("regions", []),
        "fields": payload["fields"],
        "is_template": True,
        "template_category": "bank_documents",
        "pinned": True,
    }
    _validate_schema_region_references(schema_json)
    existing = _find_active_by_name(db, Schema, payload["name"], workspace_id=workspace_id)
    legacy_schema = _find_legacy_bank_poc_schema(db, workspace_id=workspace_id)
    if existing:
        _archive_legacy_bank_poc_schema(db, existing.id, workspace_id=workspace_id)
        current_schema_json = existing.schema_json
        existing.display_name = payload.get("display_name")
        existing.description = payload.get("description")
        existing.schema_json = json.dumps(schema_json, ensure_ascii=False)
        existing.is_template = True
        existing.template_category = "bank_documents"
        existing.pinned = True
        existing.archived = False
        if current_schema_json != existing.schema_json:
            existing.current_version = (existing.current_version or 1) + 1
            log_audit_event(
                db,
                entity_type="schema",
                entity_id=existing.id,
                action="updated",
                message=f"Updated schema {existing.name} from bank PoC template",
                metadata={"template": "bank_documents_poc", "field_count": len(payload["fields"])},
            )
        return existing, False
    if legacy_schema:
        current_schema_json = legacy_schema.schema_json
        legacy_schema.name = payload["name"]
        legacy_schema.display_name = payload.get("display_name")
        legacy_schema.description = payload.get("description")
        legacy_schema.schema_json = json.dumps(schema_json, ensure_ascii=False)
        legacy_schema.is_template = True
        legacy_schema.template_category = "bank_documents"
        legacy_schema.pinned = True
        legacy_schema.archived = False
        if current_schema_json != legacy_schema.schema_json:
            legacy_schema.current_version = (legacy_schema.current_version or 1) + 1
        log_audit_event(
            db,
            entity_type="schema",
            entity_id=legacy_schema.id,
            action="updated",
            message=f"Updated legacy schema {legacy_schema.name} from bank PoC template",
            metadata={"template": "bank_documents_poc", "field_count": len(payload["fields"])},
        )
        return legacy_schema, False

    schema = Schema(
        workspace_id=workspace_id,
        name=payload["name"],
        display_name=payload.get("display_name"),
        description=payload.get("description"),
        current_version=1,
        schema_json=json.dumps(schema_json, ensure_ascii=False),
        is_template=True,
        template_category="bank_documents",
        pinned=True,
        ephemeral=False,
        archived=False,
    )
    db.add(schema)
    db.flush()
    log_audit_event(
        db,
        entity_type="schema",
        entity_id=schema.id,
        action="created",
        message=f"Created schema {schema.name} from bank PoC template",
        metadata={"template": "bank_documents_poc", "field_count": len(payload["fields"])},
    )
    return schema, True


def _seed_bank_poc_classifier(db: Session, payload: dict[str, Any], *, workspace_id: str | None) -> tuple[DocumentClassifier, bool]:
    config = {
        "name": payload["name"],
        "description": payload.get("description"),
        "allow_unknown": bool(payload.get("allow_unknown", True)),
        "classes": payload["classes"],
    }
    existing = _find_active_by_name(db, DocumentClassifier, payload["name"], workspace_id=workspace_id)
    if existing:
        next_config_json = json.dumps(config, ensure_ascii=False)
        if (
            existing.description != payload.get("description")
            or existing.allow_unknown != bool(payload.get("allow_unknown", True))
            or existing.config_json != next_config_json
        ):
            existing.description = payload.get("description")
            existing.allow_unknown = bool(payload.get("allow_unknown", True))
            existing.config_json = next_config_json
            log_audit_event(
                db,
                entity_type="document_classifier",
                entity_id=existing.id,
                action="updated",
                message=f"Updated classifier {existing.name} from bank PoC template",
                metadata={"template": "bank_documents_poc", "class_count": len(payload["classes"])},
            )
        return existing, False

    classifier = DocumentClassifier(
        workspace_id=workspace_id,
        name=payload["name"],
        description=payload.get("description"),
        allow_unknown=bool(payload.get("allow_unknown", True)),
        config_json=json.dumps(config, ensure_ascii=False),
        archived=False,
    )
    db.add(classifier)
    db.flush()
    log_audit_event(
        db,
        entity_type="document_classifier",
        entity_id=classifier.id,
        action="created",
        message=f"Created classifier {classifier.name} from bank PoC template",
        metadata={"template": "bank_documents_poc", "class_count": len(payload["classes"])},
    )
    return classifier, True


def _seed_bank_poc_checklist(db: Session, payload: dict[str, Any], *, workspace_id: str | None) -> tuple[RequiredFieldChecklist, bool]:
    config = {
        "name": payload["name"],
        "description": payload.get("description"),
        "regions": payload.get("regions", []),
        "items": payload["items"],
    }
    _validate_checklist_region_references(config)
    existing = _find_active_by_name(db, RequiredFieldChecklist, payload["name"], workspace_id=workspace_id)
    if existing:
        next_config_json = json.dumps(config, ensure_ascii=False)
        if existing.description != payload.get("description") or existing.config_json != next_config_json:
            existing.description = payload.get("description")
            existing.config_json = next_config_json
            log_audit_event(
                db,
                entity_type="required_field_checklist",
                entity_id=existing.id,
                action="updated",
                message=f"Updated checklist {existing.name} from bank PoC template",
                metadata={"template": "bank_documents_poc", "item_count": len(payload["items"])},
            )
        return existing, False

    checklist = RequiredFieldChecklist(
        workspace_id=workspace_id,
        name=payload["name"],
        description=payload.get("description"),
        config_json=json.dumps(config, ensure_ascii=False),
        archived=False,
    )
    db.add(checklist)
    db.flush()
    log_audit_event(
        db,
        entity_type="required_field_checklist",
        entity_id=checklist.id,
        action="created",
        message=f"Created checklist {checklist.name} from bank PoC template",
        metadata={"template": "bank_documents_poc", "item_count": len(payload["items"])},
    )
    return checklist, True


def _seed_bank_poc_workflow(
    db: Session,
    payload: dict[str, Any],
    schema: Schema,
    classifier: DocumentClassifier,
    checklist: RequiredFieldChecklist,
    *,
    workspace_id: str | None,
) -> tuple[WorkflowDefinition, bool]:
    definition = _bank_poc_workflow_definition(schema.id, classifier.id, checklist.id)
    validate_workflow_definition(definition, db, workspace_id=workspace_id)
    existing = _find_active_by_name(db, WorkflowDefinition, payload["name"], workspace_id=workspace_id)
    if existing:
        existing.archived = False
        if existing.description != payload.get("description") or existing.definition_json != json.dumps(definition, ensure_ascii=False):
            existing.description = payload.get("description")
            existing.definition_json = json.dumps(definition, ensure_ascii=False)
            log_audit_event(
                db,
                entity_type="workflow_definition",
                entity_id=existing.id,
                action="updated",
                message=f"Updated workflow {existing.name} from bank PoC template",
                metadata={"template": "bank_documents_poc"},
            )
        return existing, False

    workflow = WorkflowDefinition(
        workspace_id=workspace_id,
        name=payload["name"],
        description=payload.get("description"),
        definition_json=json.dumps(definition, ensure_ascii=False),
        archived=False,
    )
    db.add(workflow)
    db.flush()
    log_audit_event(
        db,
        entity_type="workflow_definition",
        entity_id=workflow.id,
        action="created",
        message=f"Created workflow {workflow.name} from bank PoC template",
        metadata={"template": "bank_documents_poc"},
    )
    return workflow, True


def _seed_bank_poc_sample_documents(
    db: Session,
    samples: list[dict[str, Any]],
    *,
    workspace_id: str | None,
) -> tuple[list[Document], dict[str, bool]]:
    if not samples:
        samples = [
            {"filename": "bank_00006.jpg", "source_path": "assets/sample/bank_00006.jpg"},
            {"filename": "bank_00008.jpg", "source_path": "assets/sample/bank_00008.jpg"},
            {"filename": "bank_00018.jpeg", "source_path": "assets/sample/bank_00018.jpeg"},
        ]
    documents: list[Document] = []
    created: dict[str, bool] = {}
    for sample in samples:
        filename = str(sample["filename"])
        source_path = str(sample["source_path"])
        library_path = f"bank-poc/{filename}"
        existing = (
            _scope_query(db.query(Document), Document, workspace_id)
            .filter(Document.library_path == library_path, Document.status != "deleted")
            .order_by(Document.created_at.desc())
            .first()
        )
        if existing:
            documents.append(existing)
            created[filename] = False
            continue

        local_path = PROJECT_ROOT / source_path
        if not local_path.exists():
            raise HTTPException(status_code=500, detail=f"Bank PoC sample document is missing: {source_path}")

        with local_path.open("rb") as handle:
            upload = LocalUpload(filename=filename, content_type=_sample_content_type(filename), file=handle)
            document = _create_document_from_upload(upload, db, workspace_id=workspace_id)
        document.library_path = library_path
        document.ai_summary = f"은행 PoC 샘플 문서. 출처: {source_path}"
        document.recommendation_reasoning = "사용자가 제공한 로컬 샘플 자산을 데모용으로 보관함에 등록했습니다."
        _ensure_library_folder_records(db, _library_folder_path(library_path), workspace_id=workspace_id)
        log_audit_event(
            db,
            entity_type="document",
            entity_id=document.id,
            action="seeded",
            message="Seeded bank PoC sample document",
            metadata={
                "template": "bank_documents_poc",
                "library_path": library_path,
                "source_path": source_path,
                "source_note": sample.get("source_note", "사용자 제공 로컬 샘플 자산"),
            },
        )
        documents.append(document)
        created[filename] = True
    return documents, created


def _bank_poc_sample_document_read(document: Document) -> dict[str, Any]:
    payload = _document_read(document).model_dump()
    payload["source_path"] = f"assets/sample/{document.filename}"
    payload["source_note"] = "사용자 제공 로컬 샘플 자산"
    return payload


def _sample_content_type(filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix in {".jpg", ".jpeg"}:
        return "image/jpeg"
    if suffix == ".png":
        return "image/png"
    return "application/octet-stream"


def _find_legacy_bank_poc_schema(db: Session, *, workspace_id: str | None) -> Schema | None:
    return _find_active_by_name(db, Schema, "bank_document_core_fields", workspace_id=workspace_id)


def _archive_legacy_bank_poc_schema(db: Session, keep_schema_id: str, *, workspace_id: str | None) -> None:
    legacy = _find_legacy_bank_poc_schema(db, workspace_id=workspace_id)
    if not legacy or legacy.id == keep_schema_id:
        return
    legacy.archived = True
    legacy.pinned = False
    log_audit_event(
        db,
        entity_type="schema",
        entity_id=legacy.id,
        action="archived",
        message="Archived legacy bank PoC schema after Korean schema migration",
        metadata={"template": "bank_documents_poc", "replacement_schema_id": keep_schema_id},
    )


def _find_active_by_name(db: Session, model, name: str, *, workspace_id: str | None):
    query = _scope_query(db.query(model), model, workspace_id).filter(model.name == name)
    if hasattr(model, "archived"):
        query = query.filter(model.archived == False)  # noqa: E712
    return query.order_by(model.created_at.desc()).first()


def _bank_poc_workflow_definition(schema_id: str, classifier_id: str, checklist_id: str) -> dict[str, Any]:
    return {
        "nodes": [
            {"id": "input", "position": {"x": 40, "y": 240}, "data": {"kind": "input", "label": "문서 입력"}},
            {
                "id": "classifier",
                "position": {"x": 250, "y": 240},
                "data": {"kind": "classifier", "label": "은행 문서 분류", "config": {"classifier_id": classifier_id}},
            },
            {
                "id": "branch",
                "position": {"x": 460, "y": 210},
                "data": {
                    "kind": "branch",
                    "label": "문서 종류 분기",
                    "branchKeys": [
                        "class:신청서",
                        "class:동의서",
                        "class:증빙문서",
                        "unknown",
                    ],
                },
            },
            {
                "id": "kie_application",
                "position": {"x": 670, "y": 70},
                "data": {"kind": "kie", "label": "신청서 핵심 정보 추출", "config": {"schema_id": schema_id}},
            },
            {
                "id": "required_application",
                "position": {"x": 900, "y": 70},
                "data": {"kind": "required-checker", "label": "신청서 필수 항목 확인", "config": {"checklist_id": checklist_id}},
            },
            {
                "id": "required_consent",
                "position": {"x": 670, "y": 240},
                "data": {"kind": "required-checker", "label": "동의서 필수 항목 확인", "config": {"checklist_id": checklist_id}},
            },
            {
                "id": "kie_supporting",
                "position": {"x": 670, "y": 410},
                "data": {"kind": "kie", "label": "증빙 정보 추출", "config": {"schema_id": schema_id}},
            },
            {"id": "merge", "position": {"x": 900, "y": 300}, "data": {"kind": "merge", "label": "결과 병합"}},
            {"id": "export", "position": {"x": 1110, "y": 300}, "data": {"kind": "export", "label": "Export"}},
        ],
        "edges": [
            {"id": "input-classifier", "source": "input", "target": "classifier"},
            {"id": "classifier-branch", "source": "classifier", "target": "branch"},
            {"id": "branch-application", "source": "branch", "target": "kie_application", "sourceHandle": "class:신청서"},
            {"id": "branch-consent", "source": "branch", "target": "required_consent", "sourceHandle": "class:동의서"},
            {"id": "branch-supporting", "source": "branch", "target": "kie_supporting", "sourceHandle": "class:증빙문서"},
            {"id": "branch-unknown", "source": "branch", "target": "merge", "sourceHandle": "unknown"},
            {"id": "application-kie-required", "source": "kie_application", "target": "required_application"},
            {"id": "application-required-merge", "source": "required_application", "target": "merge"},
            {"id": "consent-required-merge", "source": "required_consent", "target": "merge"},
            {"id": "supporting-kie-merge", "source": "kie_supporting", "target": "merge"},
            {"id": "merge-export", "source": "merge", "target": "export"},
        ],
    }


def _repair_image_document_if_needed(document: Document, db: Session) -> None:
    source_path = materialize_storage_ref(document.storage_path)
    if not is_supported_image(source_path) or not source_path.exists() or not document.pages:
        return

    page = document.pages[0]
    try:
        source_width, source_height = read_image_size(source_path)
    except DocumentProcessingError:
        return
    page_path = materialize_storage_ref(page.image_path)
    if page.width == source_width and page.height == source_height and page_path.exists():
        return

    try:
        page_info = rasterize_image_page(source_path, scratch_dir_for_ref(page.image_path, "repair"))
    except DocumentProcessingError:
        return

    next_width = int(page_info["width"])
    next_height = int(page_info["height"])
    if page.width == next_width and page.height == next_height and page_path.exists():
        return

    if get_settings().storage_backend.strip().lower() == "s3":
        image_path = Path(str(page_info["image_path"]))
        page.image_path = persist_artifact(image_path, f"documents/{document.id}/pages/{image_path.name}", _image_media_type(image_path))
    else:
        page.image_path = str(page_info["image_path"])
    page.width = next_width
    page.height = next_height
    document.page_count = 1
    db.commit()
    db.refresh(document)


def _raw_extraction_read(raw: RawExtraction) -> RawExtractionRead:
    return RawExtractionRead(
        id=raw.id,
        filename=raw.filename,
        source_format=raw.source_format,
        size_bytes=raw.size_bytes,
        status=raw.status,
        pdf_url=f"/api/raw-extractions/{raw.id}/pdf" if raw.pdf_path else None,
        html_url=f"/api/raw-extractions/{raw.id}/html" if raw.html_path else None,
        warnings=json.loads(raw.warnings or "[]"),
        error_message=raw.error_message,
        created_at=raw.created_at,
        updated_at=raw.updated_at,
    )


def _create_document_from_upload(file: UploadFile, db: Session, *, workspace_id: str | None = None) -> Document:
    try:
        document, original_path = _create_uploaded_document(file, db, workspace_id=workspace_id)
        _preprocess_document_pages(document, original_path, db)
    except DocumentProcessingError as exc:
        raise HTTPException(status_code=exc.status_code, detail=str(exc)) from exc
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Failed to process uploaded document") from exc
    return document


def _persist_document_artifacts(original_path: Path, pages: list[dict[str, int | str]]) -> tuple[str, list[dict[str, int | str]]]:
    settings = get_settings()
    if settings.storage_backend.strip().lower() != "s3":
        return str(original_path), pages
    document_key = original_path.parent.name
    original_ref = _persist_original_artifact(original_path)
    return original_ref, _persist_document_pages(document_key, pages)


def _create_uploaded_document(file: UploadFile, db: Session, *, workspace_id: str | None = None) -> tuple[Document, Path]:
    filename, original_path, size_bytes = save_upload_file(file)
    document = Document(
        workspace_id=workspace_id,
        filename=filename,
        mime_type=file.content_type or "application/octet-stream",
        size_bytes=size_bytes,
        page_count=0,
        storage_path=_persist_original_artifact(original_path),
        status="preprocessing",
    )
    db.add(document)
    db.flush()
    return document, original_path


def _create_failed_upload_document(file: UploadFile, message: str, db: Session, *, workspace_id: str | None = None) -> Document:
    return _create_failed_upload_document_record(file.filename or "upload", file.content_type or "application/octet-stream", message, db, workspace_id=workspace_id)


def _create_failed_upload_document_record(filename: str, mime_type: str, message: str, db: Session, *, workspace_id: str | None = None) -> Document:
    document = Document(
        workspace_id=workspace_id,
        filename=filename,
        mime_type=mime_type,
        size_bytes=0,
        page_count=0,
        storage_path="",
        status="failed",
        error_message=message,
    )
    db.add(document)
    db.flush()
    return document


def _preprocess_document_pages(document: Document, original_path: Path, db: Session, *, raise_errors: bool = True) -> bool:
    try:
        pages = rasterize_document(original_path)
        pages = _persist_document_pages(original_path.parent.name, pages)
    except DocumentProcessingError as exc:
        document.status = "failed"
        document.error_message = str(exc)
        if raise_errors:
            raise
        return False
    except Exception as exc:
        document.status = "failed"
        document.error_message = "Failed to process uploaded document"
        if raise_errors:
            raise DocumentProcessingError("Failed to process uploaded document") from exc
        return False

    for existing in list(document.pages):
        db.delete(existing)
    db.flush()
    persisted_pages: list[dict[str, int | str]] = []
    for page in pages:
        image_path = Path(str(page["image_path"]))
        persisted_pages.append({**page, "image_path": str(image_path)})
        db.add(
            DocumentPage(
                document_id=document.id,
                page_number=int(page["page_number"]),
                image_path=str(page["image_path"]),
                width=int(page["width"]),
                height=int(page["height"]),
            )
        )
    document.page_count = len(persisted_pages)
    document.status = "ready"
    document.error_message = None
    return True


def _upload_library_paths(form: FormData, files: list[UploadFile]) -> list[str | None]:
    raw_values = list(form.getlist("library_paths") or form.getlist("relative_paths") or [])
    paths = [str(value) if value is not None else "" for value in raw_values]
    if len(paths) != len(files):
        paths = [getattr(file, "filename", "") or "" for file in files]
    return [_normalize_library_path(path) or (file.filename or "") for path, file in zip(paths, files, strict=False)]


def _normalize_library_path(path: str | None) -> str:
    if not path:
        return ""
    normalized = str(path).replace("\\", "/").strip().strip("/")
    parts = [part.strip() for part in normalized.split("/") if part.strip() and part.strip() not in {".", ".."}]
    return "/".join(parts)


def _library_folder_path(path: str | None) -> str:
    normalized = _normalize_library_path(path)
    if not normalized or "/" not in normalized:
        return ""
    return normalized.rsplit("/", 1)[0]


def _library_folder_ancestors(path: str) -> list[str]:
    normalized = _normalize_library_path(path)
    if not normalized:
        return [""]
    parts = normalized.split("/")
    ancestors = [""]
    for index in range(1, len(parts) + 1):
        ancestors.append("/".join(parts[:index]))
    return ancestors


def _library_basename(path: str | None, fallback: str) -> str:
    normalized = _normalize_library_path(path)
    if not normalized:
        return fallback
    return normalized.rsplit("/", 1)[-1] or fallback


def _join_library_path(folder_path: str, name: str) -> str:
    folder = _normalize_library_path(folder_path)
    clean_name = _normalize_library_path(name).rsplit("/", 1)[-1] if name else ""
    return f"{folder}/{clean_name}" if folder else clean_name


def _document_library_query(
    db: Session,
    *,
    include_deleted: bool = False,
    status: str | None = None,
    q: str | None = None,
    library_path: str | None = None,
    workspace_id: str | None = None,
):
    query = _scope_query(db.query(Document), Document, workspace_id)
    if not include_deleted:
        query = query.filter(Document.status != "deleted")
    if status:
        query = query.filter(Document.status == status)
    if q:
        pattern = f"%{q.strip()}%"
        query = query.filter((Document.filename.ilike(pattern)) | (Document.library_path.ilike(pattern)))
    if library_path is not None:
        normalized_path = _normalize_library_path(library_path)
        if normalized_path:
            query = query.filter(Document.library_path.ilike(f"{normalized_path}/%"))
        else:
            query = query.filter((Document.library_path.is_(None)) | (Document.library_path == "") | (~Document.library_path.contains("/")))
    return query


def _ensure_library_folder_records(db: Session, folder_path: str, *, workspace_id: str | None = None) -> None:
    for path in _library_folder_ancestors(folder_path):
        if not path:
            continue
        if not _library_folder_record_exists(db, path, workspace_id=workspace_id):
            db.add(DocumentLibraryFolder(workspace_id=workspace_id, path=path))


def _library_folder_record_exists(db: Session, path: str, *, workspace_id: str | None = None) -> bool:
    if any(isinstance(item, DocumentLibraryFolder) and item.path == path and item.workspace_id == workspace_id for item in db.new):
        return True
    if any(isinstance(item, DocumentLibraryFolder) and item.path == path and item.workspace_id == workspace_id for item in db.dirty):
        return True
    query = db.query(DocumentLibraryFolder.id).filter(DocumentLibraryFolder.path == path)
    query = _scope_query(query, DocumentLibraryFolder, workspace_id)
    return query.one_or_none() is not None


def _normalize_library_clipboard_payload(payload: DocumentLibraryClipboardRequest) -> tuple[list[str], list[str], str]:
    document_ids = list(dict.fromkeys(document_id for document_id in payload.document_ids if document_id))
    folder_paths = list(dict.fromkeys(_normalize_library_path(path) for path in payload.folder_paths if _normalize_library_path(path)))
    target_folder = _normalize_library_path(payload.target_folder)
    if not document_ids and not folder_paths:
        raise HTTPException(status_code=422, detail="document_ids or folder_paths is required")
    return document_ids, folder_paths, target_folder


def _folder_path_contains(folder_path: str, candidate: str) -> bool:
    folder = _normalize_library_path(folder_path)
    target = _normalize_library_path(candidate)
    return bool(folder and (target == folder or target.startswith(f"{folder}/")))


def _existing_library_folder_paths(db: Session, *, workspace_id: str | None = None) -> set[str]:
    paths = {row[0] for row in _scope_query(db.query(DocumentLibraryFolder.path), DocumentLibraryFolder, workspace_id).all()}
    for row in _scope_query(db.query(Document.library_path), Document, workspace_id).filter(Document.status != "deleted", Document.library_path.isnot(None)).all():
        folder_path = _library_folder_path(row[0])
        paths.update(_library_folder_ancestors(folder_path))
    paths.discard("")
    return paths


def _copy_name(name: str, attempt: int) -> str:
    path = Path(name)
    suffix = path.suffix
    stem = path.name.removesuffix(suffix) if suffix else path.name
    marker = "copy" if attempt == 1 else f"copy {attempt}"
    return f"{stem} {marker}{suffix}"


def _unique_library_document_path(db: Session, desired_path: str, *, exclude_document_ids: set[str] | None = None, workspace_id: str | None = None) -> str:
    normalized = _normalize_library_path(desired_path)
    folder = _library_folder_path(normalized)
    name = _library_basename(normalized, "document")
    excluded = exclude_document_ids or set()
    candidate = normalized
    attempt = 0
    while True:
        query = _scope_query(db.query(Document.id), Document, workspace_id).filter(Document.status != "deleted", Document.library_path == candidate)
        if excluded:
            query = query.filter(~Document.id.in_(excluded))
        if query.first() is None:
            return candidate
        attempt += 1
        candidate = _join_library_path(folder, _copy_name(name, attempt))


def _unique_library_folder_path(db: Session, desired_path: str, *, exclude_prefixes: set[str] | None = None, workspace_id: str | None = None) -> str:
    normalized = _normalize_library_path(desired_path)
    parent = _library_folder_path(normalized)
    name = _library_basename(normalized, "folder")
    excluded = exclude_prefixes or set()
    existing_paths = _existing_library_folder_paths(db, workspace_id=workspace_id)
    existing_paths = {
        path
        for path in existing_paths
        if not any(path == excluded_path or path.startswith(f"{excluded_path}/") for excluded_path in excluded if excluded_path)
    }
    candidate = normalized
    attempt = 0
    while candidate in existing_paths:
        attempt += 1
        candidate = _join_library_path(parent, _copy_name(name, attempt))
    return candidate


def _documents_in_library_folder(db: Session, folder_path: str, *, workspace_id: str | None = None) -> list[Document]:
    normalized = _normalize_library_path(folder_path)
    if not normalized:
        return []
    return (
        _scope_query(db.query(Document), Document, workspace_id)
        .filter(Document.status.notin_(["deleted", "failed"]), Document.library_path.ilike(f"{normalized}/%"))
        .order_by(Document.library_path.asc(), Document.created_at.asc(), Document.id.asc())
        .all()
    )


def _selected_library_documents(db: Session, document_ids: list[str], *, workspace_id: str | None = None) -> list[Document]:
    if not document_ids:
        return []
    documents = _scope_query(db.query(Document), Document, workspace_id).filter(Document.id.in_(document_ids)).all()
    by_id = {document.id: document for document in documents}
    missing = [document_id for document_id in document_ids if document_id not in by_id]
    if missing:
        raise HTTPException(status_code=404, detail=f"Document not found: {missing[0]}")
    blocked = [document for document in documents if document.status in {"deleted", "failed"}]
    if blocked:
        raise HTTPException(status_code=409, detail=f"Document cannot be copied or moved: {blocked[0].filename}")
    return [by_id[document_id] for document_id in document_ids]


def _move_library_entries(payload: DocumentLibraryClipboardRequest, db: Session, *, workspace_id: str | None = None) -> DocumentLibraryActionRead:
    document_ids, folder_paths, target_folder = _normalize_library_clipboard_payload(payload)
    if target_folder:
        _ensure_library_folder_records(db, target_folder, workspace_id=workspace_id)
    moved_documents: list[Document] = []
    moved_document_ids: set[str] = set()

    for folder_path in folder_paths:
        if _folder_path_contains(folder_path, target_folder):
            raise HTTPException(status_code=409, detail="A folder cannot be moved into itself")
        source_parent = _library_folder_path(folder_path)
        if source_parent == target_folder:
            new_folder_path = folder_path
        else:
            desired_folder_path = _join_library_path(target_folder, _library_basename(folder_path, "folder"))
            new_folder_path = _unique_library_folder_path(db, desired_folder_path, exclude_prefixes={folder_path}, workspace_id=workspace_id)
        if new_folder_path != folder_path:
            _move_library_folder_records(db, folder_path, new_folder_path, workspace_id=workspace_id)
            for document in _documents_in_library_folder(db, folder_path, workspace_id=workspace_id):
                suffix = _normalize_library_path(document.library_path).removeprefix(f"{folder_path}/")
                document.library_path = _unique_library_document_path(db, _join_library_path(new_folder_path, suffix), exclude_document_ids={document.id}, workspace_id=workspace_id)
                moved_documents.append(document)
                moved_document_ids.add(document.id)
        _ensure_library_folder_records(db, new_folder_path, workspace_id=workspace_id)

    for document in _selected_library_documents(db, document_ids, workspace_id=workspace_id):
        if document.id in moved_document_ids:
            continue
        basename = _library_basename(document.library_path, document.filename)
        document.library_path = _unique_library_document_path(db, _join_library_path(target_folder, basename), exclude_document_ids={document.id}, workspace_id=workspace_id)
        moved_documents.append(document)
        moved_document_ids.add(document.id)

    log_audit_event(
        db,
        entity_type="document_library",
        entity_id=target_folder or "root",
        action="moved",
        message=f"Moved {len(moved_document_ids)} document(s) in the document library",
        metadata={"target_folder": target_folder, "document_count": len(moved_document_ids), "folder_paths": folder_paths},
    )
    db.flush()
    return DocumentLibraryActionRead(
            documents=[_document_read(document) for document in moved_documents],
            folders=_document_tree_read(db, workspace_id=workspace_id).folders,
        )


def _move_library_folder_records(db: Session, source_folder: str, target_folder: str, *, workspace_id: str | None = None) -> None:
    source = _normalize_library_path(source_folder)
    target = _normalize_library_path(target_folder)
    rows = (
        _scope_query(db.query(DocumentLibraryFolder), DocumentLibraryFolder, workspace_id)
        .filter((DocumentLibraryFolder.path == source) | (DocumentLibraryFolder.path.ilike(f"{source}/%")))
        .order_by(DocumentLibraryFolder.path.asc())
        .all()
    )
    for row in rows:
        suffix = row.path.removeprefix(source).lstrip("/")
        row.path = _join_library_path(target, suffix) if suffix else target


def _copy_library_entries(payload: DocumentLibraryClipboardRequest, db: Session, *, workspace_id: str | None = None) -> tuple[DocumentLibraryActionRead, list[str]]:
    document_ids, folder_paths, target_folder = _normalize_library_clipboard_payload(payload)
    if target_folder:
        _ensure_library_folder_records(db, target_folder, workspace_id=workspace_id)
    copied_documents: list[Document] = []
    conversion_job_ids: list[str] = []

    for folder_path in folder_paths:
        if _folder_path_contains(folder_path, target_folder):
            raise HTTPException(status_code=409, detail="A folder cannot be copied into itself")
        desired_folder_path = _join_library_path(target_folder, _library_basename(folder_path, "folder"))
        new_folder_path = _unique_library_folder_path(db, desired_folder_path, workspace_id=workspace_id)
        _ensure_library_folder_records(db, new_folder_path, workspace_id=workspace_id)
        for explicit_folder in sorted(
            [path for path in _existing_library_folder_paths(db, workspace_id=workspace_id) if path == folder_path or path.startswith(f"{folder_path}/")]
        ):
            suffix = explicit_folder.removeprefix(folder_path).lstrip("/")
            _ensure_library_folder_records(db, _join_library_path(new_folder_path, suffix) if suffix else new_folder_path, workspace_id=workspace_id)
        for document in _documents_in_library_folder(db, folder_path, workspace_id=workspace_id):
            suffix = _normalize_library_path(document.library_path).removeprefix(f"{folder_path}/")
            copied_document, conversion_job_id = _copy_library_document(
                document,
                _unique_library_document_path(db, _join_library_path(new_folder_path, suffix), workspace_id=workspace_id),
                db,
            )
            copied_documents.append(copied_document)
            if conversion_job_id:
                conversion_job_ids.append(conversion_job_id)

    for document in _selected_library_documents(db, document_ids, workspace_id=workspace_id):
        basename = _library_basename(document.library_path, document.filename)
        copied_document, conversion_job_id = _copy_library_document(
            document,
            _unique_library_document_path(db, _join_library_path(target_folder, basename), workspace_id=workspace_id),
            db,
        )
        copied_documents.append(copied_document)
        if conversion_job_id:
            conversion_job_ids.append(conversion_job_id)

    log_audit_event(
        db,
        entity_type="document_library",
        entity_id=target_folder or "root",
        action="copied",
        message=f"Copied {len(copied_documents)} document(s) in the document library",
        metadata={"target_folder": target_folder, "document_count": len(copied_documents), "folder_paths": folder_paths},
    )
    db.flush()
    return (
        DocumentLibraryActionRead(
            documents=[_document_read(document) for document in copied_documents],
            folders=_document_tree_read(db, workspace_id=workspace_id).folders,
        ),
        conversion_job_ids,
    )


def _copy_library_document(source: Document, library_path: str, db: Session) -> tuple[Document, str | None]:
    if source.status in {"deleted", "failed"}:
        raise HTTPException(status_code=409, detail=f"Document cannot be copied: {source.filename}")
    source_original = materialize_storage_ref(source.storage_path)
    if not source_original.exists():
        raise HTTPException(status_code=404, detail=f"Document payload is missing: {source.filename}")

    local_dir = get_settings().resolved_storage_dir / uuid4().hex
    local_dir.mkdir(parents=True, exist_ok=True)
    original_suffix = source_original.suffix or Path(source.filename).suffix
    local_original = local_dir / f"original{original_suffix or '.bin'}"
    shutil.copy2(source_original, local_original)

    copied = Document(
        workspace_id=source.workspace_id,
        filename=source.filename,
        mime_type=source.mime_type,
        size_bytes=source.size_bytes,
        page_count=0,
        storage_path="",
        library_path=_normalize_library_path(library_path) or source.filename,
        status="queued" if source.status in DOCUMENT_CONVERTING_STATUSES else source.status,
        error_message=None,
        document_type=source.document_type,
        language=source.language,
        ai_summary=source.ai_summary,
        recommendation_reasoning=source.recommendation_reasoning,
    )
    db.add(copied)
    db.flush()
    copied.storage_path = persist_artifact(local_original, f"documents/{copied.id}/original{local_original.suffix}")

    if source.status == "ready":
        page_dir = local_dir / "pages"
        page_dir.mkdir(parents=True, exist_ok=True)
        for page in source.pages:
            source_page = materialize_storage_ref(page.image_path)
            if not source_page.exists():
                raise HTTPException(status_code=404, detail=f"Document page payload is missing: {source.filename}")
            local_page = page_dir / source_page.name
            if local_page.exists():
                local_page = page_dir / f"page_{page.page_number:04d}{source_page.suffix or '.jpg'}"
            shutil.copy2(source_page, local_page)
            page_ref = persist_artifact(local_page, f"documents/{copied.id}/pages/{local_page.name}", _image_media_type(local_page))
            db.add(
                DocumentPage(
                    document_id=copied.id,
                    page_number=page.page_number,
                    image_path=page_ref,
                    width=page.width,
                    height=page.height,
                )
            )
        copied.page_count = source.page_count
        return copied, None

    copied.page_count = 0
    copied.status = "queued"
    job = DocumentConversionJob(document_id=copied.id, status="queued")
    db.add(job)
    db.flush()
    return copied, job.id


def _delete_library_document_payload(document: Document, db: Session) -> Document:
    if document.status == "deleted":
        return document
    _delete_document_storage(document)
    for page in list(document.pages):
        db.delete(page)
    document.page_count = 0
    document.status = "deleted"
    document.error_message = "Original document was deleted from the document library"
    document.deleted_at = datetime.utcnow()
    return document


def _create_queued_library_document(file: UploadFile, db: Session, *, library_path: str | None = None, workspace_id: str | None = None) -> tuple[Document, str]:
    filename, original_path, size_bytes = save_upload_file(file)
    display_path = _normalize_library_path(library_path) or filename
    document = Document(
        workspace_id=workspace_id,
        filename=filename,
        mime_type=file.content_type or "application/octet-stream",
        size_bytes=size_bytes,
        page_count=0,
        storage_path=_persist_original_artifact(original_path),
        library_path=display_path,
        status="queued",
    )
    db.add(document)
    db.flush()
    job = DocumentConversionJob(document_id=document.id, status="queued")
    db.add(job)
    db.flush()
    return document, job.id


def _enqueue_document_conversion_job(job_id: str) -> None:
    with _document_conversion_queue_lock:
        if job_id in _document_conversion_enqueued:
            return
        _document_conversion_enqueued.add(job_id)
    _document_conversion_queue.put(job_id)


def _start_document_conversion_workers() -> threading.Event:
    stop_event = threading.Event()
    _enqueue_pending_document_conversion_jobs()
    worker_count = max(1, int(get_settings().preprocess_max_workers or 1))

    def _run() -> None:
        while not stop_event.is_set():
            try:
                job_id = _document_conversion_queue.get(timeout=0.5)
            except queue.Empty:
                continue
            try:
                _run_document_conversion_job(job_id)
            finally:
                with _document_conversion_queue_lock:
                    _document_conversion_enqueued.discard(job_id)
                _document_conversion_queue.task_done()

    for index in range(worker_count):
        thread = threading.Thread(target=_run, name=f"document-conversion-{index + 1}", daemon=True)
        thread.start()
    return stop_event


def _enqueue_pending_document_conversion_jobs() -> None:
    db = SessionLocal()
    try:
        pending = (
            db.query(DocumentConversionJob)
            .filter(DocumentConversionJob.status.in_(["queued", "running"]))
            .order_by(DocumentConversionJob.created_at.asc(), DocumentConversionJob.id.asc())
            .all()
        )
        for job in pending:
            if job.status == "running":
                job.status = "queued"
                job.started_at = None
        db.commit()
        for job in pending:
            _enqueue_document_conversion_job(job.id)
    finally:
        db.close()


def _run_document_conversion_job(job_id: str) -> None:
    db = SessionLocal()
    document_id: str | None = None
    try:
        job = db.get(DocumentConversionJob, job_id)
        if not job or job.status not in {"queued", "running"}:
            return
        document = db.get(Document, job.document_id)
        if not document or document.status == "deleted":
            job.status = "canceled"
            job.completed_at = datetime.utcnow()
            db.commit()
            return
        document_id = document.id
        now = datetime.utcnow()
        job.status = "running"
        job.started_at = job.started_at or now
        job.attempts = (job.attempts or 0) + 1
        document.status = "preprocessing"
        document.error_message = None
        db.commit()

        original_path = materialize_storage_ref(document.storage_path)
        if not original_path.exists():
            raise DocumentProcessingError("Original document file is missing", status_code=404)
        _preprocess_document_pages(document, original_path, db, raise_errors=True)
        job.status = "completed"
        job.error_message = None
        job.completed_at = datetime.utcnow()
        log_audit_event(
            db,
            entity_type="document",
            entity_id=document.id,
            action="ready",
            message=f"Document conversion completed for {document.filename}",
            metadata={"conversion_job_id": job.id, "page_count": document.page_count},
        )
        db.commit()
    except DocumentProcessingError as exc:
        db.rollback()
        _mark_document_conversion_failed(job_id, str(exc))
    except Exception:
        db.rollback()
        _mark_document_conversion_failed(job_id, "Failed to process uploaded document")
    finally:
        db.close()
    if document_id:
        _activate_ready_document_waiters(document_id)


def _mark_document_conversion_failed(job_id: str, message: str) -> None:
    db = SessionLocal()
    try:
        job = db.get(DocumentConversionJob, job_id)
        if not job:
            return
        document = db.get(Document, job.document_id)
        job.status = "failed"
        job.error_message = message
        job.completed_at = datetime.utcnow()
        if document and document.status != "deleted":
            document.status = "failed"
            document.error_message = message
        log_audit_event(
            db,
            entity_type="document_conversion_job",
            entity_id=job.id,
            action="failed",
            message=message,
            metadata={"document_id": job.document_id},
        )
        db.commit()
    finally:
        db.close()


def _activate_ready_document_waiters(document_id: str) -> None:
    extraction_batches: list[tuple[str, str]] = []
    classification_batches: list[tuple[str, str]] = []
    required_batches: list[tuple[str, str]] = []
    extraction_jobs: list[str] = []
    classification_jobs: list[str] = []
    required_jobs: list[str] = []
    workflow_runs: list[tuple[str, int]] = []
    db = SessionLocal()
    try:
        document = db.get(Document, document_id)
        if not document or document.status != "ready":
            return

        for job in db.query(ExtractionJob).filter(ExtractionJob.document_id == document_id, ExtractionJob.status == "waiting_for_document").all():
            job.status = "queued"
            job.error_message = None
            item = db.query(BatchItem).filter(BatchItem.job_id == job.id).one_or_none()
            if item and item.batch and item.batch.status == "running":
                extraction_batches.append((item.batch_id, job.id))
            elif not item:
                extraction_jobs.append(job.id)

        for job in db.query(ClassificationJob).filter(ClassificationJob.document_id == document_id, ClassificationJob.status == "waiting_for_document").all():
            job.status = "queued"
            job.error_message = None
            item = db.query(ClassificationBatchItem).filter(ClassificationBatchItem.job_id == job.id).one_or_none()
            if item and item.batch and item.batch.status == "running":
                classification_batches.append((item.batch_id, job.id))
            elif not item:
                classification_jobs.append(job.id)

        for job in db.query(RequiredFieldCheckJob).filter(RequiredFieldCheckJob.document_id == document_id, RequiredFieldCheckJob.status == "waiting_for_document").all():
            job.status = "queued"
            job.error_message = None
            item = db.query(RequiredFieldCheckBatchItem).filter(RequiredFieldCheckBatchItem.job_id == job.id).one_or_none()
            if item and item.batch and item.batch.status == "running":
                required_batches.append((item.batch_id, job.id))
            elif not item:
                required_jobs.append(job.id)

        for item in db.query(WorkflowRunItem).filter(WorkflowRunItem.document_id == document_id, WorkflowRunItem.status == "waiting_for_document").all():
            run = item.run
            if not run or run.status != "running":
                continue
            item.status = "queued"
            item.error_message = None
            item.completed_at = None
            item.execution_generation = run.execution_generation
            workflow_runs.append((run.id, run.execution_generation))

        db.commit()
    finally:
        db.close()

    for batch_id, job_id in extraction_batches:
        threading.Thread(target=run_batch_jobs, args=(batch_id, [job_id]), daemon=True).start()
    for batch_id, job_id in classification_batches:
        threading.Thread(target=run_classification_batch, args=(batch_id, [job_id]), daemon=True).start()
    for batch_id, job_id in required_batches:
        threading.Thread(target=run_required_field_check_batch, args=(batch_id, [job_id]), daemon=True).start()
    for job_id in extraction_jobs:
        threading.Thread(target=run_extraction_job, args=(job_id,), daemon=True).start()
    for job_id in classification_jobs:
        threading.Thread(target=run_classification_job, args=(job_id,), daemon=True).start()
    for job_id in required_jobs:
        threading.Thread(target=run_required_field_check_job, args=(job_id,), daemon=True).start()
    for run_id, generation in sorted(set(workflow_runs)):
        threading.Thread(target=run_workflow_run, args=(run_id, generation), daemon=True).start()


def _execution_job_status_for_document(document: Document) -> str:
    if document.status == "ready":
        return "queued"
    if document.status in DOCUMENT_CONVERTING_STATUSES:
        return "waiting_for_document"
    return "blocked"


def _document_not_executable_message(document: Document) -> str:
    if document.status == "deleted":
        return "Original document was deleted"
    if document.status == "failed":
        return document.error_message or "Document conversion failed"
    return f"Document is not executable in status {document.status}"


def _raise_document_not_executable(document: Document) -> None:
    if document.status == "deleted":
        raise HTTPException(status_code=410, detail=_document_not_executable_message(document))
    raise HTTPException(status_code=422, detail=_document_not_executable_message(document))


def _documents_for_selection(document_ids: list[str], db: Session, *, workspace_id: str | None = None) -> list[Document]:
    ordered_ids: list[str] = []
    seen: set[str] = set()
    for document_id in document_ids:
        if document_id not in seen:
            ordered_ids.append(document_id)
            seen.add(document_id)
    documents = _scope_query(db.query(Document), Document, workspace_id).filter(Document.id.in_(ordered_ids)).all()
    by_id = {document.id: document for document in documents}
    missing = [document_id for document_id in ordered_ids if document_id not in by_id]
    if missing:
        raise HTTPException(status_code=404, detail={"message": "Some documents were not found", "document_ids": missing})
    return [by_id[document_id] for document_id in ordered_ids]


def _create_extraction_batch_from_documents(schema: Schema, documents: list[Document], db: Session) -> tuple[Batch, list[str]]:
    batch = Batch(workspace_id=schema.workspace_id, schema_id=schema.id, schema_version=1, status="running", total_count=len(documents))
    db.add(batch)
    db.flush()
    ready_job_ids: list[str] = []
    waiting_count = 0
    for index, document in enumerate(documents):
        job_status = _execution_job_status_for_document(document)
        error_message = None
        completed_at = None
        if job_status == "blocked":
            job_status = "failed"
            error_message = _document_not_executable_message(document)
            completed_at = datetime.utcnow()
        elif job_status == "queued":
            ready_job_ids.append("")
        else:
            waiting_count += 1
        job = ExtractionJob(
            workspace_id=schema.workspace_id,
            document_id=document.id,
            schema_id=schema.id,
            schema_version=1,
            status=job_status,
            error_message=error_message,
            completed_at=completed_at,
        )
        db.add(job)
        db.flush()
        if job_status == "queued":
            ready_job_ids[-1] = job.id
        db.add(
            BatchItem(
                batch_id=batch.id,
                document_id=document.id,
                job_id=job.id,
                filename=document.filename,
                upload_index=index,
                client_file_id=f"library:{document.id}",
            )
        )
    _finish_batch_immediately_if_no_active_documents(batch, waiting_count, ready_job_ids)
    log_audit_event(
        db,
        entity_type="batch",
        entity_id=batch.id,
        action="created_from_documents",
        message=f"Created KIE batch from {len(documents)} library document(s)",
        metadata={"schema_id": schema.id, "ready_count": len(ready_job_ids), "waiting_count": waiting_count},
    )
    db.commit()
    db.refresh(batch)
    return batch, ready_job_ids


def _create_classification_batch_from_documents(classifier: DocumentClassifier, documents: list[Document], db: Session) -> tuple[ClassificationBatch, list[str]]:
    batch = ClassificationBatch(workspace_id=classifier.workspace_id, classifier_id=classifier.id, status="running", total_count=len(documents))
    db.add(batch)
    db.flush()
    ready_job_ids: list[str] = []
    waiting_count = 0
    for index, document in enumerate(documents):
        job_status = _execution_job_status_for_document(document)
        error_message = None
        completed_at = None
        if job_status == "blocked":
            job_status = "failed"
            error_message = _document_not_executable_message(document)
            completed_at = datetime.utcnow()
        elif job_status == "queued":
            ready_job_ids.append("")
        else:
            waiting_count += 1
        job = ClassificationJob(
            workspace_id=classifier.workspace_id,
            document_id=document.id,
            classifier_id=classifier.id,
            status=job_status,
            error_message=error_message,
            completed_at=completed_at,
        )
        db.add(job)
        db.flush()
        if job_status == "queued":
            ready_job_ids[-1] = job.id
        db.add(
            ClassificationBatchItem(
                batch_id=batch.id,
                document_id=document.id,
                job_id=job.id,
                filename=document.filename,
                upload_index=index,
                client_file_id=f"library:{document.id}",
            )
        )
    _finish_batch_immediately_if_no_active_documents(batch, waiting_count, ready_job_ids)
    log_audit_event(
        db,
        entity_type="classification_batch",
        entity_id=batch.id,
        action="created_from_documents",
        message=f"Created classification batch from {len(documents)} library document(s)",
        metadata={"classifier_id": classifier.id, "ready_count": len(ready_job_ids), "waiting_count": waiting_count},
    )
    db.commit()
    db.refresh(batch)
    return batch, ready_job_ids


def _create_required_field_batch_from_documents(checklist: RequiredFieldChecklist, documents: list[Document], db: Session) -> tuple[RequiredFieldCheckBatch, list[str]]:
    batch = RequiredFieldCheckBatch(workspace_id=checklist.workspace_id, checklist_id=checklist.id, status="running", total_count=len(documents))
    db.add(batch)
    db.flush()
    ready_job_ids: list[str] = []
    waiting_count = 0
    for index, document in enumerate(documents):
        job_status = _execution_job_status_for_document(document)
        error_message = None
        completed_at = None
        if job_status == "blocked":
            job_status = "failed"
            error_message = _document_not_executable_message(document)
            completed_at = datetime.utcnow()
        elif job_status == "queued":
            ready_job_ids.append("")
        else:
            waiting_count += 1
        job = RequiredFieldCheckJob(
            workspace_id=checklist.workspace_id,
            document_id=document.id,
            checklist_id=checklist.id,
            status=job_status,
            error_message=error_message,
            completed_at=completed_at,
        )
        db.add(job)
        db.flush()
        if job_status == "queued":
            ready_job_ids[-1] = job.id
        db.add(
            RequiredFieldCheckBatchItem(
                batch_id=batch.id,
                document_id=document.id,
                job_id=job.id,
                filename=document.filename,
                upload_index=index,
                client_file_id=f"library:{document.id}",
            )
        )
    _finish_batch_immediately_if_no_active_documents(batch, waiting_count, ready_job_ids)
    log_audit_event(
        db,
        entity_type="required_field_check_batch",
        entity_id=batch.id,
        action="created_from_documents",
        message=f"Created required field check batch from {len(documents)} library document(s)",
        metadata={"checklist_id": checklist.id, "ready_count": len(ready_job_ids), "waiting_count": waiting_count},
    )
    db.commit()
    db.refresh(batch)
    return batch, ready_job_ids


def _finish_batch_immediately_if_no_active_documents(batch: Any, waiting_count: int, ready_job_ids: list[str]) -> None:
    if ready_job_ids or waiting_count:
        return
    batch.status = "completed_with_errors"
    batch.completed_at = datetime.utcnow()


def _create_workflow_run_from_documents(workflow: WorkflowDefinition, documents: list[Document], db: Session) -> tuple[WorkflowRun, bool]:
    now = datetime.utcnow()
    run = WorkflowRun(
        workspace_id=workflow.workspace_id,
        workflow_id=workflow.id,
        workflow_name=workflow.name,
        workflow_definition_json=workflow.definition_json,
        status="running",
        total_count=len(documents),
        created_at=now,
        started_at=now,
        inference_started_at=now,
        execution_generation=1,
    )
    db.add(run)
    db.flush()
    has_ready_items = False
    has_waiting_items = False
    for index, document in enumerate(documents):
        item_status = _execution_job_status_for_document(document)
        error_message = None
        completed_at = None
        if item_status == "blocked":
            item_status = "failed"
            error_message = _document_not_executable_message(document)
            completed_at = now
        elif item_status == "queued":
            has_ready_items = True
        else:
            has_waiting_items = True
        db.add(
            WorkflowRunItem(
                run_id=run.id,
                document_id=document.id,
                filename=document.filename,
                upload_index=index,
                status=item_status,
                error_message=error_message,
                client_file_id=f"library:{document.id}",
                execution_generation=run.execution_generation,
                result_json=json.dumps(
                    {
                        "document_id": document.id,
                        "filename": document.filename,
                        "node_results": {},
                        "error_message": error_message,
                    },
                    ensure_ascii=False,
                ),
                completed_at=completed_at,
            )
        )
    if not has_ready_items and not has_waiting_items:
        run.status = "completed_with_errors"
        run.completed_at = now
        run.inference_started_at = None
    log_audit_event(
        db,
        entity_type="workflow_run",
        entity_id=run.id,
        action="created_from_documents",
        message=f"Created workflow run from {len(documents)} library document(s)",
        metadata={"workflow_id": workflow.id, "ready_count": has_ready_items, "waiting_count": has_waiting_items},
    )
    db.commit()
    db.refresh(run)
    return run, has_ready_items


def _persist_original_artifact(original_path: Path) -> str:
    if get_settings().storage_backend.strip().lower() != "s3":
        return str(original_path)
    document_key = original_path.parent.name
    return persist_artifact(original_path, f"documents/{document_key}/original{original_path.suffix}")


def _persist_document_pages(document_key: str, pages: list[dict[str, int | str]]) -> list[dict[str, int | str]]:
    if get_settings().storage_backend.strip().lower() != "s3":
        return pages
    persisted_pages: list[dict[str, int | str]] = []
    for page in pages:
        image_path = Path(str(page["image_path"]))
        page_ref = persist_artifact(image_path, f"documents/{document_key}/pages/{image_path.name}", _image_media_type(image_path))
        persisted_pages.append({**page, "image_path": page_ref})
    return persisted_pages


def _image_media_type(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".jpg", ".jpeg"}:
        return "image/jpeg"
    if suffix == ".png":
        return "image/png"
    return "application/octet-stream"


def _schema_read(schema: Schema) -> SchemaRead:
    schema_data = _schema_data(schema)
    return SchemaRead(
        id=schema.id,
        name=schema.name,
        display_name=schema.display_name,
        description=schema.description,
        is_template=schema.is_template,
        template_category=schema.template_category,
        pinned=schema.pinned,
        ephemeral=schema.ephemeral,
        archived=schema.archived,
        regions=schema_data.get("regions", []),
        fields=schema_data["fields"],
        created_at=schema.created_at,
        updated_at=schema.updated_at,
    )


def _schema_data(schema: Schema) -> dict[str, Any]:
    if not schema.schema_json or schema.schema_json == "{}":
        raise HTTPException(status_code=500, detail="Schema data is missing")
    return json.loads(schema.schema_json)


def _classifier_read(classifier: DocumentClassifier) -> DocumentClassifierRead:
    data = _classifier_data(classifier)
    return DocumentClassifierRead(
        id=classifier.id,
        name=classifier.name,
        description=classifier.description,
        allow_unknown=classifier.allow_unknown,
        archived=classifier.archived,
        classes=data["classes"],
        created_at=classifier.created_at,
        updated_at=classifier.updated_at,
    )


def _classifier_data(classifier: DocumentClassifier) -> dict[str, Any]:
    if not classifier.config_json or classifier.config_json == "{}":
        raise HTTPException(status_code=500, detail="Document classifier data is missing")
    return json.loads(classifier.config_json)


def _classification_job_read(job: ClassificationJob) -> ClassificationJobRead:
    return ClassificationJobRead(
        job_id=job.id,
        document_id=job.document_id,
        classifier_id=job.classifier_id,
        status=job.status,
        error_message=job.error_message,
        result_id=job.result_id,
        result=ClassificationResultRead(**classification_result_to_dict(job.result)) if job.result else None,
        created_at=job.created_at,
        started_at=job.started_at,
        completed_at=job.completed_at,
    )


def _checklist_read(checklist: RequiredFieldChecklist) -> RequiredFieldChecklistRead:
    data = _checklist_data(checklist)
    return RequiredFieldChecklistRead(
        id=checklist.id,
        name=checklist.name,
        description=checklist.description,
        archived=checklist.archived,
        regions=data.get("regions", []),
        items=data["items"],
        created_at=checklist.created_at,
        updated_at=checklist.updated_at,
    )


def _checklist_data(checklist: RequiredFieldChecklist) -> dict[str, Any]:
    if not checklist.config_json or checklist.config_json == "{}":
        raise HTTPException(status_code=500, detail="Required field checklist data is missing")
    return json.loads(checklist.config_json)


def _required_field_checklist_recommendation_read(payload: dict[str, Any]) -> RequiredFieldChecklistRecommendationRead:
    recommendation = RequiredFieldChecklistRecommendationRead(**payload)
    region_ids = {region.id for region in recommendation.regions}
    seen_items: set[str] = set()
    unique_items = []
    for item in recommendation.items:
        if item.item_name in seen_items:
            continue
        seen_items.add(item.item_name)
        if item.region_id and item.region_id not in region_ids:
            item = item.model_copy(update={"region_id": None})
        unique_items.append(item)
    if not unique_items:
        raise ValueError("checklist recommendation must include at least one item")
    return RequiredFieldChecklistRecommendationRead(
        name=recommendation.name.strip() or "ai_recommended_checklist",
        description=recommendation.description,
        reasoning=recommendation.reasoning,
        regions=recommendation.regions,
        items=unique_items,
    )


def _save_workflow_ai_draft_images(files: list[UploadFile]) -> list[Path]:
    if not files:
        raise HTTPException(status_code=422, detail="At least one sample image is required")
    if len(files) > WORKFLOW_AI_DRAFT_MAX_IMAGES:
        raise HTTPException(status_code=413, detail=f"AI workflow draft supports up to {WORKFLOW_AI_DRAFT_MAX_IMAGES} sample images")

    temp_dir = Path(tempfile.mkdtemp(prefix="workflow_ai_draft_"))
    saved_paths: list[Path] = []
    try:
        for index, file in enumerate(files, start=1):
            suffix = Path(file.filename or "").suffix.lower()
            if suffix not in WORKFLOW_AI_DRAFT_IMAGE_EXTENSIONS:
                raise HTTPException(status_code=400, detail="Only PNG, JPG, and JPEG sample images are supported")
            image_path = temp_dir / f"sample_{index}{suffix}"
            size = 0
            with image_path.open("wb") as destination:
                while True:
                    chunk = file.file.read(1024 * 1024)
                    if not chunk:
                        break
                    size += len(chunk)
                    if size > get_settings().upload_max_file_bytes:
                        raise HTTPException(status_code=413, detail="Uploaded sample image is too large")
                    destination.write(chunk)
            if size <= 0:
                raise HTTPException(status_code=400, detail="Uploaded sample image is empty")
            read_image_size(image_path)
            saved_paths.append(image_path)
        return saved_paths
    except HTTPException:
        shutil.rmtree(temp_dir, ignore_errors=True)
        raise
    except DocumentProcessingError as exc:
        shutil.rmtree(temp_dir, ignore_errors=True)
        raise HTTPException(status_code=400, detail=str(exc)) from exc


def _workflow_ai_draft_definition(*, include_checklist: bool) -> dict[str, Any]:
    nodes = [
        {
            "id": "ai_input",
            "type": "workflow",
            "position": {"x": 0, "y": 150},
            "data": {"kind": "input", "label": "문서 입력", "config": {}},
        },
        {
            "id": "ai_kie",
            "type": "workflow",
            "position": {"x": 270, "y": 110},
            "data": {"kind": "kie", "label": "핵심 정보 추출", "config": {}},
        },
        {
            "id": "ai_merge",
            "type": "workflow",
            "position": {"x": 820 if include_checklist else 560, "y": 150},
            "data": {"kind": "merge", "label": "결과 병합", "config": {}},
        },
        {
            "id": "ai_export",
            "type": "workflow",
            "position": {"x": 1080 if include_checklist else 800, "y": 150},
            "data": {"kind": "export", "label": "Export", "config": {}},
        },
    ]
    if include_checklist:
        nodes.insert(
            2,
            {
                "id": "ai_required",
                "type": "workflow",
                "position": {"x": 560, "y": 110},
                "data": {"kind": "required-checker", "label": "필수 항목 확인", "config": {}},
            },
        )
        edges = [
            {"id": "ai_input-out-ai_kie", "source": "ai_input", "target": "ai_kie"},
            {"id": "ai_kie-out-ai_required", "source": "ai_kie", "target": "ai_required"},
            {"id": "ai_required-out-ai_merge", "source": "ai_required", "target": "ai_merge"},
            {"id": "ai_merge-out-ai_export", "source": "ai_merge", "target": "ai_export"},
        ]
    else:
        edges = [
            {"id": "ai_input-out-ai_kie", "source": "ai_input", "target": "ai_kie"},
            {"id": "ai_kie-out-ai_merge", "source": "ai_kie", "target": "ai_merge"},
            {"id": "ai_merge-out-ai_export", "source": "ai_merge", "target": "ai_export"},
        ]
    return {"nodes": nodes, "edges": edges}


def _workflow_ai_draft_name(schema_draft: SchemaCreate, *, sample_count: int) -> str:
    base = schema_draft.display_name or schema_draft.name or "AI 생성 워크플로우"
    return f"{base} 워크플로우 초안 · 샘플 {sample_count}장"


def _available_scoped_name(db: Session, model, desired: str, *, workspace_id: str | None) -> str:
    base = desired.strip() or "AI 초안"
    query = db.query(model.name).filter(model.name == base, model.archived == False)  # noqa: E712
    if hasattr(model, "ephemeral"):
        query = query.filter(model.ephemeral == False)  # noqa: E712
    if workspace_id is not None and hasattr(model, "workspace_id"):
        query = query.filter(model.workspace_id == workspace_id)
    if not query.first():
        return base
    existing_names = {
        row[0]
        for row in _scope_query(
            db.query(model.name).filter(model.archived == False),  # noqa: E712
            model,
            workspace_id,
        ).all()
    }
    return _duplicate_name(base, existing_names)


def _validate_checklist_region_references(checklist_data: dict[str, Any]) -> None:
    item_names = [item.get("item_name") for item in checklist_data.get("items", []) if isinstance(item, dict)]
    if len(item_names) != len(set(item_names)):
        raise HTTPException(status_code=422, detail="required field item_name values must be unique")
    region_ids = [region.get("id") for region in checklist_data.get("regions", []) if isinstance(region, dict)]
    if len(region_ids) != len(set(region_ids)):
        raise HTTPException(status_code=422, detail="required field region ids must be unique")
    region_id_set = set(region_ids)
    missing_region_ids = sorted(
        {
            item.get("region_id")
            for item in checklist_data.get("items", [])
            if isinstance(item, dict) and item.get("region_id") and item.get("region_id") not in region_id_set
        }
    )
    if missing_region_ids:
        raise HTTPException(
            status_code=422,
            detail=f"required field item region_id values are missing from regions: {', '.join(missing_region_ids)}",
        )


def _required_field_job_read(job: RequiredFieldCheckJob) -> RequiredFieldCheckJobRead:
    return RequiredFieldCheckJobRead(
        job_id=job.id,
        document_id=job.document_id,
        checklist_id=job.checklist_id,
        status=job.status,
        error_message=job.error_message,
        result_id=job.result_id,
        result=RequiredFieldCheckResultRead(**required_field_result_to_dict(job.result)) if job.result else None,
        created_at=job.created_at,
        started_at=job.started_at,
        completed_at=job.completed_at,
    )


def _validate_schema_region_references(schema_data: dict[str, Any]) -> None:
    region_ids = {region.get("id") for region in schema_data.get("regions", []) if isinstance(region, dict)}
    missing_region_ids = sorted(
        {
            field.get("region_id")
            for field in schema_data.get("fields", [])
            if isinstance(field, dict) and field.get("region_id") and field.get("region_id") not in region_ids
        }
    )
    if missing_region_ids:
        raise HTTPException(
            status_code=422,
            detail=f"schema field region_id values are missing from regions: {', '.join(missing_region_ids)}",
        )


def _raise_if_schema_name_conflicts(db: Session, name: str, schema_id: str | None = None, workspace_id: str | None = None) -> None:
    normalized = name.strip()
    if not normalized:
        return
    query = db.query(Schema).filter(Schema.name == normalized, Schema.ephemeral == False, Schema.archived == False)  # noqa: E712
    if workspace_id is not None:
        query = query.filter(Schema.workspace_id == workspace_id)
    if schema_id:
        query = query.filter(Schema.id != schema_id)
    if query.first():
        raise HTTPException(status_code=409, detail=f"Schema name already exists: {normalized}")


def _duplicate_name(name: str, existing_names: set[str], max_length: int = 120) -> str:
    base = name.strip() or "schema"
    for index in range(1, 10000):
        suffix = f" ({index})"
        truncated_base = base[: max(1, max_length - len(suffix))].rstrip()
        candidate = f"{truncated_base}{suffix}"
        if candidate not in existing_names:
            return candidate
    raise HTTPException(status_code=409, detail=f"Could not create a duplicate name for: {base}")


def _merge_duplicate_schema_names_into(db: Session, schema: Schema, name: str, workspace_id: str | None = None) -> None:
    normalized = name.strip()
    if not normalized:
        return
    duplicates = (
        db.query(Schema)
        .filter(Schema.name == normalized, Schema.ephemeral == False, Schema.archived == False, Schema.id != schema.id)  # noqa: E712
    )
    if workspace_id is not None:
        duplicates = duplicates.filter(Schema.workspace_id == workspace_id)
    for duplicate in duplicates.all():
        db.query(ExtractionJob).filter(ExtractionJob.schema_id == duplicate.id).update(
            {ExtractionJob.schema_id: schema.id},
            synchronize_session=False,
        )
        db.query(Batch).filter(Batch.schema_id == duplicate.id).update(
            {Batch.schema_id: schema.id},
            synchronize_session=False,
        )
        db.query(ExportPreset).filter(ExportPreset.schema_id == duplicate.id).update(
            {ExportPreset.schema_id: schema.id},
            synchronize_session=False,
        )
        log_audit_event(
            db,
            entity_type="schema",
            entity_id=schema.id,
            action="merged_duplicate",
            message=f"Merged duplicate schema {duplicate.id} into {schema.id}",
            metadata={"duplicate_schema_id": duplicate.id, "name": normalized},
        )
        db.delete(duplicate)


def _schema_recommendation_read(payload: dict[str, Any]) -> SchemaRecommendationRead:
    recommendation = SchemaRecommendationRead(**payload)
    seen: set[str] = set()
    unique_fields = []
    for field in recommendation.fields:
        if field.key_name in seen:
            continue
        seen.add(field.key_name)
        unique_fields.append(field)
    return SchemaRecommendationRead(
        name=recommendation.name.strip() or "ai_recommended_schema",
        display_name=recommendation.display_name,
        description=recommendation.description,
        document_type=recommendation.document_type,
        language=recommendation.language,
        reasoning=recommendation.reasoning,
        fields=unique_fields,
    )


def _job_read(job: ExtractionJob) -> ExtractionJobRead:
    return ExtractionJobRead(
        job_id=job.id,
        document_id=job.document_id,
        schema_id=job.schema_id,
        status=job.status,
        error_message=job.error_message,
        result_id=job.result_id,
        result=result_to_dict(job.result) if job.result else None,
        created_at=job.created_at,
        started_at=job.started_at,
        completed_at=job.completed_at,
    )


def _batch_read(batch: Batch, *, include_items: bool = True, db: Session | None = None) -> BatchRead:
    if include_items or db is None:
        items = [_batch_item_read(item) for item in _sorted_batch_items(batch.items)]
        counters = _owner_counters(batch.total_count, [item.status for item in items], batch.status)
    else:
        items = []
        counters = _owner_counters_from_status_counts(
            batch.total_count,
            _module_batch_status_counts(db, BatchItem, ExtractionJob, "batch_id", batch.id),
            batch.status,
        )
    return BatchRead(
        id=batch.id,
        schema_id=batch.schema_id,
        status=counters["status"],
        total_count=batch.total_count,
        completed_count=counters["completed_count"],
        failed_count=counters["failed_count"],
        canceled_count=counters["canceled_count"],
        uploaded_count=counters["uploaded_count"],
        preprocessing_count=counters["preprocessing_count"],
        ready_count=counters["ready_count"],
        queued_count=counters["queued_count"],
        running_count=counters["running_count"],
        needs_review_count=counters["needs_review_count"],
        progress_phase=counters["progress_phase"],
        progress=counters["progress"],
        items=items if include_items else [],
        created_at=batch.created_at,
        completed_at=batch.completed_at,
    )


def _batch_item_read(item: BatchItem) -> BatchItemRead:
    return BatchItemRead(
        id=item.id,
        document_id=item.document_id,
        job_id=item.job_id,
        filename=item.filename,
        upload_index=item.upload_index,
        status=item.job.status if item.job else "unknown",
        result_id=item.job.result_id if item.job else None,
        error_message=item.job.error_message if item.job else None,
        created_at=item.created_at,
    )


def _classification_batch_read(batch: ClassificationBatch, *, include_items: bool = True, db: Session | None = None) -> ClassificationBatchRead:
    if include_items or db is None:
        items = [_classification_batch_item_read(item) for item in _sorted_module_items(batch.items)]
        counters = _owner_counters(batch.total_count, [item.status for item in items], batch.status)
    else:
        items = []
        counters = _owner_counters_from_status_counts(
            batch.total_count,
            _module_batch_status_counts(db, ClassificationBatchItem, ClassificationJob, "batch_id", batch.id),
            batch.status,
        )
    return ClassificationBatchRead(
        id=batch.id,
        classifier_id=batch.classifier_id,
        status=counters["status"],
        total_count=batch.total_count,
        completed_count=counters["completed_count"],
        failed_count=counters["failed_count"],
        canceled_count=counters["canceled_count"],
        uploaded_count=counters["uploaded_count"],
        preprocessing_count=counters["preprocessing_count"],
        ready_count=counters["ready_count"],
        queued_count=counters["queued_count"],
        running_count=counters["running_count"],
        needs_review_count=counters["needs_review_count"],
        progress_phase=counters["progress_phase"],
        progress=counters["progress"],
        items=items if include_items else [],
        created_at=batch.created_at,
        completed_at=batch.completed_at,
    )


def _classification_batch_item_read(item: ClassificationBatchItem) -> ClassificationBatchItemRead:
    return ClassificationBatchItemRead(
        id=item.id,
        document_id=item.document_id,
        job_id=item.job_id,
        filename=item.filename,
        upload_index=item.upload_index,
        status=item.job.status if item.job else "unknown",
        result_id=item.job.result_id if item.job else None,
        error_message=item.job.error_message if item.job else None,
        created_at=item.created_at,
    )


def _required_field_batch_read(batch: RequiredFieldCheckBatch, *, include_items: bool = True, db: Session | None = None) -> RequiredFieldCheckBatchRead:
    if include_items or db is None:
        items = [_required_field_batch_item_read(item) for item in _sorted_module_items(batch.items)]
        counters = _owner_counters(batch.total_count, [item.status for item in items], batch.status)
    else:
        items = []
        counters = _owner_counters_from_status_counts(
            batch.total_count,
            _module_batch_status_counts(db, RequiredFieldCheckBatchItem, RequiredFieldCheckJob, "batch_id", batch.id),
            batch.status,
        )
    return RequiredFieldCheckBatchRead(
        id=batch.id,
        checklist_id=batch.checklist_id,
        status=counters["status"],
        total_count=batch.total_count,
        completed_count=counters["completed_count"],
        failed_count=counters["failed_count"],
        canceled_count=counters["canceled_count"],
        uploaded_count=counters["uploaded_count"],
        preprocessing_count=counters["preprocessing_count"],
        ready_count=counters["ready_count"],
        queued_count=counters["queued_count"],
        running_count=counters["running_count"],
        needs_review_count=counters["needs_review_count"],
        progress_phase=counters["progress_phase"],
        progress=counters["progress"],
        items=items if include_items else [],
        created_at=batch.created_at,
        completed_at=batch.completed_at,
    )


def _required_field_batch_item_read(item: RequiredFieldCheckBatchItem) -> RequiredFieldCheckBatchItemRead:
    return RequiredFieldCheckBatchItemRead(
        id=item.id,
        document_id=item.document_id,
        job_id=item.job_id,
        filename=item.filename,
        upload_index=item.upload_index,
        status=item.job.status if item.job else "unknown",
        result_id=item.job.result_id if item.job else None,
        error_message=item.job.error_message if item.job else None,
        created_at=item.created_at,
    )


def _module_batch_status_counts(db: Session, item_model: Any, job_model: Any, owner_field: str, owner_id: str) -> dict[str, int]:
    owner_column = getattr(item_model, owner_field)
    rows = (
        db.query(job_model.status, func.count(item_model.id))
        .select_from(item_model)
        .join(job_model, item_model.job_id == job_model.id)
        .filter(owner_column == owner_id)
        .group_by(job_model.status)
        .all()
    )
    return {status: int(count) for status, count in rows}


def _owner_counters(total_count: int, statuses: list[str], owner_status: str) -> dict[str, Any]:
    return _owner_counters_from_status_counts(total_count, dict(Counter(statuses)), owner_status)


def _owner_counters_from_status_counts(total_count: int, status_counts: dict[str, int], owner_status: str) -> dict[str, Any]:
    completed_statuses = {"completed", "needs_review"}
    terminal_statuses = {"completed", "needs_review", "failed", "canceled"}
    uploaded_count = sum(status_counts.values())
    preprocessing_count = sum(status_counts.get(status, 0) for status in {"uploading", "preprocessing", "waiting_for_document"})
    queued_count = status_counts.get("queued", 0)
    running_count = status_counts.get("running", 0)
    needs_review_count = status_counts.get("needs_review", 0)
    completed_count = sum(status_counts.get(status, 0) for status in completed_statuses)
    failed_count = status_counts.get("failed", 0)
    canceled_count = status_counts.get("canceled", 0)
    finished_count = sum(status_counts.get(status, 0) for status in terminal_statuses)
    ready_count = max(0, uploaded_count - preprocessing_count)
    if owner_status in {"canceled", "failed"} and not uploaded_count:
        status = owner_status
        progress_phase = owner_status
    elif total_count and finished_count >= total_count:
        if canceled_count == total_count:
            status = "canceled"
        elif failed_count or canceled_count:
            status = "completed_with_errors"
        else:
            status = "completed"
        progress_phase = status
    elif uploaded_count < total_count:
        status = "uploading"
        progress_phase = "uploading"
    elif preprocessing_count:
        status = "preprocessing"
        progress_phase = "preprocessing"
    elif running_count or owner_status == "running":
        status = "running"
        progress_phase = "running"
    elif canceled_count:
        status = "canceling"
        progress_phase = "running"
    elif queued_count:
        status = "queued"
        progress_phase = "queued"
    else:
        status = owner_status
        progress_phase = owner_status
    progress = finished_count / total_count if total_count else 0
    return {
        "status": status,
        "uploaded_count": uploaded_count,
        "preprocessing_count": preprocessing_count,
        "ready_count": ready_count,
        "queued_count": queued_count,
        "running_count": running_count,
        "completed_count": completed_count,
        "failed_count": failed_count,
        "canceled_count": canceled_count,
        "needs_review_count": needs_review_count,
        "progress_phase": progress_phase,
        "progress": progress,
    }


def _cancel_module_batch(batch: Any, entity_type: str, db: Session) -> None:
    canceled_count = 0
    now = datetime.utcnow()
    for item in batch.items:
        if item.job and item.job.status in {"queued", "running", "waiting_for_document"}:
            item.job.status = "canceled"
            item.job.error_message = "Canceled by user"
            item.job.completed_at = now
            canceled_count += 1
    if canceled_count:
        _close_batch_if_all_jobs_terminal(batch, now)
        if batch.status not in {"canceled", "completed", "completed_with_errors"}:
            batch.status = "cancel_requested"
        action = "cancel_requested"
        message = f"Cancel requested for {canceled_count} running or queued job(s)"
        metadata = {"canceled_count": canceled_count}
    else:
        action = "cancel_skipped"
        message = "No running or queued batch jobs to cancel"
        metadata = {}
    log_audit_event(db, entity_type=entity_type, entity_id=batch.id, action=action, message=message, metadata=metadata)


def _close_batch_if_all_jobs_terminal(batch: Any, completed_at: datetime) -> None:
    terminal_statuses = {"completed", "needs_review", "failed", "canceled"}
    jobs = [item.job for item in batch.items if item.job]
    if not jobs or any(job.status not in terminal_statuses for job in jobs):
        return
    statuses = [job.status for job in jobs]
    if all(status == "canceled" for status in statuses):
        batch.status = "canceled"
    elif any(status in {"failed", "canceled"} for status in statuses):
        batch.status = "completed_with_errors"
    else:
        batch.status = "completed"
    batch.completed_at = completed_at


def _create_restarted_workflow_run(
    source_run: WorkflowRun,
    workflow: WorkflowDefinition,
    now: datetime,
    db: Session,
) -> tuple[WorkflowRun, int]:
    new_run = WorkflowRun(
        workspace_id=source_run.workspace_id,
        workflow_id=workflow.id,
        workflow_name=workflow.name,
        workflow_definition_json=workflow.definition_json,
        restarted_from_run_id=source_run.id,
        status="running",
        total_count=source_run.total_count,
        created_at=now,
        upload_duration_ms=source_run.upload_duration_ms or _workflow_upload_duration_ms(source_run),
        started_at=now,
        inference_started_at=now,
        execution_generation=1,
    )
    db.add(new_run)
    db.flush()

    queued_count = 0
    for source_item in _sorted_workflow_items(source_run.items):
        document = source_item.document or db.get(Document, source_item.document_id)
        ready = bool(document and document.status == "ready" and document.pages)
        status = "queued" if ready else "failed"
        message = None if ready else (document.error_message if document else source_item.error_message or "Document preprocessing was interrupted")
        result_json = (
            _initial_workflow_item_payload(document)
            if ready and document
            else json.dumps(
                {
                    "document_id": source_item.document_id,
                    "filename": source_item.filename,
                    "node_results": {},
                    "error_message": message,
                },
                ensure_ascii=False,
            )
        )
        db.add(
            WorkflowRunItem(
                run_id=new_run.id,
                document_id=source_item.document_id,
                filename=source_item.filename,
                upload_index=source_item.upload_index,
                status=status,
                error_message=message,
                client_file_id=source_item.client_file_id,
                upload_duration_ms=source_item.upload_duration_ms,
                execution_generation=new_run.execution_generation,
                result_json=result_json,
                completed_at=None if ready else now,
            )
        )
        if ready:
            queued_count += 1
    return new_run, queued_count


def _create_waiting_workflow_run(
    source_run: WorkflowRun,
    workflow: WorkflowDefinition,
    now: datetime,
    db: Session,
) -> tuple[WorkflowRun, int]:
    group_id = _ensure_workflow_queue_group(source_run, db)
    queue_order = _next_workflow_queue_order(db, group_id)
    new_run = WorkflowRun(
        workspace_id=source_run.workspace_id,
        workflow_id=workflow.id,
        workflow_name=workflow.name,
        workflow_definition_json=workflow.definition_json,
        workflow_run_group_id=group_id,
        queued_from_run_id=source_run.id,
        queue_order=queue_order,
        status="waiting",
        total_count=source_run.total_count,
        created_at=now,
        upload_duration_ms=source_run.upload_duration_ms or _workflow_upload_duration_ms(source_run),
        execution_generation=0,
    )
    db.add(new_run)
    db.flush()

    queued_count = 0
    for source_item in _sorted_workflow_items(source_run.items):
        document = source_item.document or db.get(Document, source_item.document_id)
        ready = bool(document and document.status == "ready" and document.pages)
        status = "queued" if ready else "failed"
        message = None if ready else (document.error_message if document else source_item.error_message or "Document preprocessing was interrupted")
        result_json = (
            _initial_workflow_item_payload(document)
            if ready and document
            else json.dumps(
                {
                    "document_id": source_item.document_id,
                    "filename": source_item.filename,
                    "node_results": {},
                    "error_message": message,
                },
                ensure_ascii=False,
            )
        )
        db.add(
            WorkflowRunItem(
                run_id=new_run.id,
                document_id=source_item.document_id,
                filename=source_item.filename,
                upload_index=source_item.upload_index,
                status=status,
                error_message=message,
                client_file_id=source_item.client_file_id,
                upload_duration_ms=source_item.upload_duration_ms,
                execution_generation=new_run.execution_generation,
                result_json=result_json,
                completed_at=None if ready else now,
            )
        )
        if ready:
            queued_count += 1
    return new_run, queued_count


def _ensure_workflow_queue_group(run: WorkflowRun, db: Session) -> str:
    group_id = run.workflow_run_group_id or run.id
    run.workflow_run_group_id = group_id
    if run.queue_order is None:
        run.queue_order = 1 if run.id == group_id else _next_workflow_queue_order(db, group_id)
    db.flush()
    return group_id


def _next_workflow_queue_order(db: Session, group_id: str) -> int:
    orders = [
        order
        for (order,) in db.query(WorkflowRun.queue_order).filter(WorkflowRun.workflow_run_group_id == group_id).all()
        if order is not None
    ]
    return max(orders, default=0) + 1


def _validate_workflow_enqueue_source(run: WorkflowRun) -> None:
    if run.status not in WORKFLOW_ENQUEUE_BLOCKED_STATUSES:
        return
    if run.status == "waiting":
        detail = "Waiting workflow runs cannot be enqueued again"
    elif run.status == "canceled":
        detail = "Canceled workflow runs cannot be enqueued"
    else:
        detail = "Failed workflow runs cannot be enqueued"
    raise HTTPException(status_code=409, detail=detail)


def _validate_waiting_workflow_run_can_start(run: WorkflowRun, db: Session) -> None:
    group_id = run.workflow_run_group_id or run.id
    first_waiting = (
        db.query(WorkflowRun)
        .filter(WorkflowRun.workflow_run_group_id == group_id, WorkflowRun.status == "waiting")
        .order_by(WorkflowRun.queue_order.asc(), WorkflowRun.created_at.asc(), WorkflowRun.id.asc())
        .first()
    )
    if not first_waiting or first_waiting.id != run.id:
        raise HTTPException(status_code=409, detail="Only the first waiting workflow run can be started")

    run_position = _workflow_queue_position(run)
    active_predecessors = [
        candidate
        for candidate in db.query(WorkflowRun).filter(WorkflowRun.workflow_run_group_id == group_id).all()
        if candidate.id != run.id
        and candidate.status not in WORKFLOW_RUN_TERMINAL_STATUSES
        and _workflow_queue_position(candidate) < run_position
    ]
    if active_predecessors:
        raise HTTPException(status_code=409, detail="Previous workflow runs in this queue are still active")


def _workflow_queue_position(run: WorkflowRun) -> tuple[int, datetime, str]:
    return (run.queue_order if run.queue_order is not None else 0, run.created_at or datetime.min, run.id)


def _cancel_waiting_workflow_run(run: WorkflowRun, db: Session) -> None:
    now = datetime.utcnow()
    message = "Removed from workflow run queue"
    canceled_count = WorkflowRunLifecycle(run).cancel_waiting(now, message=message)
    log_audit_event(
        db,
        entity_type="workflow_run",
        entity_id=run.id,
        action="queue_canceled",
        message=f"Canceled waiting workflow run with {canceled_count} queued item(s)",
        metadata={"workflow_id": run.workflow_id, "queue_group_id": run.workflow_run_group_id, "queue_order": run.queue_order},
    )


def _stop_workflow_run_without_deleting_documents(run: WorkflowRun, message: str, db: Session) -> None:
    now = datetime.utcnow()
    _accumulate_workflow_run_inference_duration(run, now)
    stopped_count = WorkflowRunLifecycle(run).stop_without_deleting_documents(
        now,
        message=message,
        cancel_active_job=lambda item: _cancel_workflow_item_active_jobs(item, db, now, message),
    )
    log_audit_event(
        db,
        entity_type="workflow_run",
        entity_id=run.id,
        action="stopped",
        message=f"Stopped workflow run without deleting uploaded documents; {stopped_count} active item(s) canceled",
        metadata={
            "workflow_id": run.workflow_id,
            "queue_group_id": run.workflow_run_group_id,
            "queue_order": run.queue_order,
            "stopped_count": stopped_count,
        },
    )


def _cancel_workflow_item_active_jobs(item: WorkflowRunItem, db: Session, now: datetime, message: str) -> None:
    try:
        payload = json.loads(item.result_json or "{}")
    except json.JSONDecodeError:
        return
    node_results = payload.get("node_results")
    if not isinstance(node_results, dict):
        return
    for node_result in node_results.values():
        if not isinstance(node_result, dict):
            continue
        job_id = node_result.get("job_id")
        kind = node_result.get("kind")
        if not isinstance(job_id, str):
            continue
        if kind == "classifier":
            job = db.get(ClassificationJob, job_id)
            entity_type = "classification_job"
        elif kind == "kie":
            job = db.get(ExtractionJob, job_id)
            entity_type = "extraction_job"
        elif kind == "required-checker":
            job = db.get(RequiredFieldCheckJob, job_id)
            entity_type = "required_field_check_job"
        else:
            continue
        if not job or job.status in {"completed", "needs_review", "failed", "canceled"}:
            continue
        ModuleJobLifecycle(job).cancel(message, now)
        log_audit_event(
            db,
            entity_type=entity_type,
            entity_id=job.id,
            action="canceled",
            message=message,
            metadata={"workflow_run_id": item.run_id, "workflow_run_item_id": item.id},
        )


def _unshared_workflow_document_ids(run: WorkflowRun, document_ids: list[str], db: Session) -> list[str]:
    unique_document_ids = sorted({document_id for document_id in document_ids if document_id})
    if not unique_document_ids:
        return []
    shared_document_ids = {
        document_id
        for (document_id,) in db.query(WorkflowRunItem.document_id)
        .filter(WorkflowRunItem.run_id != run.id, WorkflowRunItem.document_id.in_(unique_document_ids))
        .all()
    }
    return [document_id for document_id in unique_document_ids if document_id not in shared_document_ids]


def _workflow_upload_duration_ms(run: WorkflowRun) -> int | None:
    durations = [item.upload_duration_ms for item in run.items if item.upload_duration_ms is not None]
    if not durations:
        return None
    return sum(durations)


def _duration_ms(started_at: datetime, ended_at: datetime | None = None) -> int:
    ended = ended_at or datetime.utcnow()
    return max(0, int((ended - started_at).total_seconds() * 1000))


def _accumulate_workflow_run_inference_duration(run: WorkflowRun, ended_at: datetime) -> None:
    if not run.inference_started_at:
        return
    run.inference_duration_ms = (run.inference_duration_ms or 0) + _duration_ms(run.inference_started_at, ended_at)
    run.inference_started_at = None


def _prepare_job_batch_resume(items: list[Any]) -> None:
    now = datetime.utcnow()
    for item in items:
        job = item.job
        document = item.document
        if not job:
            continue
        if job.status == "queued":
            continue
        if job.status not in {"running", "preprocessing"}:
            continue
        if document and document.status == "ready" and document.pages:
            job.status = "queued"
            job.error_message = None
            job.started_at = None
            job.completed_at = None
        else:
            job.status = "failed"
            job.error_message = document.error_message if document else "Document preprocessing was interrupted"
            job.completed_at = now


def _discard_batch_items(batch: Any, db: Session) -> int:
    items = list(batch.items)
    document_ids = [item.document_id for item in items]
    for item in items:
        db.delete(item)
    _delete_document_payloads(document_ids, db)
    return len(items)


def _delete_document_payloads(document_ids: list[str], db: Session) -> None:
    unique_document_ids = sorted({document_id for document_id in document_ids if document_id})
    if not unique_document_ids:
        return
    _delete_jobs_for_documents(unique_document_ids, db)
    documents = db.query(Document).filter(Document.id.in_(unique_document_ids)).all()
    for document in documents:
        _delete_document_storage(document)
        for page in list(document.pages):
            db.delete(page)
        db.delete(document)


def _delete_jobs_for_documents(document_ids: list[str], db: Session) -> None:
    for model in (ExtractionJob, ClassificationJob, RequiredFieldCheckJob):
        jobs = db.query(model).filter(model.document_id.in_(document_ids)).all()
        for job in jobs:
            if job.result:
                db.delete(job.result)
            db.delete(job)


def _delete_document_storage(document: Document) -> None:
    refs = [document.storage_path, *(page.image_path for page in document.pages)]
    local_deleted = False
    if document.storage_path and not is_s3_ref(document.storage_path):
        storage_root = get_settings().resolved_storage_dir.resolve()
        document_dir = Path(document.storage_path).resolve().parent
        if document_dir != storage_root and storage_root in document_dir.parents:
            delete_storage_ref(document_dir)
            local_deleted = True
    if local_deleted:
        return
    for ref in refs:
        if ref:
            delete_storage_ref(ref)


def _sorted_batch_items(items) -> list[BatchItem]:
    return sorted(items, key=_batch_item_sort_key)


def _batch_item_sort_key(item: BatchItem) -> tuple[str, str]:
    upload_index = getattr(item, "upload_index", None)
    if upload_index is not None:
        return (f"{upload_index:012d}", item.id)
    return (f"z:{item.filename.casefold()}", item.id)


def _sorted_module_items(items) -> list[Any]:
    return sorted(items, key=_module_item_sort_key)


def _sorted_workflow_items(items) -> list[WorkflowRunItem]:
    return sorted(items, key=_module_item_sort_key)


def _module_item_sort_key(item: Any) -> tuple[str, str]:
    upload_index = getattr(item, "upload_index", None)
    if upload_index is not None:
        return (f"{upload_index:012d}", item.id)
    return (f"z:{item.filename.casefold()}", item.id)


def _upload_file_sort_key(file: UploadFile) -> tuple[str, str]:
    filename = file.filename or ""
    return (filename.casefold(), filename)


async def _read_batch_upload_form(request: Request) -> tuple[FormData, list[UploadFile]]:
    settings = get_settings()
    try:
        form = await request.form(
            max_files=_multipart_max_files(settings.upload_max_batch_files),
            max_fields=max(32, settings.upload_chunk_files * 3 + 16),
        )
    except StarletteHTTPException as exc:
        detail = str(exc.detail)
        if "Too many files" in detail or "Maximum number of files" in detail:
            raise HTTPException(status_code=413, detail=_batch_file_limit_message()) from exc
        raise HTTPException(status_code=400, detail=f"Invalid multipart upload: {detail}") from exc
    except OSError as exc:
        if exc.errno == errno.EMFILE:
            raise HTTPException(
                status_code=413,
                detail="Too many files were uploaded in a single request. Upload large batches in smaller chunks.",
            ) from exc
        raise

    files = [item for item in form.getlist("files") if _is_upload_file(item)]
    if not files:
        raise HTTPException(status_code=400, detail="At least one file is required")
    _validate_upload_file_count(files)
    return form, files


def _multipart_max_files(upload_limit: int) -> int:
    if upload_limit > 0:
        return max(1000, upload_limit)
    return 10000


def _is_upload_file(value: Any) -> bool:
    return hasattr(value, "filename") and hasattr(value, "file")


def _required_form_value(form: FormData, key: str) -> str:
    value = form.get(key)
    if not isinstance(value, str) or not value.strip():
        raise HTTPException(status_code=400, detail=f"{key} is required")
    return value.strip()


def _validate_upload_file_count(files: list[UploadFile]) -> None:
    limit = get_settings().upload_max_batch_files
    if limit > 0 and len(files) > limit:
        raise HTTPException(status_code=413, detail=_batch_file_limit_message())


def _validate_declared_batch_file_count(file_count: int) -> None:
    limit = get_settings().upload_max_batch_files
    if limit > 0 and file_count > limit:
        raise HTTPException(status_code=413, detail=_batch_file_limit_message())


def _batch_file_limit_message() -> str:
    limit = get_settings().upload_max_batch_files
    if limit > 0:
        return f"Batch file count exceeds the configured limit of {limit}"
    return "Batch file count exceeds the multipart parser limit"


def _ordered_upload_entries(form: FormData, files: list[UploadFile]) -> list[tuple[UploadFile, str | None, int | None]]:
    raw_client_ids = [value if isinstance(value, str) and value.strip() else None for value in form.getlist("client_file_ids")]
    if len(raw_client_ids) != len(files):
        raw_client_ids = [None] * len(files)
    raw_upload_indexes: list[int | None] = []
    for value in form.getlist("upload_indexes"):
        if isinstance(value, str) and value.strip():
            try:
                parsed = int(value)
                raw_upload_indexes.append(parsed if parsed >= 0 else None)
            except ValueError:
                raw_upload_indexes.append(None)
        else:
            raw_upload_indexes.append(None)
    if len(raw_upload_indexes) != len(files):
        raw_upload_indexes = [None] * len(files)
    return sorted(
        zip(files, raw_client_ids, raw_upload_indexes, strict=False),
        key=lambda entry: (entry[2] is None, entry[2] if entry[2] is not None else 0, *_upload_file_sort_key(entry[0])),
    )


def _existing_client_file_ids(db: Session, item_model: Any, owner_field: str, owner_id: str, client_file_ids: list[str | None]) -> set[str]:
    ids = [client_file_id for client_file_id in client_file_ids if client_file_id]
    if not ids:
        return set()
    rows = (
        db.query(item_model.client_file_id)
        .filter(getattr(item_model, owner_field) == owner_id, item_model.client_file_id.in_(ids))
        .all()
    )
    return {row[0] for row in rows if row[0]}


def _existing_upload_indexes(db: Session, item_model: Any, owner_field: str, owner_id: str, upload_indexes: list[int | None]) -> set[int]:
    if not hasattr(item_model, "upload_index"):
        return set()
    indexes = [upload_index for upload_index in upload_indexes if upload_index is not None]
    if not indexes:
        return set()
    rows = (
        db.query(item_model.upload_index)
        .filter(getattr(item_model, owner_field) == owner_id, item_model.upload_index.in_(indexes))
        .all()
    )
    return {row[0] for row in rows if row[0] is not None}


def _filter_new_upload_entries(
    db: Session,
    item_model: Any,
    owner_field: str,
    owner_id: str,
    entries: list[tuple[UploadFile, str | None, int | None]],
) -> list[tuple[UploadFile, str | None, int | None]]:
    existing_ids = _existing_client_file_ids(db, item_model, owner_field, owner_id, [client_id for _, client_id, _ in entries])
    existing_indexes = _existing_upload_indexes(db, item_model, owner_field, owner_id, [upload_index for _, _, upload_index in entries])
    accepted_ids = set(existing_ids)
    accepted_indexes = set(existing_indexes)
    new_entries: list[tuple[UploadFile, str | None, int | None]] = []
    for file, client_file_id, upload_index in entries:
        if client_file_id and client_file_id in accepted_ids:
            continue
        if upload_index is not None and upload_index in accepted_indexes:
            continue
        new_entries.append((file, client_file_id, upload_index))
        if client_file_id:
            accepted_ids.add(client_file_id)
        if upload_index is not None:
            accepted_indexes.add(upload_index)
    return new_entries


def _ensure_upload_append_capacity(
    db: Session,
    item_model: Any,
    owner_field: str,
    owner_id: str,
    total_count: int,
    incoming_count: int,
) -> None:
    current_count = db.query(item_model).filter(getattr(item_model, owner_field) == owner_id).count()
    if current_count + incoming_count > total_count:
        raise HTTPException(status_code=413, detail="Batch upload exceeds declared file count")


def _upload_failure_message(exc: Exception) -> str:
    if isinstance(exc, HTTPException):
        return str(exc.detail)
    if isinstance(exc, DocumentProcessingError):
        return str(exc)
    return "Failed to process uploaded document"


def _initial_workflow_item_payload(document: Document) -> str:
    return json.dumps({"document_id": document.id, "filename": document.filename, "node_results": {}}, ensure_ascii=False)


async def _append_workflow_upload_items(run: WorkflowRun, form: FormData, files: list[UploadFile], db: Session) -> None:
    entries = _filter_new_upload_entries(db, WorkflowRunItem, "run_id", run.id, _ordered_upload_entries(form, files))
    incoming_count = len(entries)
    _ensure_upload_append_capacity(db, WorkflowRunItem, "run_id", run.id, run.total_count, incoming_count)

    for file, client_file_id, upload_index in entries:
        upload_started_at = datetime.utcnow()
        try:
            db.refresh(run)
            if not _owner_accepts_uploads(run):
                continue
            document, original_path = _create_uploaded_document(file, db, workspace_id=run.workspace_id)
            item = WorkflowRunItem(
                run_id=run.id,
                document_id=document.id,
                filename=document.filename,
                upload_index=upload_index,
                status="preprocessing",
                client_file_id=client_file_id,
                result_json=_initial_workflow_item_payload(document),
            )
            db.add(item)
            log_audit_event(
                db,
                entity_type="workflow_run_item",
                entity_id=run.id,
                action="preprocessing",
                message=f"Preprocessing workflow document {document.filename}",
                metadata={"document_id": document.id},
            )
            db.commit()

            ok = _preprocess_document_pages(document, original_path, db, raise_errors=False)
            db.refresh(run)
            if not _owner_accepts_uploads(run):
                db.rollback()
                continue
            if ok:
                item.status = "queued"
                item.error_message = None
                item.upload_duration_ms = _duration_ms(upload_started_at)
                log_audit_event(
                    db,
                    entity_type="workflow_run_item",
                    entity_id=run.id,
                    action="queued",
                    message=f"Queued workflow document {document.filename}",
                    metadata={"document_id": document.id},
                )
            else:
                item.status = "failed"
                item.error_message = document.error_message
                item.completed_at = datetime.utcnow()
                item.upload_duration_ms = _duration_ms(upload_started_at, item.completed_at)
                item.result_json = json.dumps(
                    {
                        "document_id": document.id,
                        "filename": document.filename,
                        "node_results": {},
                        "error_message": document.error_message,
                    },
                    ensure_ascii=False,
                )
            db.commit()
        except Exception as exc:
            db.rollback()
            _record_failed_workflow_upload_item(run, file, client_file_id, upload_index, _duration_ms(upload_started_at), _upload_failure_message(exc), db)
        finally:
            await file.close()


def _record_failed_workflow_upload_item(
    run: WorkflowRun,
    file: UploadFile,
    client_file_id: str | None,
    upload_index: int | None,
    upload_duration_ms: int | None,
    message: str,
    db: Session,
) -> None:
    document = _create_failed_upload_document(file, message, db, workspace_id=run.workspace_id)
    item = WorkflowRunItem(
        run_id=run.id,
        document_id=document.id,
        filename=document.filename,
        upload_index=upload_index,
        status="failed",
        error_message=message,
        client_file_id=client_file_id,
        upload_duration_ms=upload_duration_ms,
        result_json=json.dumps(
            {"document_id": document.id, "filename": document.filename, "node_results": {}, "error_message": message},
            ensure_ascii=False,
        ),
        completed_at=datetime.utcnow(),
    )
    db.add(item)
    log_audit_event(
        db,
        entity_type="workflow_run_item",
        entity_id=run.id,
        action="failed",
        message=message,
        metadata={"document_id": document.id},
    )
    db.commit()


def _seal_missing_workflow_upload_items(run: WorkflowRun, db: Session) -> int:
    existing_count = db.query(WorkflowRunItem).filter(WorkflowRunItem.run_id == run.id).count()
    missing_count = max(0, run.total_count - existing_count)
    if missing_count == 0:
        return 0
    existing_indexes = {
        row[0]
        for row in db.query(WorkflowRunItem.upload_index)
        .filter(WorkflowRunItem.run_id == run.id, WorkflowRunItem.upload_index.isnot(None))
        .all()
    }
    candidate_indexes = [index for index in range(run.total_count) if index not in existing_indexes]
    if len(candidate_indexes) < missing_count:
        candidate_indexes.extend([None] * (missing_count - len(candidate_indexes)))
    now = datetime.utcnow()
    message = "Upload was not received before execution was restarted"
    for ordinal, upload_index in enumerate(candidate_indexes[:missing_count], start=1):
        suffix = upload_index + 1 if upload_index is not None else existing_count + ordinal
        filename = f"missing_upload_{suffix:05d}"
        document = _create_failed_upload_document_record(
            filename,
            "application/octet-stream",
            message,
            db,
            workspace_id=run.workspace_id,
        )
        db.add(
            WorkflowRunItem(
                run_id=run.id,
                document_id=document.id,
                filename=document.filename,
                upload_index=upload_index,
                status="failed",
                error_message=message,
                client_file_id=f"missing:{run.id}:{upload_index}" if upload_index is not None else None,
                result_json=json.dumps(
                    {"document_id": document.id, "filename": document.filename, "node_results": {}, "error_message": message},
                    ensure_ascii=False,
                ),
                completed_at=now,
            )
        )
    log_audit_event(
        db,
        entity_type="workflow_run",
        entity_id=run.id,
        action="upload_sealed",
        message=f"Marked {missing_count} missing upload slot(s) as failed",
        metadata={"workflow_id": run.workflow_id, "missing_count": missing_count},
    )
    db.flush()
    return missing_count


async def _append_extraction_batch_items(batch: Batch, form: FormData, files: list[UploadFile], db: Session) -> None:
    entries = _filter_new_upload_entries(db, BatchItem, "batch_id", batch.id, _ordered_upload_entries(form, files))
    incoming_count = len(entries)
    _ensure_upload_append_capacity(db, BatchItem, "batch_id", batch.id, batch.total_count, incoming_count)

    for file, client_file_id, upload_index in entries:
        try:
            db.refresh(batch)
            if not _owner_accepts_uploads(batch):
                continue
            document, original_path = _create_uploaded_document(file, db, workspace_id=batch.workspace_id)
            job = ExtractionJob(
                workspace_id=batch.workspace_id,
                document_id=document.id,
                schema_id=batch.schema_id,
                schema_version=batch.schema_version,
                status="preprocessing",
            )
            db.add(job)
            db.flush()
            db.add(
                BatchItem(
                    batch_id=batch.id,
                    document_id=document.id,
                    job_id=job.id,
                    filename=document.filename,
                    upload_index=upload_index,
                    client_file_id=client_file_id,
                )
            )
            log_audit_event(
                db,
                entity_type="document",
                entity_id=document.id,
                action="uploaded",
                message=f"Batch uploaded {document.filename}",
                metadata={"batch_id": batch.id, "filename": document.filename},
            )
            db.commit()

            ok = _preprocess_document_pages(document, original_path, db, raise_errors=False)
            db.refresh(batch)
            if not _owner_accepts_uploads(batch):
                db.rollback()
                continue
            if ok:
                job.status = "queued"
                job.error_message = None
                log_audit_event(
                    db,
                    entity_type="extraction_job",
                    entity_id=job.id,
                    action="created",
                    message="Batch extraction job created",
                    metadata={"batch_id": batch.id, "document_id": document.id, "schema_id": batch.schema_id},
                )
            else:
                job.status = "failed"
                job.error_message = document.error_message
                job.completed_at = datetime.utcnow()
            db.commit()
        except Exception as exc:
            db.rollback()
            _record_failed_extraction_batch_item(batch, file, client_file_id, upload_index, _upload_failure_message(exc), db)
        finally:
            await file.close()


def _record_failed_extraction_batch_item(
    batch: Batch,
    file: UploadFile,
    client_file_id: str | None,
    upload_index: int | None,
    message: str,
    db: Session,
) -> None:
    document = _create_failed_upload_document(file, message, db, workspace_id=batch.workspace_id)
    job = ExtractionJob(
        workspace_id=batch.workspace_id,
        document_id=document.id,
        schema_id=batch.schema_id,
        schema_version=batch.schema_version,
        status="failed",
        error_message=message,
        completed_at=datetime.utcnow(),
    )
    db.add(job)
    db.flush()
    db.add(
        BatchItem(
            batch_id=batch.id,
            document_id=document.id,
            job_id=job.id,
            filename=document.filename,
            upload_index=upload_index,
            client_file_id=client_file_id,
        )
    )
    log_audit_event(db, entity_type="extraction_job", entity_id=job.id, action="failed", message=message, metadata={"batch_id": batch.id})
    db.commit()


async def _append_classification_batch_items(batch: ClassificationBatch, form: FormData, files: list[UploadFile], db: Session) -> None:
    entries = _filter_new_upload_entries(db, ClassificationBatchItem, "batch_id", batch.id, _ordered_upload_entries(form, files))
    incoming_count = len(entries)
    _ensure_upload_append_capacity(db, ClassificationBatchItem, "batch_id", batch.id, batch.total_count, incoming_count)

    for file, client_file_id, upload_index in entries:
        try:
            db.refresh(batch)
            if not _owner_accepts_uploads(batch):
                continue
            document, original_path = _create_uploaded_document(file, db, workspace_id=batch.workspace_id)
            job = ClassificationJob(
                workspace_id=batch.workspace_id,
                document_id=document.id,
                classifier_id=batch.classifier_id,
                status="preprocessing",
            )
            db.add(job)
            db.flush()
            db.add(
                ClassificationBatchItem(
                    batch_id=batch.id,
                    document_id=document.id,
                    job_id=job.id,
                    filename=document.filename,
                    upload_index=upload_index,
                    client_file_id=client_file_id,
                )
            )
            db.commit()

            ok = _preprocess_document_pages(document, original_path, db, raise_errors=False)
            db.refresh(batch)
            if not _owner_accepts_uploads(batch):
                db.rollback()
                continue
            if ok:
                job.status = "queued"
                job.error_message = None
                log_audit_event(
                    db,
                    entity_type="classification_job",
                    entity_id=job.id,
                    action="queued",
                    message="Queued classification batch job",
                    metadata={"batch_id": batch.id, "document_id": document.id, "classifier_id": batch.classifier_id},
                )
            else:
                job.status = "failed"
                job.error_message = document.error_message
                job.completed_at = datetime.utcnow()
            db.commit()
        except Exception as exc:
            db.rollback()
            _record_failed_classification_batch_item(batch, file, client_file_id, upload_index, _upload_failure_message(exc), db)
        finally:
            await file.close()


def _record_failed_classification_batch_item(
    batch: ClassificationBatch,
    file: UploadFile,
    client_file_id: str | None,
    upload_index: int | None,
    message: str,
    db: Session,
) -> None:
    document = _create_failed_upload_document(file, message, db, workspace_id=batch.workspace_id)
    job = ClassificationJob(
        workspace_id=batch.workspace_id,
        document_id=document.id,
        classifier_id=batch.classifier_id,
        status="failed",
        error_message=message,
        completed_at=datetime.utcnow(),
    )
    db.add(job)
    db.flush()
    db.add(
        ClassificationBatchItem(
            batch_id=batch.id,
            document_id=document.id,
            job_id=job.id,
            filename=document.filename,
            upload_index=upload_index,
            client_file_id=client_file_id,
        )
    )
    log_audit_event(db, entity_type="classification_job", entity_id=job.id, action="failed", message=message, metadata={"batch_id": batch.id})
    db.commit()


async def _append_required_field_batch_items(batch: RequiredFieldCheckBatch, form: FormData, files: list[UploadFile], db: Session) -> None:
    entries = _filter_new_upload_entries(db, RequiredFieldCheckBatchItem, "batch_id", batch.id, _ordered_upload_entries(form, files))
    incoming_count = len(entries)
    _ensure_upload_append_capacity(db, RequiredFieldCheckBatchItem, "batch_id", batch.id, batch.total_count, incoming_count)

    for file, client_file_id, upload_index in entries:
        try:
            db.refresh(batch)
            if not _owner_accepts_uploads(batch):
                continue
            document, original_path = _create_uploaded_document(file, db, workspace_id=batch.workspace_id)
            job = RequiredFieldCheckJob(
                workspace_id=batch.workspace_id,
                document_id=document.id,
                checklist_id=batch.checklist_id,
                status="preprocessing",
            )
            db.add(job)
            db.flush()
            db.add(
                RequiredFieldCheckBatchItem(
                    batch_id=batch.id,
                    document_id=document.id,
                    job_id=job.id,
                    filename=document.filename,
                    upload_index=upload_index,
                    client_file_id=client_file_id,
                )
            )
            db.commit()

            ok = _preprocess_document_pages(document, original_path, db, raise_errors=False)
            db.refresh(batch)
            if not _owner_accepts_uploads(batch):
                db.rollback()
                continue
            if ok:
                job.status = "queued"
                job.error_message = None
                log_audit_event(
                    db,
                    entity_type="required_field_check_job",
                    entity_id=job.id,
                    action="queued",
                    message="Queued required field check batch job",
                    metadata={"batch_id": batch.id, "document_id": document.id, "checklist_id": batch.checklist_id},
                )
            else:
                job.status = "failed"
                job.error_message = document.error_message
                job.completed_at = datetime.utcnow()
            db.commit()
        except Exception as exc:
            db.rollback()
            _record_failed_required_field_batch_item(batch, file, client_file_id, upload_index, _upload_failure_message(exc), db)
        finally:
            await file.close()


def _record_failed_required_field_batch_item(
    batch: RequiredFieldCheckBatch,
    file: UploadFile,
    client_file_id: str | None,
    upload_index: int | None,
    message: str,
    db: Session,
) -> None:
    document = _create_failed_upload_document(file, message, db, workspace_id=batch.workspace_id)
    job = RequiredFieldCheckJob(
        workspace_id=batch.workspace_id,
        document_id=document.id,
        checklist_id=batch.checklist_id,
        status="failed",
        error_message=message,
        completed_at=datetime.utcnow(),
    )
    db.add(job)
    db.flush()
    db.add(
        RequiredFieldCheckBatchItem(
            batch_id=batch.id,
            document_id=document.id,
            job_id=job.id,
            filename=document.filename,
            upload_index=upload_index,
            client_file_id=client_file_id,
        )
    )
    log_audit_event(db, entity_type="required_field_check_job", entity_id=job.id, action="failed", message=message, metadata={"batch_id": batch.id})
    db.commit()


def _queued_extraction_job_ids(batch: Batch) -> list[str]:
    return [item.job_id for item in batch.items if item.job and item.job.status == "queued"]


def _owner_accepts_uploads(owner: Any) -> bool:
    return getattr(owner, "status", None) in {"uploading", "queued"}


def _queued_classification_job_ids(batch: ClassificationBatch) -> list[str]:
    return [item.job_id for item in batch.items if item.job and item.job.status == "queued"]


def _queued_required_field_job_ids(batch: RequiredFieldCheckBatch) -> list[str]:
    return [item.job_id for item in batch.items if item.job and item.job.status == "queued"]


def _validate_owner_can_start(owner: Any, items: list[Any]) -> None:
    _validate_owner_upload_complete(owner, items)
    preprocessing_count = sum(1 for item in items if _upload_item_status(item) == "preprocessing")
    if preprocessing_count:
        raise HTTPException(
            status_code=422,
            detail={"message": f"{preprocessing_count} file(s) are still preprocessing", "preprocessing_count": preprocessing_count},
        )


def _validate_owner_upload_complete(owner: Any, items: list[Any]) -> None:
    uploaded_count = len(items)
    if uploaded_count != owner.total_count:
        missing_count = max(0, owner.total_count - uploaded_count)
        raise HTTPException(
            status_code=422,
            detail={
                "message": "Upload is incomplete. Re-select the original files to continue uploading before starting execution.",
                "uploaded_count": uploaded_count,
                "total_count": owner.total_count,
                "missing_count": missing_count,
            },
        )


def _upload_item_status(item: Any) -> str:
    status = getattr(item, "status", None)
    if isinstance(status, str):
        return status
    job = getattr(item, "job", None)
    if job and isinstance(job.status, str):
        return job.status
    return "unknown"


def _extract_kie_cell_value(value: Any) -> Any:
    return value.get("value") if isinstance(value, dict) else value


def _add_kie_review_export_columns(
    row: dict[str, Any],
    column_prefix: str,
    value: Any,
    original_value: Any = None,
    reviewed_fields: set[str] | None = None,
    field_name: str | None = None,
) -> None:
    value_dict = value if isinstance(value, dict) else {}
    ai_review = value_dict.get("ai_review") if isinstance(value_dict.get("ai_review"), dict) else {}
    current = _extract_kie_cell_value(value)
    original = _extract_kie_cell_value(original_value) if original_value is not None else current
    row[column_prefix] = current
    row[f"{column_prefix}_original"] = original
    row[f"{column_prefix}_changed"] = current != original
    row[f"{column_prefix}_reviewed"] = field_name in reviewed_fields if reviewed_fields is not None and field_name else False
    row[f"{column_prefix}_warnings"] = value_dict.get("warnings", [])
    row[f"{column_prefix}_ai_review_enabled"] = bool(ai_review.get("enabled"))
    row[f"{column_prefix}_ai_review_status"] = ai_review.get("judgement_status")
    row[f"{column_prefix}_ai_corrected"] = bool(ai_review.get("corrected"))
    row[f"{column_prefix}_ai_review_reason"] = ai_review.get("judgement_reason")
    row[f"{column_prefix}_ai_review_confidence"] = ai_review.get("judgement_confidence")
    row[f"{column_prefix}_ai_initial_value"] = ai_review.get("initial_value")
    row[f"{column_prefix}_ai_initial_evidence"] = ai_review.get("initial_evidence")
    row[f"{column_prefix}_ai_correction_reason"] = ai_review.get("correction_reason")


def _kie_export_columns(field_name: str) -> list[str]:
    return [
        field_name,
        f"{field_name}_original",
        f"{field_name}_changed",
        f"{field_name}_reviewed",
        f"{field_name}_warnings",
        f"{field_name}_ai_review_enabled",
        f"{field_name}_ai_review_status",
        f"{field_name}_ai_corrected",
        f"{field_name}_ai_review_reason",
        f"{field_name}_ai_review_confidence",
        f"{field_name}_ai_initial_value",
        f"{field_name}_ai_initial_evidence",
        f"{field_name}_ai_correction_reason",
    ]


def _batch_export_row(item: BatchItem, field_names: list[str]) -> dict[str, Any]:
    job = item.job
    row: dict[str, Any] = {
        "filename": item.filename,
        "document_id": item.document_id,
        "job_id": item.job_id,
        "status": job.status if job else "unknown",
        "error_message": job.error_message if job else None,
        "warnings": [],
    }
    for field_name in field_names:
        row[field_name] = None
    if not job or not job.result:
        return row

    output = json.loads(job.result.corrected_output) if job.result.corrected_output else json.loads(job.result.validated_output)
    original_output = json.loads(job.result.validated_output)
    reviewed_fields = set(json.loads(job.result.reviewed_fields or "[]"))
    values = output.get("values", {})
    original_values = original_output.get("values", {}) if isinstance(original_output.get("values"), dict) else {}
    warnings: list[str] = []
    for field_name in field_names:
        value = values.get(field_name)
        if isinstance(value, dict):
            _add_kie_review_export_columns(row, field_name, value, original_values.get(field_name), reviewed_fields, field_name)
            warnings.extend(str(warning) for warning in value.get("warnings", []))
        else:
            _add_kie_review_export_columns(row, field_name, value, original_values.get(field_name), reviewed_fields, field_name)
    row["warnings"] = warnings
    return row


def _classification_batch_export_row(item: ClassificationBatchItem) -> dict[str, Any]:
    job = item.job
    row: dict[str, Any] = {
        "filename": item.filename,
        "document_id": item.document_id,
        "job_id": item.job_id,
        "status": job.status if job else "unknown",
        "error_message": job.error_message if job else None,
        "classification_status": None,
        "class_name": None,
        "confidence": None,
        "reason": None,
        "evidence": [],
    }
    if not job or not job.result:
        return row
    output = json.loads(job.result.corrected_output) if job.result.corrected_output else json.loads(job.result.validated_output)
    row["classification_status"] = output.get("status")
    row["class_name"] = output.get("class_name")
    row["confidence"] = output.get("confidence")
    row["reason"] = output.get("reason")
    row["evidence"] = output.get("evidence") if isinstance(output.get("evidence"), list) else []
    return row


def _required_field_batch_export_row(item: RequiredFieldCheckBatchItem, item_names: list[str]) -> dict[str, Any]:
    job = item.job
    row: dict[str, Any] = {
        "filename": item.filename,
        "document_id": item.document_id,
        "job_id": item.job_id,
        "status": job.status if job else "unknown",
        "error_message": job.error_message if job else None,
        "overall_status": None,
    }
    for item_name in item_names:
        row[f"{item_name}_status"] = None
        row[f"{item_name}_evidence"] = None
    if not job or not job.result:
        return row
    output = json.loads(job.result.corrected_output) if job.result.corrected_output else json.loads(job.result.validated_output)
    row["overall_status"] = output.get("overall_status")
    output_items = output.get("items") if isinstance(output.get("items"), list) else []
    by_name = {entry.get("item_name"): entry for entry in output_items if isinstance(entry, dict)}
    for item_name in item_names:
        entry = by_name.get(item_name, {})
        row[f"{item_name}_status"] = entry.get("status")
        row[f"{item_name}_evidence"] = entry.get("evidence")
    return row


def _csv_cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, list):
        return ";".join(str(item) for item in value)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False)
    return value


def _safe_filename_part(value: Any, fallback: str = "export") -> str:
    text_value = str(value or "").strip() or fallback
    text_value = re.sub(r'[\\/:*?"<>|\r\n\t]+', "_", text_value)
    text_value = re.sub(r"\s+", "_", text_value)
    text_value = text_value.strip("._ ")
    return (text_value or fallback)[:100]


def _export_filename(module: str, name: Any, identifier: str, extension: str) -> str:
    return "_".join(
        [
            _safe_filename_part(module, "module"),
            _safe_filename_part(name, "untitled"),
            _safe_filename_part(identifier, "job"),
        ]
    ) + f".{extension}"


def _download_headers(filename: str) -> dict[str, str]:
    ascii_filename = "".join(ch if 32 <= ord(ch) < 127 else "_" for ch in filename).replace("\\", "_").replace('"', "_")
    ascii_filename = ascii_filename or "export"
    return {"Content-Disposition": f"attachment; filename=\"{ascii_filename}\"; filename*=UTF-8''{quote(filename)}"}


def _export_job_read(job: ExportJob) -> ExportJobRead:
    return ExportJobRead(
        id=job.id,
        owner_type=job.owner_type,
        owner_id=job.owner_id,
        format=job.format,
        status=job.status,
        filename=job.filename,
        content_type=job.content_type,
        size_bytes=job.size_bytes,
        error_message=job.error_message,
        created_at=job.created_at,
        started_at=job.started_at,
        completed_at=job.completed_at,
    )


def _queue_export_job(
    db: Session,
    background_tasks: BackgroundTasks,
    *,
    owner_type: str,
    owner_id: str,
    format: str,
    retry_source_job_id: str | None = None,
    workspace_id: str | None = None,
) -> ExportJob:
    _validate_export_job_owner_exists(db, owner_type, owner_id, workspace_id=workspace_id)
    job = ExportJob(workspace_id=workspace_id, owner_type=owner_type, owner_id=owner_id, format=format, status="queued")
    db.add(job)
    db.flush()
    action = "retried" if retry_source_job_id else "created"
    message = (
        f"Retried {owner_type} export {format.upper()}"
        if retry_source_job_id
        else f"Queued {owner_type} export {format.upper()}"
    )
    metadata = {"owner_type": owner_type, "owner_id": owner_id, "format": format}
    if retry_source_job_id:
        metadata["retry_source_job_id"] = retry_source_job_id
    log_audit_event(
        db,
        entity_type="export_job",
        entity_id=job.id,
        action=action,
        message=message,
        metadata=metadata,
    )
    db.commit()
    db.refresh(job)
    background_tasks.add_task(_run_export_job, job.id)
    return job


def _validate_export_job_owner_exists(db: Session, owner_type: str, owner_id: str, *, workspace_id: str | None = None) -> None:
    if owner_type not in EXPORT_JOB_OWNER_TYPES:
        raise HTTPException(status_code=422, detail="Unsupported export owner type")
    owner_models: dict[str, tuple[Any, str]] = {
        "workflow_run": (WorkflowRun, "Workflow run not found"),
        "batch": (Batch, "Batch not found"),
        "classification_batch": (ClassificationBatch, "Classification batch not found"),
        "required_field_check_batch": (RequiredFieldCheckBatch, "Required field check batch not found"),
    }
    owner_model = owner_models.get(owner_type)
    if not owner_model:
        raise HTTPException(status_code=422, detail="Unsupported export owner type")
    model, not_found_detail = owner_model
    owner = db.get(model, owner_id)
    if not owner or (workspace_id is not None and getattr(owner, "workspace_id", None) != workspace_id):
        raise HTTPException(status_code=404, detail=not_found_detail)


def _start_export_job_worker() -> threading.Event:
    stop_event = threading.Event()
    _reset_interrupted_export_jobs()

    def _run() -> None:
        while not stop_event.is_set():
            job_id = _next_queued_export_job_id()
            if not job_id:
                stop_event.wait(EXPORT_JOB_WORKER_INTERVAL_SECONDS)
                continue
            _run_export_job(job_id)

    thread = threading.Thread(target=_run, name="export-job-worker", daemon=True)
    thread.start()
    return stop_event


def _reset_interrupted_export_jobs() -> None:
    db = SessionLocal()
    try:
        db.query(ExportJob).filter(ExportJob.status == "running").update(
            {
                ExportJob.status: "queued",
                ExportJob.started_at: None,
                ExportJob.error_message: "Recovered after server restart",
            },
            synchronize_session=False,
        )
        db.commit()
    finally:
        db.close()


def _next_queued_export_job_id() -> str | None:
    db = SessionLocal()
    try:
        row = (
            db.query(ExportJob.id)
            .filter(ExportJob.status == "queued")
            .order_by(ExportJob.created_at.asc(), ExportJob.id.asc())
            .first()
        )
        return row[0] if row else None
    finally:
        db.close()


def _run_export_job(job_id: str) -> None:
    db = SessionLocal()
    try:
        started_at = datetime.utcnow()
        claimed_count = (
            db.query(ExportJob)
            .filter(ExportJob.id == job_id, ExportJob.status == "queued")
            .update(
                {
                    ExportJob.status: "running",
                    ExportJob.started_at: started_at,
                    ExportJob.completed_at: None,
                    ExportJob.error_message: None,
                },
                synchronize_session=False,
            )
        )
        db.commit()
        if not claimed_count:
            return

        job = db.get(ExportJob, job_id)
        if not job:
            return
        artifact = _build_export_artifact(db, job.owner_type, job.owner_id, job.format)
        target = _export_job_artifact_path(job.id, artifact.filename)
        target.write_bytes(artifact.content)
        storage_ref = persist_artifact(target, f"exports/{job.id}/{artifact.filename}", artifact.content_type)

        job.storage_path = storage_ref
        job.filename = artifact.filename
        job.content_type = artifact.content_type
        job.size_bytes = len(artifact.content)
        job.status = "completed"
        job.completed_at = datetime.utcnow()
        log_audit_event(
            db,
            entity_type=job.owner_type,
            entity_id=job.owner_id,
            action="exported_async",
            message=f"Completed async export {artifact.filename}",
            metadata={"format": job.format, "export_job_id": job.id, "size_bytes": job.size_bytes},
        )
        db.commit()
    except Exception as exc:
        db.rollback()
        failed_job = db.get(ExportJob, job_id)
        if failed_job:
            failed_job.status = "failed"
            failed_job.error_message = _export_job_error_message(exc)
            failed_job.completed_at = datetime.utcnow()
            log_audit_event(
                db,
                entity_type="export_job",
                entity_id=failed_job.id,
                action="failed",
                message=failed_job.error_message,
                metadata={
                    "owner_type": failed_job.owner_type,
                    "owner_id": failed_job.owner_id,
                    "format": failed_job.format,
                },
            )
            db.commit()
    finally:
        db.close()


def _export_job_error_message(exc: Exception) -> str:
    detail = exc.detail if isinstance(exc, HTTPException) else str(exc)
    if isinstance(detail, dict | list):
        detail = json.dumps(detail, ensure_ascii=False, default=str)
    return str(detail or "Export failed")[:2000]


def _export_job_artifact_path(job_id: str, filename: str) -> Path:
    target_dir = get_settings().resolved_storage_dir / "exports" / job_id
    target_dir.mkdir(parents=True, exist_ok=True)
    return target_dir / filename


def _build_export_artifact(db: Session, owner_type: str, owner_id: str, format: str) -> ExportArtifact:
    if format not in {"json", "csv", "xlsx"}:
        raise HTTPException(status_code=422, detail="Unsupported export format")
    if owner_type == "workflow_run":
        run = db.get(WorkflowRun, owner_id)
        if not run:
            raise HTTPException(status_code=404, detail="Workflow run not found")
        return _workflow_run_export_artifact(run, format)
    if owner_type == "batch":
        batch = db.get(Batch, owner_id)
        if not batch:
            raise HTTPException(status_code=404, detail="Batch not found")
        return _batch_export_artifact(db, batch, format)
    if owner_type == "classification_batch":
        batch = db.get(ClassificationBatch, owner_id)
        if not batch:
            raise HTTPException(status_code=404, detail="Classification batch not found")
        return _classification_batch_export_artifact(db, batch, format)
    if owner_type == "required_field_check_batch":
        batch = db.get(RequiredFieldCheckBatch, owner_id)
        if not batch:
            raise HTTPException(status_code=404, detail="Required field check batch not found")
        return _required_field_batch_export_artifact(db, batch, format)
    raise HTTPException(status_code=422, detail="Unsupported export owner type")


def _workflow_run_export_artifact(run: WorkflowRun, format: str) -> ExportArtifact:
    payload = workflow_run_export_payload(run)
    workflow_name = payload.get("workflow_name") or (run.workflow.name if run.workflow else "workflow")
    filename = _export_filename("workflow", workflow_name, run.id, format)
    if format == "json":
        return ExportArtifact(_json_export_bytes(payload), filename, "application/json")
    if format == "xlsx":
        rows = payload.get("rows") if isinstance(payload.get("rows"), list) else []
        return ExportArtifact(_xlsx_bytes(rows, _table_fieldnames(rows)), filename, XLSX_MIME_TYPE)
    return ExportArtifact(_csv_export_bytes(workflow_run_export_csv(run)), filename, "text/csv; charset=utf-8")


def _batch_export_artifact(db: Session, batch: Batch, format: str) -> ExportArtifact:
    schema = db.get(Schema, batch.schema_id)
    if not schema:
        raise HTTPException(status_code=404, detail="Schema not found")
    schema_data = _schema_data(schema)
    field_names = [field["key_name"] for field in schema_data.get("fields", [])]
    rows = [_batch_export_row(item, field_names) for item in _sorted_batch_items(batch.items)]
    payload = {
        "batch_id": batch.id,
        "schema_id": batch.schema_id,
        "schema_name": schema.name,
        "status": _batch_read(batch).status,
        "total_count": batch.total_count,
        "rows": rows,
    }
    field_columns = [column for field_name in field_names for column in _kie_export_columns(field_name)]
    fieldnames = ["filename", "document_id", "job_id", "status", "error_message", *field_columns, "warnings"]
    filename = _export_filename("KIE", schema.name, batch.id, format)
    return _tabular_export_artifact(payload, rows, fieldnames, filename, format)


def _classification_batch_export_artifact(db: Session, batch: ClassificationBatch, format: str) -> ExportArtifact:
    classifier = db.get(DocumentClassifier, batch.classifier_id)
    if not classifier:
        raise HTTPException(status_code=404, detail="Document classifier not found")
    rows = [_classification_batch_export_row(item) for item in _sorted_module_items(batch.items)]
    payload = {
        "batch_id": batch.id,
        "classifier_id": batch.classifier_id,
        "classifier_name": classifier.name,
        "status": _classification_batch_read(batch).status,
        "total_count": batch.total_count,
        "rows": rows,
    }
    fieldnames = [
        "filename",
        "document_id",
        "job_id",
        "status",
        "error_message",
        "classification_status",
        "class_name",
        "confidence",
        "reason",
        "evidence",
    ]
    filename = _export_filename("classification", classifier.name, batch.id, format)
    return _tabular_export_artifact(payload, rows, fieldnames, filename, format)


def _required_field_batch_export_artifact(db: Session, batch: RequiredFieldCheckBatch, format: str) -> ExportArtifact:
    checklist = db.get(RequiredFieldChecklist, batch.checklist_id)
    if not checklist:
        raise HTTPException(status_code=404, detail="Required field checklist not found")
    item_names = [item.get("item_name") for item in _checklist_data(checklist).get("items", []) if item.get("item_name")]
    rows = [_required_field_batch_export_row(item, item_names) for item in _sorted_module_items(batch.items)]
    payload = {
        "batch_id": batch.id,
        "checklist_id": batch.checklist_id,
        "checklist_name": checklist.name,
        "status": _required_field_batch_read(batch).status,
        "total_count": batch.total_count,
        "rows": rows,
    }
    item_columns = [f"{name}_status" for name in item_names] + [f"{name}_evidence" for name in item_names]
    fieldnames = ["filename", "document_id", "job_id", "status", "error_message", "overall_status", *item_columns]
    filename = _export_filename("required_checker", checklist.name, batch.id, format)
    return _tabular_export_artifact(payload, rows, fieldnames, filename, format)


def _tabular_export_artifact(
    payload: dict[str, Any],
    rows: list[dict[str, Any]],
    fieldnames: list[str],
    filename: str,
    format: str,
) -> ExportArtifact:
    if format == "json":
        return ExportArtifact(_json_export_bytes(payload), filename, "application/json")
    if format == "xlsx":
        return ExportArtifact(_xlsx_bytes(rows, fieldnames), filename, XLSX_MIME_TYPE)

    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=fieldnames)
    writer.writeheader()
    for row in rows:
        writer.writerow({key: _csv_cell(row.get(key)) for key in fieldnames})
    return ExportArtifact(_csv_export_bytes(output.getvalue()), filename, "text/csv; charset=utf-8")


def _export_artifact_response(artifact: ExportArtifact) -> Response:
    return Response(
        content=artifact.content,
        media_type=artifact.content_type,
        headers=_download_headers(artifact.filename),
    )


def _json_export_bytes(payload: dict[str, Any]) -> bytes:
    return json.dumps(payload, ensure_ascii=False, default=str).encode("utf-8")


def _csv_export_bytes(content: str) -> bytes:
    return f"\ufeff{content}".encode("utf-8")


def _csv_download_response(content: str, filename: str) -> Response:
    return _export_artifact_response(ExportArtifact(_csv_export_bytes(content), filename, "text/csv; charset=utf-8"))


def _xlsx_bytes(rows: list[dict[str, Any]], fieldnames: list[str]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Export"
    if fieldnames:
        sheet.append(fieldnames)
    for row in rows:
        sheet.append([_xlsx_cell(row.get(field)) for field in fieldnames])

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _xlsx_download_response(rows: list[dict[str, Any]], fieldnames: list[str], filename: str) -> Response:
    return _export_artifact_response(ExportArtifact(_xlsx_bytes(rows, fieldnames), filename, XLSX_MIME_TYPE))


def _xlsx_cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (str, int, float, bool, datetime)):
        return value
    if isinstance(value, list):
        return ";".join(str(item) for item in value)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False, default=str)
    return str(value)


def _table_fieldnames(rows: list[dict[str, Any]]) -> list[str]:
    fieldnames: list[str] = []
    for row in rows:
        for key in row:
            if key not in fieldnames:
                fieldnames.append(key)
    return fieldnames


def _export_preset_read(preset: ExportPreset) -> ExportPresetRead:
    return ExportPresetRead(
        id=preset.id,
        schema_id=preset.schema_id,
        name=preset.name,
        fields=json.loads(preset.fields_json),
        created_at=preset.created_at,
        updated_at=preset.updated_at,
    )


def _apply_export_preset(payload: dict[str, Any], preset: ExportPreset) -> dict[str, Any]:
    fields = [field for field in json.loads(preset.fields_json) if field.get("include", True)]
    if not fields:
        return payload
    values = payload.get("values", {})
    next_values: dict[str, Any] = {}
    for field in fields:
        key_name = field.get("key_name")
        if not key_name or key_name not in values:
            continue
        output_key = field.get("column_name") or key_name
        next_values[output_key] = values[key_name]
    return {**payload, "values": next_values}


def _archive_search(
    db: Session,
    *,
    workspace_id: str | None = None,
    q: str | None,
    status: str | None,
    schema_id: str | None,
    document_type: str | None,
    limit: int,
) -> list[ArchiveSearchResult]:
    normalized_q = (q or "").strip().lower()
    if schema_id:
        schema = db.get(Schema, schema_id)
        _ensure_workspace_scope(schema, workspace_id, "Schema not found")
    documents = _scope_query(db.query(Document), Document, workspace_id).order_by(Document.created_at.desc()).limit(200).all()
    results: list[ArchiveSearchResult] = []
    for document in documents:
        if document_type and document.document_type != document_type:
            continue
        jobs = (
            _scope_query(db.query(ExtractionJob), ExtractionJob, workspace_id)
            .filter(ExtractionJob.document_id == document.id)
            .order_by(ExtractionJob.created_at.desc())
            .all()
        )
        if not jobs:
            if status or schema_id:
                continue
            haystack = " ".join(filter(None, [document.filename, document.document_type, document.language])).lower()
            if normalized_q and normalized_q not in haystack:
                continue
            results.append(
                ArchiveSearchResult(
                    document_id=document.id,
                    filename=document.filename,
                    document_type=document.document_type,
                    language=document.language,
                    created_at=document.created_at,
                    matched_text=document.filename,
                )
            )
            if len(results) >= limit:
                return results
            continue

        for job in jobs:
            if status and job.status != status:
                continue
            if schema_id and job.schema_id != schema_id:
                continue
            schema_name = job.schema.name if job.schema else None
            matched_text = _job_search_text(document, job, schema_name)
            if normalized_q and normalized_q not in matched_text.lower():
                continue
            results.append(
                ArchiveSearchResult(
                    document_id=document.id,
                    filename=document.filename,
                    document_type=document.document_type,
                    language=document.language,
                    job_id=job.id,
                    result_id=job.result_id,
                    schema_id=job.schema_id,
                    schema_name=schema_name,
                    status=job.status,
                    matched_text=matched_text[:240],
                    created_at=job.created_at,
                )
            )
            if len(results) >= limit:
                return results
    return results


def _job_search_text(document: Document, job: ExtractionJob, schema_name: str | None) -> str:
    parts = [document.filename, document.document_type or "", document.language or "", schema_name or "", job.status]
    if job.result:
        payload = json.loads(job.result.corrected_output) if job.result.corrected_output else json.loads(job.result.validated_output)
        for key, value in payload.get("values", {}).items():
            parts.append(str(key))
            if isinstance(value, dict):
                parts.append(str(value.get("value", "")))
                parts.append(str(value.get("evidence", "")))
    return " ".join(parts)


def _audit_event_read(event: AuditEvent) -> AuditEventRead:
    return AuditEventRead(
        id=event.id,
        entity_type=event.entity_type,
        entity_id=event.entity_id,
        action=event.action,
        message=event.message,
        metadata=json.loads(event.metadata_json),
        created_at=event.created_at,
    )


def _apply_security_headers(response: Response) -> None:
    response.headers.setdefault("X-Content-Type-Options", "nosniff")
    response.headers.setdefault("Referrer-Policy", "strict-origin-when-cross-origin")
    response.headers.setdefault("X-Frame-Options", "DENY")
    response.headers.setdefault("Permissions-Policy", "camera=(), microphone=(), geolocation=()")
    response.headers.setdefault(
        "Content-Security-Policy",
        "default-src 'self'; img-src 'self' data: blob:; style-src 'self' 'unsafe-inline'; script-src 'self'; connect-src 'self' http://localhost:* http://127.0.0.1:*",
    )


def _start_retention_cleanup_worker() -> threading.Event | None:
    settings = get_settings()
    retention_hours = settings.resolved_upload_retention_hours
    if retention_hours <= 0:
        return None

    stop_event = threading.Event()

    def _run() -> None:
        while not stop_event.is_set():
            try:
                _cleanup_expired_upload_data()
            except Exception:
                pass
            interval = max(60, int(get_settings().retention_cleanup_interval_seconds or 86400))
            stop_event.wait(interval)

    thread = threading.Thread(target=_run, name="upload-retention-cleanup", daemon=True)
    thread.start()
    return stop_event


def _cleanup_expired_upload_data() -> dict[str, Any]:
    settings = get_settings()
    retention_hours = settings.resolved_upload_retention_hours
    if retention_hours <= 0:
        return {"status": "disabled"}
    cutoff = datetime.utcnow() - timedelta(hours=retention_hours)
    db = SessionLocal()
    try:
        documents = db.query(Document).filter(Document.created_at < cutoff).all()
        raw_rows = db.query(RawExtraction).filter(RawExtraction.created_at < cutoff).all()
        export_jobs = db.query(ExportJob).filter(ExportJob.created_at < cutoff).all()
        paths = _storage_paths_for_cleanup(documents, raw_rows, export_jobs)
        counts = _delete_history_before(db, cutoff)
        db.commit()
    finally:
        db.close()

    removed_paths: list[str] = []
    for path in paths:
        try:
            delete_storage_ref(path)
            removed_paths.append(str(path))
        except Exception:
            pass
    return {"status": "cleaned", "cutoff": cutoff.isoformat(), "counts": counts, "removed_paths": removed_paths}


def _storage_paths_for_cleanup(documents: list[Document], raw_rows: list[RawExtraction], export_jobs: list[ExportJob]) -> set[str]:
    paths: set[str] = set()
    for document in documents:
        if document.storage_path:
            paths.add(_artifact_root(document.storage_path))
        for page in document.pages:
            if page.image_path and is_s3_ref(page.image_path):
                paths.add(page.image_path)
    for raw in raw_rows:
        for ref in [raw.storage_path, raw.pdf_path, raw.html_path]:
            if ref:
                paths.add(_artifact_root(ref))
    for job in export_jobs:
        if job.storage_path:
            paths.add(job.storage_path if is_s3_ref(job.storage_path) else str(Path(job.storage_path).parent))
    return paths


def _artifact_root(ref: str) -> str:
    if is_s3_ref(ref):
        return ref
    path = Path(ref)
    return str(path.parent if path.name.startswith(("original", "preview", "content")) else path)


def _delete_history_before(db: Session, cutoff: datetime) -> dict[str, int]:
    counts: dict[str, int] = {}
    expired_document_ids = [row[0] for row in db.query(Document.id).filter(Document.created_at < cutoff).all()]
    expired_extraction_job_ids = [row[0] for row in db.query(ExtractionJob.id).filter(ExtractionJob.created_at < cutoff).all()]
    expired_classification_job_ids = [row[0] for row in db.query(ClassificationJob.id).filter(ClassificationJob.created_at < cutoff).all()]
    expired_required_job_ids = [row[0] for row in db.query(RequiredFieldCheckJob.id).filter(RequiredFieldCheckJob.created_at < cutoff).all()]

    counts["classification_batch_items"] = db.query(ClassificationBatchItem).filter(ClassificationBatchItem.created_at < cutoff).delete(synchronize_session=False)
    counts["classification_batches"] = db.query(ClassificationBatch).filter(ClassificationBatch.created_at < cutoff).delete(synchronize_session=False)
    if expired_classification_job_ids:
        counts["classification_results"] = db.query(ClassificationResult).filter(ClassificationResult.job_id.in_(expired_classification_job_ids)).delete(synchronize_session=False)
    counts["classification_jobs"] = db.query(ClassificationJob).filter(ClassificationJob.created_at < cutoff).delete(synchronize_session=False)

    counts["required_field_check_batch_items"] = db.query(RequiredFieldCheckBatchItem).filter(RequiredFieldCheckBatchItem.created_at < cutoff).delete(synchronize_session=False)
    counts["required_field_check_batches"] = db.query(RequiredFieldCheckBatch).filter(RequiredFieldCheckBatch.created_at < cutoff).delete(synchronize_session=False)
    if expired_required_job_ids:
        counts["required_field_check_results"] = db.query(RequiredFieldCheckResult).filter(RequiredFieldCheckResult.job_id.in_(expired_required_job_ids)).delete(synchronize_session=False)
    counts["required_field_check_jobs"] = db.query(RequiredFieldCheckJob).filter(RequiredFieldCheckJob.created_at < cutoff).delete(synchronize_session=False)

    counts["workflow_run_items"] = db.query(WorkflowRunItem).filter(WorkflowRunItem.created_at < cutoff).delete(synchronize_session=False)
    counts["workflow_runs"] = db.query(WorkflowRun).filter(WorkflowRun.created_at < cutoff).delete(synchronize_session=False)
    counts["batch_items"] = db.query(BatchItem).filter(BatchItem.created_at < cutoff).delete(synchronize_session=False)
    counts["batches"] = db.query(Batch).filter(Batch.created_at < cutoff).delete(synchronize_session=False)
    if expired_extraction_job_ids:
        counts["extraction_results"] = db.query(ExtractionResult).filter(ExtractionResult.job_id.in_(expired_extraction_job_ids)).delete(synchronize_session=False)
    counts["extraction_jobs"] = db.query(ExtractionJob).filter(ExtractionJob.created_at < cutoff).delete(synchronize_session=False)
    counts["export_jobs"] = db.query(ExportJob).filter(ExportJob.created_at < cutoff).delete(synchronize_session=False)
    if expired_document_ids:
        counts["document_pages"] = db.query(DocumentPage).filter(DocumentPage.document_id.in_(expired_document_ids)).delete(synchronize_session=False)
    counts["documents"] = db.query(Document).filter(Document.created_at < cutoff).delete(synchronize_session=False)
    counts["raw_extractions"] = db.query(RawExtraction).filter(RawExtraction.created_at < cutoff).delete(synchronize_session=False)
    counts["audit_events"] = db.query(AuditEvent).filter(AuditEvent.created_at < cutoff).delete(synchronize_session=False)
    counts["draft_schemas"] = db.query(Schema).filter(Schema.ephemeral == True, Schema.created_at < cutoff).delete(synchronize_session=False)  # noqa: E712
    return counts


def _configure_frontend_static() -> None:
    settings = get_settings()
    if not settings.serve_frontend:
        return

    dist_dir = settings.resolved_frontend_dist_dir
    index_path = dist_dir / "index.html"
    if not index_path.exists():
        return

    assets_dir = dist_dir / "assets"
    if assets_dir.exists():
        app.mount("/assets", StaticFiles(directory=assets_dir), name="frontend_assets")

    @app.get("/", include_in_schema=False)
    def _frontend_index() -> FileResponse:
        return FileResponse(index_path)

    @app.get("/{full_path:path}", include_in_schema=False)
    def _frontend_spa(full_path: str) -> FileResponse:
        if full_path.startswith("api/"):
            raise HTTPException(status_code=404, detail="Not Found")

        requested = (dist_dir / full_path).resolve()
        if requested.is_file() and _is_relative_to(requested, dist_dir.resolve()):
            return FileResponse(requested)
        return FileResponse(index_path)


def _is_relative_to(path: Path, parent: Path) -> bool:
    try:
        path.relative_to(parent)
    except ValueError:
        return False
    return True


_configure_frontend_static()
